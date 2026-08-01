from __future__ import annotations

from collections.abc import Iterable
from dataclasses import dataclass
from datetime import datetime, time, timedelta
from io import BytesIO

from django.utils import timezone
from django.db.models import Q
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from backend.apps.assessments.models import WrongQuestion
from backend.apps.assessments.pdf_fonts import register_pdf_cjk_font


WRONG_EXPORT_SCOPES = {"active", "due", "mastered", "all"}
WRONG_EXPORT_FORMATS = {"docx", "xlsx", "pdf"}
MAX_EXPORT_QUESTIONS = 500
MAX_TEXT_LENGTH = 12_000


class WrongExportRequestError(ValueError):
    pass


@dataclass(frozen=True)
class ExportArtifact:
    content: bytes
    content_type: str
    filename: str


def build_wrong_export(user, scope: str, export_format: str) -> ExportArtifact:
    normalized_format = str(export_format or "").strip().lower()
    if normalized_format not in WRONG_EXPORT_FORMATS:
        raise WrongExportRequestError("导出格式必须是 docx、xlsx 或 pdf")
    items = list(select_wrong_questions(user, scope))
    if not items:
        raise WrongExportRequestError("当前范围没有可导出的错题")
    snapshot = build_wrong_question_snapshot(user, scope, items)
    return render_wrong_export(snapshot, normalized_format)


def render_wrong_export(
    snapshot: dict,
    export_format: str,
    *,
    enhancement: dict | None = None,
    citations: list[dict] | None = None,
    degraded_reasons: list[str] | None = None,
) -> ExportArtifact:
    normalized_format = str(export_format or "").strip().lower()
    if normalized_format not in WRONG_EXPORT_FORMATS:
        raise WrongExportRequestError("导出格式必须是 docx、xlsx 或 pdf")
    date_label = timezone.localdate().strftime("%Y%m%d")
    filename = f"wrong-questions-{snapshot['scope']}-{date_label}.{normalized_format}"
    if normalized_format == "docx":
        return ExportArtifact(
            _build_docx(snapshot, enhancement, citations, degraded_reasons),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            filename,
        )
    if normalized_format == "xlsx":
        return ExportArtifact(
            _build_xlsx(snapshot, enhancement, citations, degraded_reasons),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            filename,
        )
    return ExportArtifact(
        _build_pdf(snapshot, enhancement, citations, degraded_reasons),
        "application/pdf",
        filename,
    )


def select_wrong_questions(user, scope: str):
    normalized_scope = str(scope or "active").strip().lower()
    if normalized_scope not in WRONG_EXPORT_SCOPES:
        raise WrongExportRequestError("导出范围必须是 active、due、mastered 或 all")

    queryset = (
        WrongQuestion.objects.filter(user=user)
        .select_related("question__knowledge_point__chapter__course")
        .order_by("mastered", "review_after", "-wrong_count", "-last_wrong_at")
    )
    if normalized_scope == "active":
        queryset = queryset.filter(mastered=False)
    elif normalized_scope == "due":
        tomorrow = timezone.localdate() + timedelta(days=1)
        tomorrow_start = timezone.make_aware(
            datetime.combine(tomorrow, time.min),
            timezone.get_current_timezone(),
        )
        queryset = queryset.filter(mastered=False).filter(
            Q(review_after__isnull=True) | Q(review_after__lt=tomorrow_start)
        )
    elif normalized_scope == "mastered":
        queryset = queryset.filter(mastered=True)
    return queryset[:MAX_EXPORT_QUESTIONS]


def build_wrong_question_snapshot(user, scope: str, wrong_items: Iterable[WrongQuestion]) -> dict:
    normalized_scope = str(scope or "active").strip().lower()
    if normalized_scope not in WRONG_EXPORT_SCOPES:
        raise WrongExportRequestError("导出范围必须是 active、due、mastered 或 all")

    rows = [_snapshot_row(item) for item in list(wrong_items)[:MAX_EXPORT_QUESTIONS]]
    points = sorted({row["knowledgePoint"] for row in rows if row["knowledgePoint"]})
    return {
        "schemaVersion": "1.0",
        "scope": normalized_scope,
        "studentName": _bounded_text(getattr(user, "display_name", "") or "同学", 64),
        "generatedAt": timezone.now().isoformat(),
        "summary": {
            "totalCount": len(rows),
            "activeCount": sum(1 for row in rows if not row["mastered"]),
            "masteredCount": sum(1 for row in rows if row["mastered"]),
            "knowledgePoints": points,
        },
        "questions": rows,
    }


def _snapshot_row(wrong: WrongQuestion) -> dict:
    question = wrong.question
    point = question.knowledge_point
    chapter = point.chapter
    course = chapter.course
    options = []
    for index, option in enumerate(question.options if isinstance(question.options, list) else []):
        if isinstance(option, dict):
            key = _bounded_text(option.get("key") or chr(65 + index), 16)
            text = _bounded_text(option.get("text"), 2_000)
        else:
            key = chr(65 + index)
            text = _bounded_text(option, 2_000)
        options.append({"key": key, "text": text})
    return {
        "wrongQuestionId": str(wrong.id),
        "questionId": str(question.id),
        "course": _bounded_text(course.title, 128),
        "chapter": _bounded_text(chapter.title, 128),
        "knowledgePoint": _bounded_text(point.title, 128),
        "stem": _bounded_text(question.stem),
        "options": options,
        "lastWrongAnswer": _bounded_text(wrong.last_wrong_answer, 16),
        "correctAnswer": _bounded_text(question.correct_answer, 16),
        "errorCause": _bounded_text(question.error_cause or "待进一步分析", 500),
        "explanation": _bounded_text(question.explanation),
        "wrongCount": wrong.wrong_count,
        "consecutiveCorrect": wrong.consecutive_correct,
        "mastered": wrong.mastered,
        "lastWrongAt": wrong.last_wrong_at.isoformat(),
        "reviewAfter": wrong.review_after.isoformat() if wrong.review_after else None,
        "source": _question_source(question),
    }


def _question_source(question) -> str:
    resource = getattr(question, "source_resource", None)
    if resource is not None:
        return _bounded_text(getattr(resource, "title", "课程资料"), 255)
    return "智学引擎测评题库"


def _bounded_text(value, limit: int = MAX_TEXT_LENGTH) -> str:
    return str(value or "").strip()[:limit]


def _build_docx(
    snapshot: dict,
    enhancement: dict | None = None,
    citations: list[dict] | None = None,
    degraded_reasons: list[str] | None = None,
) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    styles = document.styles
    for name, size in (("Normal", 10.5), ("Title", 22), ("Heading 1", 15), ("Heading 2", 12)):
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.2

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("错题复习清单")
    run.bold = True
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(15, 23, 42)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(
        f"{snapshot['studentName']} · {scope_label(snapshot['scope'])} · "
        f"{_display_datetime(snapshot['generatedAt'])}"
    )

    summary = snapshot["summary"]
    table = document.add_table(rows=2, cols=4)
    table.style = "Table Grid"
    labels = ("错题总数", "待巩固", "已掌握", "重点知识点")
    values = (
        str(summary["totalCount"]),
        str(summary["activeCount"]),
        str(summary["masteredCount"]),
        "、".join(summary["knowledgePoints"][:5]) or "暂无",
    )
    for index, label in enumerate(labels):
        table.cell(0, index).text = label
        table.cell(1, index).text = values[index]
        for paragraph in table.cell(0, index).paragraphs:
            paragraph.runs[0].bold = True
        _shade_cell(table.cell(0, index), "E2E8F0")

    current_point = None
    for number, row in enumerate(snapshot["questions"], start=1):
        if row["knowledgePoint"] != current_point:
            current_point = row["knowledgePoint"]
            document.add_heading(current_point, level=1)
        document.add_heading(f"{number}. {row['stem']}", level=2)
        for option in row["options"]:
            document.add_paragraph(f"{option['key']}. {option['text']}", style="Normal")
        details = (
            f"上次错误答案：{row['lastWrongAnswer'] or '未记录'}    正确答案：{row['correctAnswer']}    "
            f"错误次数：{row['wrongCount']}    下次复习：{_display_datetime(row['reviewAfter'])}"
        )
        detail_paragraph = document.add_paragraph(details)
        detail_paragraph.runs[0].bold = True
        document.add_paragraph(f"错因：{row['errorCause']}")
        document.add_paragraph(f"解析：{row['explanation']}")
        document.add_paragraph(f"来源：{row['source']} · {row['course']} / {row['chapter']}")

    if degraded_reasons:
        notice = document.add_paragraph("基础版：AI 增强暂未完整完成，原题、原答案和原始解析仍可正常复习。")
        notice.runs[0].bold = True
    sections = _enhancement_sections(enhancement)
    if sections:
        document.add_heading("AI 深度诊断", level=1)
    for section_data in sections:
        document.add_heading(_bounded_text(section_data.get("knowledgePoint") or "未分类", 128), level=2)
        document.add_paragraph(f"共性诊断：{_bounded_text(section_data.get('groupDiagnosis'))}")
        gaps = section_data.get("prerequisiteGaps") if isinstance(section_data.get("prerequisiteGaps"), list) else []
        document.add_paragraph(f"前置知识缺口：{'、'.join(_bounded_text(item, 500) for item in gaps) or '未发现明确缺口'}")
        for analysis in section_data.get("questionAnalyses") or []:
            document.add_paragraph(f"题目 {analysis.get('questionId') or '-'}：{_bounded_text(analysis.get('deepErrorDiagnosis'))}")
            for step in analysis.get("solutionSteps") or []:
                document.add_paragraph(_bounded_text(step), style="List Bullet")
            document.add_paragraph(f"自检：{_bounded_text(analysis.get('selfCheck'))}")
            document.add_paragraph(f"复习建议：{_bounded_text(analysis.get('reviewAdvice'))}")
        practices = section_data.get("transferPractice") if isinstance(section_data.get("transferPractice"), list) else []
        if practices:
            document.add_heading("迁移练习", level=2)
        for index, practice in enumerate(practices, start=1):
            document.add_paragraph(f"{index}. {_bounded_text(practice.get('prompt'))}")
            document.add_paragraph(f"答案：{_bounded_text(practice.get('answer'))}")
            document.add_paragraph(f"解析：{_bounded_text(practice.get('explanation'))}")
    if citations:
        document.add_heading("资料引用", level=1)
        for citation in citations:
            document.add_paragraph(
                f"[{_bounded_text(citation.get('sourceId'), 32)}] "
                f"{_bounded_text(citation.get('title'), 255)} · {_bounded_text(citation.get('location'), 500)}\n"
                f"{_bounded_text(citation.get('snippet'), 600)}"
            )

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("由智学引擎生成 · 仅用于个人复习")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


XLSX_HEADERS = [
    "错题ID", "知识点", "章节", "题干", "选项", "上次错误答案", "正确答案",
    "错因", "解析", "错误次数", "连续答对次数", "掌握状态", "上次错误时间",
    "下次复习时间", "题目来源",
]


def _build_xlsx(
    snapshot: dict,
    enhancement: dict | None = None,
    citations: list[dict] | None = None,
    degraded_reasons: list[str] | None = None,
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "错题明细"
    sheet.append(XLSX_HEADERS)
    for row in snapshot["questions"]:
        sheet.append([
            row["wrongQuestionId"],
            row["knowledgePoint"],
            row["chapter"],
            row["stem"],
            "\n".join(f"{item['key']}. {item['text']}" for item in row["options"]),
            row["lastWrongAnswer"],
            row["correctAnswer"],
            row["errorCause"],
            row["explanation"],
            row["wrongCount"],
            row["consecutiveCorrect"],
            "已掌握" if row["mastered"] else "待巩固",
            _excel_datetime(row["lastWrongAt"]),
            _excel_datetime(row["reviewAfter"]),
            f"{row['source']} · {row['course']} / {row['chapter']}",
        ])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(name="Microsoft YaHei", color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    widths = [38, 18, 18, 48, 42, 16, 14, 28, 48, 12, 14, 14, 20, 20, 36]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for date_index in (13, 14):
            row[date_index - 1].number_format = "yyyy-mm-dd hh:mm"
    analysis_sheet = workbook.create_sheet("AI分析")
    analysis_sheet.append([
        "题目ID", "知识点", "共性诊断", "前置知识缺口", "深层错因",
        "解题步骤", "自检", "复习建议", "证据状态", "引用编号",
    ])
    practice_sheet = workbook.create_sheet("迁移练习与引用")
    practice_sheet.append(["记录类型", "知识点", "内容", "答案", "解析/片段", "难度/位置", "引用编号"])
    sections = _enhancement_sections(enhancement)
    for section_data in sections:
        point = _bounded_text(section_data.get("knowledgePoint"), 128)
        gaps = "\n".join(_bounded_text(item, 500) for item in section_data.get("prerequisiteGaps") or [])
        for analysis in section_data.get("questionAnalyses") or []:
            analysis_sheet.append([
                _bounded_text(analysis.get("questionId"), 64), point,
                _bounded_text(section_data.get("groupDiagnosis")), gaps,
                _bounded_text(analysis.get("deepErrorDiagnosis")),
                "\n".join(_bounded_text(item, 500) for item in analysis.get("solutionSteps") or []),
                _bounded_text(analysis.get("selfCheck")), _bounded_text(analysis.get("reviewAdvice")),
                _bounded_text(section_data.get("evidenceStatus"), 32),
                ", ".join(_bounded_text(item, 32) for item in analysis.get("citationIds") or []),
            ])
        for practice in section_data.get("transferPractice") or []:
            practice_sheet.append([
                "迁移练习", point, _bounded_text(practice.get("prompt")),
                _bounded_text(practice.get("answer")), _bounded_text(practice.get("explanation")),
                _bounded_text(practice.get("difficulty"), 32),
                ", ".join(_bounded_text(item, 32) for item in practice.get("citationIds") or []),
            ])
    for citation in citations or []:
        practice_sheet.append([
            "资料引用", _bounded_text(citation.get("knowledgePoint"), 128),
            _bounded_text(citation.get("title"), 255), "",
            _bounded_text(citation.get("snippet"), 600), _bounded_text(citation.get("location"), 500),
            _bounded_text(citation.get("sourceId"), 32),
        ])
    if degraded_reasons and analysis_sheet.max_row == 1:
        analysis_sheet.append(["", "", "基础版", "", "AI 增强暂未完整完成", "", "", "", "unverified", ""])
    for extra_sheet in (analysis_sheet, practice_sheet):
        extra_sheet.freeze_panes = "A2"
        extra_sheet.auto_filter.ref = extra_sheet.dimensions
        for cell in extra_sheet[1]:
            cell.font = Font(name="Microsoft YaHei", color="FFFFFF", bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for column in range(1, extra_sheet.max_column + 1):
            extra_sheet.column_dimensions[get_column_letter(column)].width = 24
        for row in extra_sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Microsoft YaHei", size=10)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _build_pdf(
    snapshot: dict,
    enhancement: dict | None = None,
    citations: list[dict] | None = None,
    degraded_reasons: list[str] | None = None,
) -> bytes:
    pdf_font = register_pdf_cjk_font()
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=18 * mm,
        title="智学引擎错题复习清单",
        author="智学引擎",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ChineseTitle", parent=styles["Title"], fontName=pdf_font, fontSize=20,
        leading=26, alignment=TA_CENTER, textColor=colors.HexColor("#0F172A"), spaceAfter=8,
    )
    heading_style = ParagraphStyle(
        "ChineseHeading", parent=styles["Heading2"], fontName=pdf_font, fontSize=13,
        leading=18, textColor=colors.HexColor("#1F4E78"), spaceBefore=10, spaceAfter=5,
    )
    body_style = ParagraphStyle(
        "ChineseBody", parent=styles["BodyText"], fontName=pdf_font, fontSize=9.5,
        leading=15, textColor=colors.HexColor("#334155"), spaceAfter=4,
    )
    story = [
        Paragraph("错题复习清单", title_style),
        Paragraph(
            _escape_pdf_text(f"{snapshot['studentName']} · {scope_label(snapshot['scope'])} · {_display_datetime(snapshot['generatedAt'])}"),
            body_style,
        ),
        Spacer(1, 5 * mm),
    ]
    summary = snapshot["summary"]
    summary_table = Table(
        [["错题总数", "待巩固", "已掌握", "重点知识点"], [
            summary["totalCount"], summary["activeCount"], summary["masteredCount"],
            Paragraph(_escape_pdf_text("、".join(summary["knowledgePoints"][:5]) or "暂无"), body_style),
        ]],
        colWidths=[26 * mm, 26 * mm, 26 * mm, 88 * mm],
    )
    summary_table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), pdf_font),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E2E8F0")),
        ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#0F172A")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (2, -1), "CENTER"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.extend([summary_table, Spacer(1, 4 * mm)])
    current_point = None
    for number, row in enumerate(snapshot["questions"], start=1):
        if row["knowledgePoint"] != current_point:
            current_point = row["knowledgePoint"]
            story.append(Paragraph(_escape_pdf_text(current_point), heading_style))
        question_parts = [Paragraph(_escape_pdf_text(f"{number}. {row['stem']}"), heading_style)]
        question_parts.extend(
            Paragraph(_escape_pdf_text(f"{option['key']}. {option['text']}"), body_style)
            for option in row["options"]
        )
        question_parts.extend([
            Paragraph(_escape_pdf_text(
                f"上次错误答案：{row['lastWrongAnswer'] or '未记录'}　正确答案：{row['correctAnswer']}　"
                f"错误次数：{row['wrongCount']}　下次复习：{_display_datetime(row['reviewAfter'])}"
            ), body_style),
            Paragraph(_escape_pdf_text(f"错因：{row['errorCause']}"), body_style),
            Paragraph(_escape_pdf_text(f"解析：{row['explanation']}"), body_style),
            Paragraph(_escape_pdf_text(f"来源：{row['source']} · {row['course']} / {row['chapter']}"), body_style),
        ])
        story.append(KeepTogether(question_parts))
        if number < len(snapshot["questions"]):
            story.append(Spacer(1, 2 * mm))

    if degraded_reasons:
        story.append(Paragraph(_escape_pdf_text("基础版：AI 增强暂未完整完成，原始错题内容保持不变。"), heading_style))
    sections = _enhancement_sections(enhancement)
    if sections:
        story.append(Paragraph("AI 深度诊断", title_style))
    for section_data in sections:
        story.append(Paragraph(_escape_pdf_text(section_data.get("knowledgePoint") or "未分类"), heading_style))
        story.append(Paragraph(_escape_pdf_text(f"共性诊断：{section_data.get('groupDiagnosis') or ''}"), body_style))
        for analysis in section_data.get("questionAnalyses") or []:
            story.append(Paragraph(_escape_pdf_text(
                f"题目 {analysis.get('questionId') or '-'}：{analysis.get('deepErrorDiagnosis') or ''}"
            ), body_style))
            for step in analysis.get("solutionSteps") or []:
                story.append(Paragraph(_escape_pdf_text(f"• {step}"), body_style))
        practices = section_data.get("transferPractice") if isinstance(section_data.get("transferPractice"), list) else []
        if practices:
            story.append(Paragraph("迁移练习", heading_style))
        for practice in practices:
            story.append(Paragraph(_escape_pdf_text(f"题目：{practice.get('prompt') or ''}"), body_style))
            story.append(Paragraph(_escape_pdf_text(f"答案：{practice.get('answer') or ''}"), body_style))
            story.append(Paragraph(_escape_pdf_text(f"解析：{practice.get('explanation') or ''}"), body_style))
    if citations:
        story.append(Paragraph("资料引用", title_style))
        for citation in citations:
            story.append(Paragraph(_escape_pdf_text(
                f"[{citation.get('sourceId') or ''}] {citation.get('title') or ''} · "
                f"{citation.get('location') or ''}：{citation.get('snippet') or ''}"
            ), body_style))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont(pdf_font, 8)
        canvas.setFillColor(colors.HexColor("#64748B"))
        canvas.drawCentredString(A4[0] / 2, 9 * mm, f"由智学引擎生成 · 第 {doc.page} 页")
        canvas.restoreState()

    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def _enhancement_sections(enhancement: dict | None) -> list[dict]:
    if not isinstance(enhancement, dict):
        return []
    sections = enhancement.get("sections")
    return [item for item in sections if isinstance(item, dict)] if isinstance(sections, list) else []


def scope_label(scope: str) -> str:
    return {"active": "待巩固", "due": "今日到期", "mastered": "已掌握", "all": "全部错题"}[scope]


def _display_datetime(value) -> str:
    parsed = _parse_datetime(value)
    return timezone.localtime(parsed).strftime("%Y-%m-%d %H:%M") if parsed else "立即复习"


def _excel_datetime(value):
    parsed = _parse_datetime(value)
    return timezone.localtime(parsed).replace(tzinfo=None) if parsed else None


def _parse_datetime(value):
    if not value:
        return None
    if isinstance(value, datetime):
        return value
    try:
        return datetime.fromisoformat(str(value))
    except ValueError:
        return None


def _shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    properties.append(shade)


def _escape_pdf_text(value: str) -> str:
    return str(value).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
