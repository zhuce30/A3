from __future__ import annotations

from datetime import timedelta
from io import BytesIO
import copy
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

import json

from django.test import Client, TestCase, override_settings
from django.utils import timezone
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from backend.apps.accounts.models import User
from backend.apps.assessments.models import Question, WrongQuestion, WrongQuestionExportJob
from backend.apps.assessments.content_enhancer import EnhancementResult
from backend.apps.assessments.wrong_exports import ExportArtifact, WrongExportRequestError
from backend.apps.learning.models import Chapter, Course, KnowledgePoint
from backend.apps.profiles.models import LearnerProfile


class WrongExportDomainTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            phone="13800001111",
            password="TestPassword123",
            display_name="测试同学",
        )
        self.other_user = User.objects.create_user(
            phone="13800002222",
            password="TestPassword123",
            display_name="其他同学",
        )
        LearnerProfile.objects.create(user=self.user, weak_points=["进程调度"])
        course = Course.objects.create(code="EXPORT-001", title="操作系统导出测试")
        chapter = Chapter.objects.create(course=course, order=1, title="处理机管理")
        self.point = KnowledgePoint.objects.create(
            chapter=chapter,
            code="EXPORT-KP-001",
            order=1,
            title="进程调度",
        )
        now = timezone.now()
        later_today = timezone.localtime(now).replace(hour=23, minute=0, second=0, microsecond=0)
        self.future = self._wrong(self.user, "未来复习题", review_after=now + timedelta(days=2))
        self.due = self._wrong(self.user, "今日到期题", review_after=now - timedelta(hours=1), wrong_count=3)
        self.due_later_today = self._wrong(
            self.user,
            "今天稍晚到期题",
            review_after=later_today,
        )
        self.unscheduled = self._wrong(self.user, "未安排复习题", review_after=None)
        self.mastered = self._wrong(self.user, "已掌握题", review_after=now - timedelta(days=2), mastered=True)
        self.other = self._wrong(self.other_user, "其他用户题", review_after=now - timedelta(days=1))

    def _wrong(
        self,
        user,
        stem: str,
        *,
        review_after,
        mastered: bool = False,
        wrong_count: int = 1,
    ) -> WrongQuestion:
        question = Question.objects.create(
            knowledge_point=self.point,
            stem=stem,
            options=[
                {"key": "A", "text": "先来先服务"},
                {"key": "B", "text": "时间片轮转"},
            ],
            correct_answer="B",
            explanation="时间片轮转会周期性地把处理机分配给就绪进程。",
            error_cause="混淆了非抢占式与抢占式调度",
        )
        wrong = WrongQuestion.objects.create(
            user=user,
            question=question,
            wrong_count=wrong_count,
            mastered=mastered,
            last_wrong_answer="A",
            review_after=review_after,
        )
        return wrong

    def test_wrong_export_job_persists_snapshot_audit_and_artifact_metadata(self):
        job = WrongQuestionExportJob.objects.create(
            user=self.user,
            scope="active",
            format="docx",
            snapshot_hash="a" * 64,
            snapshot={"questions": [{"questionId": "q-1"}]},
            enhancement={"sections": []},
            citations=[{"sourceId": "S1"}],
            llm_audit={"promptVersion": "wrong-export-v1"},
            degraded=True,
            degraded_reasons=["timeout"],
            artifact_path="exports/private.docx",
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            filename="wrong-questions-active.docx",
        )
        job.refresh_from_db()
        self.assertEqual(job.user, self.user)
        self.assertEqual(job.status, "queued")
        self.assertEqual(job.progress, 0)
        self.assertEqual(job.enhancement, {"sections": []})
        self.assertEqual(job.llm_audit["promptVersion"], "wrong-export-v1")
        self.assertEqual(job.degraded_reasons, ["timeout"])
        self.assertTrue(job.created_at)
        self.assertTrue(job.updated_at)

    def test_create_wrong_export_job_validates_and_reuses_only_matching_active_job(self):
        from backend.apps.assessments.wrong_export_jobs import create_wrong_export_job

        first, created = create_wrong_export_job(self.user, "active", "docx")
        reused, reused_created = create_wrong_export_job(self.user, "active", "docx")
        xlsx, xlsx_created = create_wrong_export_job(self.user, "active", "xlsx")
        other, other_created = create_wrong_export_job(self.other_user, "active", "docx")

        self.assertTrue(created)
        self.assertFalse(reused_created)
        self.assertEqual(first.id, reused.id)
        self.assertTrue(xlsx_created)
        self.assertNotEqual(first.id, xlsx.id)
        self.assertTrue(other_created)
        self.assertNotEqual(first.id, other.id)
        self.assertNotIn("generatedAt", first.snapshot_hash)
        with self.assertRaises(WrongExportRequestError):
            create_wrong_export_job(self.user, "private", "docx")
        with self.assertRaises(WrongExportRequestError):
            create_wrong_export_job(self.user, "active", "zip")

    def test_retry_wrong_export_job_clones_failed_snapshot(self):
        from backend.apps.assessments.wrong_export_jobs import create_wrong_export_job, retry_wrong_export_job

        original, _ = create_wrong_export_job(self.user, "active", "pdf")
        original.status = "failed"
        original.save(update_fields=["status", "updated_at"])
        retried = retry_wrong_export_job(original)
        self.assertNotEqual(original.id, retried.id)
        self.assertEqual(original.snapshot_hash, retried.snapshot_hash)
        self.assertEqual(original.snapshot, retried.snapshot)

    @patch("backend.apps.assessments.wrong_export_jobs._render_export")
    @patch("backend.apps.assessments.wrong_export_jobs.enhance_wrong_export")
    @patch("backend.apps.assessments.wrong_export_jobs.search_knowledge_base")
    def test_generate_wrong_export_job_persists_enhancement_and_artifact(self, search, enhance, render):
        from backend.apps.assessments.wrong_export_jobs import (
            _save_job,
            create_wrong_export_job,
            generate_wrong_export_job,
        )

        search.return_value = []
        enhance.return_value = EnhancementResult(
            content={"sections": []},
            citations=[],
            verified=False,
            degraded=True,
            degraded_reasons=["configuration"],
            audit={"promptVersion": "wrong-export-v1", "attempts": 1},
        )
        render.return_value = ExportArtifact(b"docx-bytes", "application/test", "wrong.docx")
        with TemporaryDirectory() as directory, override_settings(A3_WRONG_EXPORT_ROOT=directory):
            job, _ = create_wrong_export_job(self.user, "active", "docx")
            with patch("backend.apps.assessments.wrong_export_jobs._save_job", wraps=_save_job) as save_job:
                result = generate_wrong_export_job(str(job.id))
            job.refresh_from_db()
            self.assertTrue(result["success"])
            self.assertEqual(job.status, "succeeded")
            self.assertEqual(job.stage, "completed")
            self.assertEqual(job.progress, 100)
            self.assertTrue(job.degraded)
            self.assertEqual(job.degraded_reasons, ["configuration"])
            self.assertEqual(job.llm_audit["promptVersion"], "wrong-export-v1")
            self.assertEqual(Path(job.artifact_path).read_bytes(), b"docx-bytes")
            stages = [call.kwargs.get("stage") for call in save_job.call_args_list if call.kwargs.get("stage")]
            self.assertEqual(stages, ["retrieval", "enhancing", "rendering", "completed"])

    @patch("backend.apps.assessments.wrong_export_jobs._render_export")
    def test_cancelled_wrong_export_job_is_terminal_and_writes_nothing(self, render):
        from backend.apps.assessments.wrong_export_jobs import create_wrong_export_job, generate_wrong_export_job

        with TemporaryDirectory() as directory, override_settings(A3_WRONG_EXPORT_ROOT=directory):
            job, _ = create_wrong_export_job(self.user, "active", "docx")
            job.status = "cancelled"
            job.save(update_fields=["status", "updated_at"])
            result = generate_wrong_export_job(str(job.id))
            self.assertTrue(result["cancelled"])
            render.assert_not_called()
            self.assertFalse(any(Path(directory).rglob("*.docx")))

    def test_wrong_export_serialization_and_resolution_hide_server_path(self):
        from backend.apps.assessments.wrong_export_jobs import (
            create_wrong_export_job,
            resolve_wrong_export_artifact,
            serialize_wrong_export_job,
        )

        with TemporaryDirectory() as directory, override_settings(A3_WRONG_EXPORT_ROOT=directory):
            job, _ = create_wrong_export_job(self.user, "active", "docx")
            path = Path(directory) / str(self.user.id) / str(job.id) / "wrong.docx"
            path.parent.mkdir(parents=True)
            path.write_bytes(b"ok")
            job.status = "succeeded"
            job.artifact_path = str(path)
            job.filename = "wrong.docx"
            job.save(update_fields=["status", "artifact_path", "filename", "updated_at"])
            payload = serialize_wrong_export_job(job)
            self.assertNotIn(directory, str(payload))
            self.assertIn("downloadUrl", payload)
            self.assertEqual(resolve_wrong_export_artifact(job), path.resolve())

    def test_select_wrong_questions_supports_all_scopes_and_owner_isolation(self):
        from backend.apps.assessments.wrong_exports import select_wrong_questions

        self.assertEqual(
            {item.id for item in select_wrong_questions(self.user, "active")},
            {self.future.id, self.due.id, self.due_later_today.id, self.unscheduled.id},
        )
        self.assertEqual(
            {item.id for item in select_wrong_questions(self.user, "due")},
            {self.due.id, self.due_later_today.id, self.unscheduled.id},
        )
        self.assertEqual(
            {item.id for item in select_wrong_questions(self.user, "mastered")},
            {self.mastered.id},
        )
        self.assertEqual(
            {item.id for item in select_wrong_questions(self.user, "all")},
            {self.future.id, self.due.id, self.due_later_today.id, self.unscheduled.id, self.mastered.id},
        )

    def test_select_wrong_questions_rejects_unknown_scope(self):
        from backend.apps.assessments.wrong_exports import WrongExportRequestError, select_wrong_questions

        with self.assertRaisesMessage(WrongExportRequestError, "导出范围"):
            select_wrong_questions(self.user, "private")

    def test_snapshot_contains_answers_sources_and_stable_summary(self):
        from backend.apps.assessments.wrong_exports import build_wrong_question_snapshot, select_wrong_questions

        snapshot = build_wrong_question_snapshot(
            self.user,
            "active",
            select_wrong_questions(self.user, "active"),
        )

        self.assertEqual(snapshot["schemaVersion"], "1.0")
        self.assertEqual(snapshot["scope"], "active")
        self.assertEqual(snapshot["studentName"], "测试同学")
        self.assertEqual(snapshot["summary"]["totalCount"], 4)
        self.assertEqual(snapshot["summary"]["activeCount"], 4)
        self.assertEqual(snapshot["summary"]["masteredCount"], 0)
        self.assertEqual(snapshot["summary"]["knowledgePoints"], ["进程调度"])
        row = next(item for item in snapshot["questions"] if item["stem"] == "今日到期题")
        self.assertEqual(row["correctAnswer"], "B")
        self.assertEqual(row["lastWrongAnswer"], "A")
        self.assertIn("时间片轮转", row["explanation"])
        self.assertEqual(row["chapter"], "处理机管理")
        self.assertEqual(row["course"], "操作系统导出测试")
        self.assertNotIn("phone", snapshot)

    def test_build_docx_export_produces_openable_chinese_document(self):
        from backend.apps.assessments.wrong_exports import build_wrong_export

        artifact = build_wrong_export(self.user, "active", "docx")
        document = Document(BytesIO(artifact.content))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)

        self.assertEqual(
            artifact.content_type,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        self.assertTrue(artifact.filename.endswith(".docx"))
        self.assertNotIn("/", artifact.filename)
        self.assertIn("错题复习清单", text)
        self.assertIn("今日到期题", text)
        self.assertIn("正确答案：B", text)
        self.assertIn("时间片轮转", text)

    def test_render_ai_enhanced_exports_preserves_originals_and_adds_stable_sections(self):
        from backend.apps.assessments.wrong_exports import (
            build_wrong_question_snapshot,
            render_wrong_export,
            select_wrong_questions,
        )

        snapshot = build_wrong_question_snapshot(self.user, "active", select_wrong_questions(self.user, "active"))
        enhancement = {
            "sections": [{
                "knowledgePoint": "进程调度",
                "groupDiagnosis": "调度规则辨析不清。",
                "prerequisiteGaps": ["就绪队列"],
                "confusions": [],
                "questionAnalyses": [{
                    "questionId": snapshot["questions"][0]["questionId"],
                    "deepErrorDiagnosis": "把轮转误认为固定优先级。",
                    "solutionSteps": ["识别算法", "模拟队列"],
                    "selfCheck": "检查每个进程是否轮流执行。",
                    "reviewAdvice": "隔天重新练习。",
                    "citationIds": ["S1"],
                }],
                "transferPractice": [{
                    "prompt": "模拟三个进程的首轮顺序。",
                    "answer": "P1、P2、P3",
                    "explanation": "按就绪队列轮转。",
                    "difficulty": "基础",
                    "citationIds": ["S1"],
                }],
                "evidenceStatus": "verified",
            }]
        }
        citations = [{"sourceId": "S1", "title": "操作系统资料", "location": "第 3 页", "snippet": "时间片轮转说明。"}]
        original_snapshot = copy.deepcopy(snapshot)
        original_enhancement = copy.deepcopy(enhancement)

        docx = render_wrong_export(snapshot, "docx", enhancement=enhancement, citations=citations)
        document = Document(BytesIO(docx.content))
        docx_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("正确答案：B", docx_text)
        self.assertIn("AI 深度诊断", docx_text)
        self.assertIn("迁移练习", docx_text)
        self.assertIn("资料引用", docx_text)

        xlsx = render_wrong_export(snapshot, "xlsx", enhancement=enhancement, citations=citations)
        workbook = load_workbook(BytesIO(xlsx.content))
        self.assertEqual(workbook.sheetnames, ["错题明细", "AI分析", "迁移练习与引用"])
        self.assertIn("正确答案", [cell.value for cell in workbook["错题明细"][1]])
        self.assertIn("深层错因", [cell.value for cell in workbook["AI分析"][1]])

        pdf = render_wrong_export(snapshot, "pdf", enhancement=enhancement, citations=citations)
        pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(pdf.content)).pages)
        self.assertIn("正确答案", pdf_text)
        self.assertIn("AI 深度诊断", pdf_text)
        self.assertIn("资料引用", pdf_text)
        self.assertEqual(snapshot, original_snapshot)
        self.assertEqual(enhancement, original_enhancement)

    def test_render_degraded_export_has_visible_notice(self):
        from backend.apps.assessments.wrong_exports import (
            build_wrong_question_snapshot,
            render_wrong_export,
            select_wrong_questions,
        )

        snapshot = build_wrong_question_snapshot(self.user, "active", select_wrong_questions(self.user, "active"))
        artifact = render_wrong_export(
            snapshot,
            "docx",
            enhancement={"sections": []},
            degraded_reasons=["timeout"],
        )
        document = Document(BytesIO(artifact.content))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("基础版", text)

    def test_build_xlsx_export_has_exact_headers_and_typed_dates(self):
        from backend.apps.assessments.wrong_exports import build_wrong_export

        artifact = build_wrong_export(self.user, "active", "xlsx")
        workbook = load_workbook(BytesIO(artifact.content))
        sheet = workbook["错题明细"]
        headers = [cell.value for cell in sheet[1]]

        self.assertEqual(artifact.content_type, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertEqual(headers[0:5], ["错题ID", "知识点", "章节", "题干", "选项"])
        self.assertIn("正确答案", headers)
        self.assertIn("下次复习时间", headers)
        self.assertEqual(sheet.max_row, 5)
        review_column = headers.index("下次复习时间") + 1
        stem_column = headers.index("题干") + 1
        due_row = next(
            row
            for row in range(2, sheet.max_row + 1)
            if sheet.cell(row=row, column=stem_column).value == "今日到期题"
        )
        self.assertIsNotNone(sheet.cell(row=due_row, column=review_column).value)
        self.assertEqual(sheet.freeze_panes, "A2")

    def test_build_pdf_export_is_parseable_and_has_metadata(self):
        from backend.apps.assessments.wrong_exports import build_wrong_export

        artifact = build_wrong_export(self.user, "active", "pdf")
        reader = PdfReader(BytesIO(artifact.content))

        self.assertEqual(artifact.content_type, "application/pdf")
        self.assertTrue(artifact.filename.endswith(".pdf"))
        self.assertGreaterEqual(len(reader.pages), 1)
        self.assertEqual(reader.metadata.title, "智学引擎错题复习清单")
        self.assertGreater(len(artifact.content), 1000)

    def test_build_wrong_export_rejects_invalid_format_and_empty_scope(self):
        from backend.apps.assessments.wrong_exports import WrongExportRequestError, build_wrong_export

        with self.assertRaisesMessage(WrongExportRequestError, "导出格式"):
            build_wrong_export(self.user, "active", "csv")
        WrongQuestion.objects.filter(user=self.user).update(mastered=False)
        with self.assertRaisesMessage(WrongExportRequestError, "没有可导出的错题"):
            build_wrong_export(self.user, "mastered", "docx")


class WrongExportApiTests(WrongExportDomainTests):
    def setUp(self):
        super().setUp()
        self.client = Client()
        self.client.force_login(self.user)

    def test_export_requires_authentication(self):
        self.client.logout()
        response = self.client.post(
            "/api/assessments/wrong-questions/export/",
            data=json.dumps({"scope": "active", "format": "docx"}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 401)
        self.assertFalse(response.json()["success"])

    def test_export_rejects_invalid_payload_and_empty_scope(self):
        invalid = self.client.post(
            "/api/assessments/wrong-questions/export/",
            data=json.dumps({"scope": "private", "format": "csv"}),
            content_type="application/json",
        )
        self.assertEqual(invalid.status_code, 400)
        self.assertFalse(invalid.json()["success"])
        WrongQuestion.objects.filter(user=self.user).update(mastered=False)
        empty = self.client.post(
            "/api/assessments/wrong-questions/export/",
            data=json.dumps({"scope": "mastered", "format": "pdf"}),
            content_type="application/json",
        )
        self.assertEqual(empty.status_code, 409)
        self.assertIn("没有可导出的错题", empty.json()["error"])

    def test_export_returns_real_binary_files_with_safe_headers(self):
        expected = {
            "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "pdf": "application/pdf",
        }
        for export_format, content_type in expected.items():
            with self.subTest(export_format=export_format):
                response = self.client.post(
                    "/api/assessments/wrong-questions/export/",
                    data=json.dumps({"scope": "active", "format": export_format}),
                    content_type="application/json",
                )
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response["Content-Type"], content_type)
                self.assertIn(f".{export_format}", response["Content-Disposition"])
                self.assertNotIn("D:\\", response["Content-Disposition"])
                self.assertEqual(response["X-Content-Type-Options"], "nosniff")
                self.assertGreater(len(response.content), 1000)
