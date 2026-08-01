from __future__ import annotations

from datetime import timedelta
import json
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from django.test import Client, TestCase, override_settings
from django.utils import timezone
from docx import Document
from pypdf import PdfReader

from backend.apps.accounts.models import User
from backend.apps.assessments.content_enhancer import EnhancementResult
from backend.apps.assessments.content_prompts import REVIEW_HANDOUT_PROMPT_VERSION
from backend.apps.assessments.models import Question, ReviewHandoutJob, WrongQuestion
from backend.apps.learning.models import Chapter, Course, KnowledgePoint
from backend.apps.profiles.models import LearnerProfile
from backend.domains.knowledge_base.schemas import KnowledgeBaseSearchResult


class ReviewHandoutDomainTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            phone="13800003333",
            password="TestPassword123",
            display_name="讲义同学",
        )
        LearnerProfile.objects.create(user=self.user, weak_points=["死锁"])
        course = Course.objects.create(code="HANDOUT-001", title="操作系统")
        chapter = Chapter.objects.create(course=course, order=1, title="进程管理")
        self.deadlock = KnowledgePoint.objects.create(
            chapter=chapter, code="HANDOUT-KP-1", order=1, title="死锁"
        )
        self.scheduling = KnowledgePoint.objects.create(
            chapter=chapter, code="HANDOUT-KP-2", order=2, title="进程调度"
        )
        self._wrong(
            self.deadlock,
            "死锁条件题",
            wrong_count=4,
            review_after=timezone.now() - timedelta(days=2),
        )
        self._wrong(
            self.scheduling,
            "调度算法题",
            wrong_count=2,
            review_after=timezone.now() + timedelta(days=2),
        )

    def _wrong(self, point, stem, *, wrong_count, review_after):
        question = Question.objects.create(
            knowledge_point=point,
            stem=stem,
            options=[
                {"key": "A", "text": "选项一"},
                {"key": "B", "text": "选项二"},
            ],
            correct_answer="B",
            explanation=f"{point.title}的可信题目解析。",
            error_cause=f"{point.title}概念混淆",
        )
        return WrongQuestion.objects.create(
            user=self.user,
            question=question,
            wrong_count=wrong_count,
            last_wrong_answer="A",
            review_after=review_after,
        )

    def _enhancement(self, *, degraded_second=False, total_degraded=False):
        sections = []
        citations = []
        for index, point in enumerate(("死锁", "进程调度"), start=1):
            degraded = total_degraded or (degraded_second and index == 2)
            if degraded:
                sections.append({
                    "knowledgePoint": point,
                    "learningDiagnosis": "模型增强暂不可用，本节根据可信原题解析生成基础内容。",
                    "coreConcept": f"{point}的可信题目解析。",
                    "prerequisites": [],
                    "commonErrorPatterns": [f"{point}概念混淆"],
                    "method": ["阅读原题解析", "遮住答案重新作答"],
                    "workedExample": {
                        "prompt": f"{point}原题",
                        "steps": [f"{point}的可信题目解析。"],
                        "answer": "B",
                    },
                    "tieredPractice": [],
                    "reviewActions": ["完成一次重练"],
                    "citationIds": [],
                    "evidenceStatus": "unverified",
                })
                continue
            source_id = f"S{index}"
            sections.append({
                "knowledgePoint": point,
                "learningDiagnosis": f"{point}条件识别不完整。",
                "coreConcept": f"{point}核心概念。",
                "prerequisites": ["基本定义"],
                "commonErrorPatterns": ["忽略必要条件"],
                "method": ["识别条件", "逐项核对", "验证结论"],
                "workedExample": {
                    "prompt": f"{point}例题",
                    "steps": ["先列条件", "再判断"],
                    "answer": "示例答案",
                },
                "tieredPractice": [
                    {"level": "基础", "prompt": "基础练习", "answer": "基础答案", "explanation": "基础解析"},
                    {"level": "巩固", "prompt": "巩固练习", "answer": "巩固答案", "explanation": "巩固解析"},
                    {"level": "迁移", "prompt": "迁移练习", "answer": "迁移答案", "explanation": "迁移解析"},
                ],
                "reviewActions": ["当天重练", "两天后复查"],
                "citationIds": [source_id],
                "evidenceStatus": "verified",
            })
            citations.append({
                "sourceId": source_id,
                "knowledgePoint": point,
                "title": "操作系统课程讲义",
                "location": f"第 {17 + index} 页",
                "snippet": f"{point}课程资料证据。",
                "score": 0.91,
            })
        degraded = total_degraded or degraded_second
        return EnhancementResult(
            content={"sections": sections},
            citations=citations,
            verified=not degraded,
            degraded=degraded,
            degraded_reasons=["authentication"] if degraded else [],
            audit={
                "promptVersion": REVIEW_HANDOUT_PROMPT_VERSION,
                "provider": "openai-compatible",
                "model": "assessment-model",
                "attempts": 2,
                "latencyMs": 42,
                "finishReason": "stop",
                "usage": {"totalTokens": 321},
                "errorCategories": ["authentication"] if degraded else [],
            },
        )

    def test_review_handout_job_ai_audit_fields_have_backward_compatible_defaults(self):
        job = ReviewHandoutJob.objects.create(
            user=self.user,
            snapshot_hash="b" * 64,
            snapshot={},
            ranked_weak_points=[],
        )
        self.assertEqual(job.enhancement, {})
        self.assertEqual(job.llm_audit, {})
        self.assertEqual(job.degraded_reasons, [])
        self.assertEqual(job.prompt_version, "")

    def test_rank_weak_points_uses_wrong_history_due_state_and_profile(self):
        from backend.apps.assessments.review_handouts import create_snapshot, rank_weak_points

        snapshot = create_snapshot(self.user, "active")
        ranked = rank_weak_points(snapshot)

        self.assertEqual(ranked[0]["knowledgePoint"], "死锁")
        self.assertGreater(ranked[0]["score"], ranked[1]["score"])
        self.assertEqual(ranked[0]["activeCount"], 1)
        self.assertEqual(ranked[0]["wrongCount"], 4)
        self.assertTrue(ranked[0]["profileMatched"])
        self.assertTrue(ranked[0]["dueCount"])

    def test_create_review_handout_job_reuses_active_identical_snapshot(self):
        from backend.apps.assessments.review_handouts import create_review_handout_job

        first, first_created = create_review_handout_job(self.user, "active")
        second, second_created = create_review_handout_job(self.user, "active")

        self.assertTrue(first_created)
        self.assertFalse(second_created)
        self.assertEqual(first.id, second.id)
        self.assertEqual(first.status, "queued")
        self.assertEqual(first.snapshot_hash, second.snapshot_hash)
        self.assertEqual(first.snapshot["studentName"], "讲义同学")
        self.assertEqual(len(first.ranked_weak_points), 2)

    def test_serialize_review_handout_job_never_exposes_server_paths(self):
        from backend.apps.assessments.review_handouts import (
            create_review_handout_job,
            serialize_review_handout_job,
        )

        job, _ = create_review_handout_job(self.user, "active")
        job.status = "succeeded"
        job.docx_path = "D:/private/review.docx"
        job.pdf_path = "D:/private/review.pdf"
        job.save(update_fields=["status", "docx_path", "pdf_path", "updated_at"])

        payload = serialize_review_handout_job(job)
        rendered = str(payload)
        self.assertNotIn("D:/private", rendered)
        self.assertTrue(payload["docxDownloadUrl"].endswith("/download/docx/"))
        self.assertTrue(payload["pdfDownloadUrl"].endswith("/download/pdf/"))

    @patch("backend.apps.assessments.review_handouts.enhance_review_handout")
    @patch("backend.apps.assessments.review_handouts.search_knowledge_base")
    def test_generate_review_handout_creates_verified_docx_and_pdf(self, search, enhance):
        from backend.apps.assessments.review_handouts import create_review_handout_job, generate_review_handout

        search.return_value = [
            KnowledgeBaseSearchResult(
                score=0.91,
                text="死锁形成需要互斥、请求保持、不可剥夺和循环等待四个必要条件。",
                metadata={"title": "操作系统课程讲义", "page_number": 18, "source_path": "course/os.pdf"},
            )
        ]
        enhance.return_value = self._enhancement()
        job, _ = create_review_handout_job(self.user, "active")
        with TemporaryDirectory() as artifact_root, override_settings(A3_REVIEW_HANDOUT_ROOT=artifact_root):
            result = generate_review_handout(str(job.id))
            job.refresh_from_db()

            self.assertTrue(result["success"])
            self.assertEqual(job.status, "succeeded")
            self.assertEqual(job.progress, 100)
            self.assertTrue(job.verified)
            self.assertFalse(job.degraded)
            self.assertEqual(job.enhancement["sections"][0]["learningDiagnosis"], "死锁条件识别不完整。")
            self.assertEqual(job.llm_audit["model"], "assessment-model")
            self.assertEqual(job.prompt_version, REVIEW_HANDOUT_PROMPT_VERSION)
            self.assertEqual(job.degraded_reasons, [])
            self.assertEqual(len(job.citations), 2)
            self.assertTrue(Path(job.docx_path).is_file())
            self.assertTrue(Path(job.pdf_path).is_file())
            text = "\n".join(paragraph.text for paragraph in Document(job.docx_path).paragraphs)
            self.assertIn("个性化复习讲义", text)
            self.assertIn("核心概念", text)
            self.assertIn("学习诊断", text)
            self.assertIn("解题方法", text)
            self.assertIn("基础练习", text)
            self.assertIn("巩固练习", text)
            self.assertIn("迁移练习", text)
            self.assertIn("7 天复习计划", text)
            self.assertIn("资料引用", text)
            self.assertIn("答案与解析", text)
            self.assertGreaterEqual(len(PdfReader(job.pdf_path).pages), 1)
        search.assert_called()
        enhance.assert_called_once()

    @patch("backend.apps.assessments.review_handouts.enhance_review_handout")
    @patch("backend.apps.assessments.review_handouts.search_knowledge_base", return_value=[])
    def test_partial_model_failure_keeps_verified_and_degraded_sections(self, _search, enhance):
        from backend.apps.assessments.review_handouts import create_review_handout_job, generate_review_handout

        enhance.return_value = self._enhancement(degraded_second=True)
        job, _ = create_review_handout_job(self.user, "active")
        with TemporaryDirectory() as artifact_root, override_settings(A3_REVIEW_HANDOUT_ROOT=artifact_root):
            result = generate_review_handout(str(job.id))
            job.refresh_from_db()

            self.assertTrue(result["success"])
            self.assertEqual(job.status, "succeeded")
            self.assertTrue(job.degraded)
            self.assertFalse(job.verified)
            self.assertIn("authentication", job.degraded_reasons)
            self.assertTrue(set(job.degraded_reasons).issubset({
                "authentication", "configuration", "disabled", "empty_content", "invalid_response",
                "missing_evidence", "provider_error", "rate_limit", "retrieval_error", "timeout",
            }))
            self.assertNotIn("secret", str(job.degraded_reasons).lower())
            text = "\n".join(paragraph.text for paragraph in Document(job.docx_path).paragraphs)
            self.assertIn("已核验", text)
            self.assertIn("未核验", text)
            self.assertIn("进程调度的可信题目解析", text)

    @patch("backend.apps.assessments.review_handouts.enhance_review_handout")
    @patch("backend.apps.assessments.review_handouts.search_knowledge_base", return_value=[])
    def test_total_model_failure_still_generates_docx_and_pdf(self, _search, enhance):
        from backend.apps.assessments.review_handouts import create_review_handout_job, generate_review_handout

        enhance.return_value = self._enhancement(total_degraded=True)
        job, _ = create_review_handout_job(self.user, "active")
        with TemporaryDirectory() as artifact_root, override_settings(A3_REVIEW_HANDOUT_ROOT=artifact_root):
            result = generate_review_handout(str(job.id))
            job.refresh_from_db()

            self.assertTrue(result["success"])
            self.assertTrue(Path(job.docx_path).is_file())
            self.assertTrue(Path(job.pdf_path).is_file())
            self.assertTrue(job.degraded)
            text = "\n".join(paragraph.text for paragraph in Document(job.docx_path).paragraphs)
            self.assertIn("死锁的可信题目解析", text)
            self.assertIn("进程调度的可信题目解析", text)

    def test_seven_day_plan_prioritizes_active_due_weaknesses(self):
        from backend.apps.assessments.review_handouts import build_seven_day_plan

        ranked = [
            {"knowledgePoint": "死锁", "score": 90, "activeCount": 1, "dueCount": 1},
            {"knowledgePoint": "进程调度", "score": 70, "activeCount": 1, "dueCount": 0},
            {"knowledgePoint": "内存管理", "score": 30, "activeCount": 0, "dueCount": 0},
        ]
        plan = build_seven_day_plan(ranked)

        self.assertEqual([item["day"] for item in plan], list(range(1, 8)))
        early_points = {task["knowledgePoint"] for day in plan[:2] for task in day["tasks"]}
        self.assertIn("死锁", early_points)
        self.assertIn("进程调度", early_points)
        self.assertTrue(any(task["type"] == "spaced_review" for day in plan[2:] for task in day["tasks"]))

    @patch("backend.apps.assessments.review_handouts.search_knowledge_base", side_effect=RuntimeError("retriever offline"))
    def test_generate_review_handout_degrades_to_trusted_question_explanations(self, _search):
        from backend.apps.assessments.review_handouts import create_review_handout_job, generate_review_handout

        job, _ = create_review_handout_job(self.user, "active")
        with TemporaryDirectory() as artifact_root, override_settings(A3_REVIEW_HANDOUT_ROOT=artifact_root):
            result = generate_review_handout(str(job.id))
            job.refresh_from_db()

            self.assertTrue(result["success"])
            self.assertEqual(job.status, "succeeded")
            self.assertTrue(job.degraded)
            self.assertFalse(job.verified)
            self.assertIn("资料检索暂不可用", job.message)
            text = "\n".join(paragraph.text for paragraph in Document(job.docx_path).paragraphs)
            self.assertIn("可信题目解析", text)

    @patch("backend.apps.assessments.review_handouts.search_knowledge_base")
    def test_cancelled_job_is_terminal_and_late_worker_writes_nothing(self, search):
        from backend.apps.assessments.review_handouts import create_review_handout_job, generate_review_handout

        job, _ = create_review_handout_job(self.user, "active")
        job.status = "cancelled"
        job.progress = 100
        job.save(update_fields=["status", "progress", "updated_at"])
        with TemporaryDirectory() as artifact_root, override_settings(A3_REVIEW_HANDOUT_ROOT=artifact_root):
            result = generate_review_handout(str(job.id))
            job.refresh_from_db()

            self.assertTrue(result["cancelled"])
            self.assertEqual(job.status, "cancelled")
            self.assertFalse(job.docx_path)
            self.assertFalse(job.pdf_path)
            self.assertEqual(list(Path(artifact_root).rglob("*")), [])
        search.assert_not_called()


class ReviewHandoutApiTests(ReviewHandoutDomainTests):
    def setUp(self):
        super().setUp()
        self.client = Client()
        self.client.force_login(self.user)

    @patch("backend.apps.assessments.views.current_app.send_task")
    def test_create_list_and_detail_dispatch_to_resource_generation_queue(self, send_task):
        send_task.return_value.id = "celery-handout-001"
        created = self.client.post(
            "/api/assessments/wrong-questions/review-handouts/",
            data=json.dumps({"scope": "active"}),
            content_type="application/json",
        )
        self.assertEqual(created.status_code, 202)
        job = created.json()["job"]
        send_task.assert_called_once_with(
            "assessments.generate_review_handout",
            args=[job["jobId"]],
            queue="resource-generation",
        )
        listed = self.client.get("/api/assessments/wrong-questions/review-handouts/")
        detail = self.client.get(f"/api/assessments/wrong-questions/review-handouts/{job['jobId']}/")
        self.assertEqual(listed.status_code, 200)
        self.assertEqual(listed.json()["jobs"][0]["jobId"], job["jobId"])
        self.assertEqual(detail.status_code, 200)
        self.assertEqual(detail.json()["job"]["snapshotId"], job["snapshotId"])

    @patch("backend.apps.assessments.views.current_app.control.revoke")
    def test_cancel_is_idempotent_and_revokes_worker(self, revoke):
        from backend.apps.assessments.review_handouts import create_review_handout_job

        job, _ = create_review_handout_job(self.user, "active")
        job.celery_task_id = "celery-handout-cancel"
        job.save(update_fields=["celery_task_id", "updated_at"])
        url = f"/api/assessments/wrong-questions/review-handouts/{job.id}/cancel/"
        first = self.client.post(url, data="{}", content_type="application/json")
        second = self.client.post(url, data="{}", content_type="application/json")
        self.assertEqual(first.status_code, 200)
        self.assertEqual(second.status_code, 200)
        self.assertEqual(second.json()["job"]["status"], "cancelled")
        revoke.assert_called_once_with("celery-handout-cancel", terminate=True)

    @patch("backend.apps.assessments.views.current_app.send_task")
    def test_retry_creates_new_job_from_failed_snapshot(self, send_task):
        from backend.apps.assessments.review_handouts import create_review_handout_job

        send_task.return_value.id = "celery-handout-retry"
        job, _ = create_review_handout_job(self.user, "active")
        job.status = "failed"
        job.error_message = "旧任务失败"
        job.save(update_fields=["status", "error_message", "updated_at"])
        response = self.client.post(
            f"/api/assessments/wrong-questions/review-handouts/{job.id}/retry/",
            data="{}",
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 202)
        self.assertNotEqual(response.json()["job"]["jobId"], str(job.id))

    def test_other_user_cannot_read_cancel_or_download_job(self):
        from backend.apps.assessments.review_handouts import create_review_handout_job

        job, _ = create_review_handout_job(self.user, "active")
        other = User.objects.create_user(phone="13800004444", password="TestPassword123")
        self.client.force_login(other)
        base = f"/api/assessments/wrong-questions/review-handouts/{job.id}"
        self.assertEqual(self.client.get(f"{base}/").status_code, 404)
        self.assertEqual(self.client.post(f"{base}/cancel/", data="{}", content_type="application/json").status_code, 404)
        self.assertEqual(self.client.get(f"{base}/download/docx/").status_code, 404)

    def test_download_returns_owned_completed_artifacts(self):
        from backend.apps.assessments.review_handouts import create_review_handout_job, generate_review_handout

        job, _ = create_review_handout_job(self.user, "active")
        with TemporaryDirectory() as artifact_root, override_settings(A3_REVIEW_HANDOUT_ROOT=artifact_root):
            with patch("backend.apps.assessments.review_handouts.search_knowledge_base", return_value=[]):
                generate_review_handout(str(job.id))
            job.refresh_from_db()
            docx = self.client.get(f"/api/assessments/wrong-questions/review-handouts/{job.id}/download/docx/")
            pdf = self.client.get(f"/api/assessments/wrong-questions/review-handouts/{job.id}/download/pdf/")
            self.assertEqual(docx.status_code, 200)
            self.assertEqual(pdf.status_code, 200)
            self.assertIn("review-handout.docx", docx["Content-Disposition"])
            self.assertEqual(pdf["Content-Type"], "application/pdf")
            self.assertGreater(sum(len(chunk) for chunk in docx.streaming_content), 1000)
            self.assertGreater(sum(len(chunk) for chunk in pdf.streaming_content), 1000)
            for response in (docx, pdf):
                for closer in response._resource_closers:
                    closer()
                response._resource_closers.clear()
