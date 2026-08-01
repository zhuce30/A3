from __future__ import annotations

import hashlib
import json
import logging
import os
import shutil
import uuid
from collections import defaultdict
from io import BytesIO
from pathlib import Path

from django.conf import settings
from django.db import transaction
from django.utils import timezone
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from backend.apps.assessments.content_enhancer import (
    EnhancementCancelled,
    enhance_review_handout,
)
from backend.apps.assessments.models import ReviewHandoutJob
from backend.apps.assessments.pdf_fonts import register_pdf_cjk_font
from backend.apps.assessments.wrong_exports import (
    WrongExportRequestError,
    build_wrong_question_snapshot,
    select_wrong_questions,
)
from backend.domains.knowledge_base.retriever import search_knowledge_base


ACTIVE_JOB_STATUSES = {"queued", "running"}
TERMINAL_JOB_STATUSES = {"succeeded", "failed", "cancelled"}
SAFE_DEGRADED_REASONS = {
    "authentication",
    "configuration",
    "disabled",
    "empty_content",
    "invalid_response",
    "claim_not_supported",
    "missing_evidence",
    "provider_error",
    "rate_limit",
    "retrieval_error",
    "timeout",
}
logger = logging.getLogger(__name__)


class ReviewHandoutCancelled(Exception):
    pass


def create_snapshot(user, scope: str) -> dict:
    items = list(select_wrong_questions(user, scope))
    if not items:
        raise WrongExportRequestError("当前范围没有可生成讲义的错题")
    snapshot = build_wrong_question_snapshot(user, scope, items)
    profile = getattr(user, "learner_profile", None)
    weak_points = profile.weak_points if profile and isinstance(profile.weak_points, list) else []
    snapshot["profileWeakPoints"] = [str(item).strip()[:128] for item in weak_points if str(item).strip()][:20]
    return snapshot


def rank_weak_points(snapshot: dict) -> list[dict]:
    now = timezone.now()
    profile_points = set(snapshot.get("profileWeakPoints") or [])
    grouped: dict[str, list[dict]] = defaultdict(list)
    for row in snapshot.get("questions") or []:
        point = str(row.get("knowledgePoint") or "未分类")
        grouped[point].append(row)
    ranked = []
    for point, rows in grouped.items():
        active_count = sum(1 for row in rows if not row.get("mastered"))
        wrong_count = sum(max(0, int(row.get("wrongCount") or 0)) for row in rows)
        due_count = sum(1 for row in rows if _is_due(row.get("reviewAfter"), now))
        profile_matched = point in profile_points
        recency = max((_recency_score(row.get("lastWrongAt"), now) for row in rows), default=0)
        score = round(active_count * 20 + wrong_count * 8 + due_count * 15 + recency + (18 if profile_matched else 0), 2)
        ranked.append({
            "knowledgePoint": point,
            "activeCount": active_count,
            "wrongCount": wrong_count,
            "dueCount": due_count,
            "profileMatched": profile_matched,
            "score": score,
            "questionIds": [row.get("wrongQuestionId") for row in rows],
        })
    return sorted(ranked, key=lambda item: (-item["score"], item["knowledgePoint"]))[:12]


@transaction.atomic
def create_review_handout_job(user, scope: str) -> tuple[ReviewHandoutJob, bool]:
    snapshot = create_snapshot(user, scope)
    digest_payload = {key: value for key, value in snapshot.items() if key != "generatedAt"}
    snapshot_hash = hashlib.sha256(
        json.dumps(digest_payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    ).hexdigest()
    existing = ReviewHandoutJob.objects.select_for_update().filter(
        user=user,
        scope=snapshot["scope"],
        snapshot_hash=snapshot_hash,
        status__in=ACTIVE_JOB_STATUSES,
    ).first()
    if existing:
        return existing, False
    job = ReviewHandoutJob.objects.create(
        user=user,
        scope=snapshot["scope"],
        snapshot_hash=snapshot_hash,
        snapshot=snapshot,
        ranked_weak_points=rank_weak_points(snapshot),
        status="queued",
        stage="queued",
        message="复习讲义已进入生成队列",
    )
    return job, True


def retry_review_handout_job(job: ReviewHandoutJob) -> ReviewHandoutJob:
    if job.status not in {"failed", "cancelled"}:
        raise WrongExportRequestError("只有失败或已取消的讲义任务可以重试")
    return ReviewHandoutJob.objects.create(
        user=job.user,
        scope=job.scope,
        snapshot_hash=job.snapshot_hash,
        snapshot=job.snapshot,
        ranked_weak_points=job.ranked_weak_points,
        status="queued",
        stage="queued",
        message="复习讲义已重新进入生成队列",
    )


def serialize_review_handout_job(job: ReviewHandoutJob) -> dict:
    base = f"/api/assessments/wrong-questions/review-handouts/{job.id}"
    payload = {
        "jobId": str(job.id),
        "scope": job.scope,
        "status": job.status,
        "progress": job.progress,
        "stage": job.stage,
        "message": job.message,
        "snapshotId": job.snapshot_hash[:16],
        "traceId": job.trace_id,
        "citations": job.citations if isinstance(job.citations, list) else [],
        "verified": job.verified,
        "degraded": job.degraded,
        "degradedReasons": job.degraded_reasons if isinstance(job.degraded_reasons, list) else [],
        "promptVersion": job.prompt_version,
        "llmAudit": job.llm_audit if isinstance(job.llm_audit, dict) else {},
        "error": job.error_message if job.status == "failed" else "",
        "createdAt": job.created_at.isoformat(),
        "updatedAt": job.updated_at.isoformat(),
    }
    if job.status == "succeeded" and job.docx_path:
        payload["docxDownloadUrl"] = f"{base}/download/docx/"
    if job.status == "succeeded" and job.pdf_path:
        payload["pdfDownloadUrl"] = f"{base}/download/pdf/"
    return payload


def generate_review_handout(job_id: str) -> dict:
    job = ReviewHandoutJob.objects.select_related("user").get(id=job_id)
    if job.status == "cancelled":
        return {"success": False, "cancelled": True}
    if job.status == "succeeded":
        return {"success": True, "jobId": str(job.id)}
    artifact_dir = _artifact_directory(job)
    try:
        _save_job(job, status="running", progress=10, stage="retrieval", message="正在检索本地课程资料", started_at=job.started_at or timezone.now())
        evidence_by_point, retrieval_reasons = _collect_handout_evidence(job)
        _save_job(job, progress=35, stage="enhancing", message="正在生成个性化讲解与分层练习")
        result = enhance_review_handout(
            job.snapshot,
            job.ranked_weak_points[:8],
            evidence_by_point,
            cancel_check=lambda: _job_cancelled(job.id),
        )
        reasons = _safe_degraded_reasons([*result.degraded_reasons, *retrieval_reasons])
        degraded = result.degraded or bool(reasons)
        verified = result.verified and not degraded
        sections = result.content.get("sections") if isinstance(result.content, dict) else []
        sections = sections if isinstance(sections, list) else []
        plan = build_seven_day_plan(job.ranked_weak_points)
        enhancement = {"sections": sections, "sevenDayPlan": plan}
        _save_job(job, progress=65, stage="rendering", message="正在排版讲义与 7 天复习计划")
        artifact_dir.mkdir(parents=True, exist_ok=True)
        docx_path = artifact_dir / "review-handout.docx"
        pdf_path = artifact_dir / "review-handout.pdf"
        _atomic_write(docx_path, _build_handout_docx(job.snapshot, sections, result.citations, plan, reasons))
        _atomic_write(pdf_path, _build_handout_pdf(job.snapshot, sections, result.citations, plan, reasons))
        _save_job(
            job,
            status="succeeded",
            progress=100,
            stage="completed",
            message=(
                "资料检索暂不可用，已根据可信题目解析生成复习提纲"
                if "retrieval_error" in reasons
                else "部分内容已降级，讲义仍已根据可信错题记录生成"
                if degraded
                else "复习讲义已生成并通过引用核验"
            ),
            enhancement=enhancement,
            citations=result.citations,
            llm_audit=result.audit,
            degraded_reasons=reasons,
            prompt_version=str(result.audit.get("promptVersion") or ""),
            verified=verified,
            degraded=degraded,
            docx_path=str(docx_path),
            pdf_path=str(pdf_path),
            trace_id=job.trace_id or uuid.uuid4().hex,
            error_message="",
            finished_at=timezone.now(),
        )
        return {"success": True, "jobId": str(job.id)}
    except (ReviewHandoutCancelled, EnhancementCancelled):
        shutil.rmtree(artifact_dir, ignore_errors=True)
        return {"success": False, "cancelled": True}
    except Exception as exc:
        logger.exception("Review handout generation failed", extra={"review_handout_job_id": str(job.id)})
        shutil.rmtree(artifact_dir, ignore_errors=True)
        message = "复习讲义生成失败，请稍后重试。"
        try:
            _save_job(job, status="failed", progress=100, stage="failed", message=message, error_message=message, finished_at=timezone.now())
        except ReviewHandoutCancelled:
            return {"success": False, "cancelled": True}
        return {"success": False, "error": message}


def _collect_handout_evidence(job: ReviewHandoutJob) -> tuple[dict[str, list[dict]], list[str]]:
    evidence_by_point: dict[str, list[dict]] = {}
    reasons = []
    source_index = 1
    for weakness in job.ranked_weak_points[:8]:
        point = weakness["knowledgePoint"]
        evidence = []
        try:
            results = search_knowledge_base(query=f"{point} 核心概念 常见错误", top_k=3)
            for result in results[:3]:
                metadata = result.metadata if isinstance(result.metadata, dict) else {}
                citation = {
                    "sourceId": f"S{source_index}",
                    "title": str(metadata.get("title") or metadata.get("source_name") or "课程资料")[:255],
                    "location": _citation_location(metadata),
                    "snippet": str(result.text or "").strip()[:600],
                    "score": round(float(result.score or 0), 4),
                }
                if citation["snippet"]:
                    evidence.append(citation)
                    source_index += 1
        except Exception:
            reasons.append("retrieval_error")
        if not evidence:
            reasons.append("missing_evidence")
        evidence_by_point[point] = evidence
    return evidence_by_point, _safe_degraded_reasons(reasons)


def build_seven_day_plan(ranked_weak_points: list[dict]) -> list[dict]:
    ranked = sorted(
        [dict(item) for item in ranked_weak_points if isinstance(item, dict)],
        key=lambda item: (-float(item.get("score") or 0), str(item.get("knowledgePoint") or "")),
    )[:8]
    days = [{"day": day, "tasks": []} for day in range(1, 8)]
    if not ranked:
        return days

    for item in ranked:
        point = str(item.get("knowledgePoint") or "未分类")
        first_day = 1 if int(item.get("dueCount") or 0) > 0 else 2 if int(item.get("activeCount") or 0) > 0 else 3
        days[first_day - 1]["tasks"].append({
            "knowledgePoint": point,
            "type": "focused_practice",
            "action": "复习核心概念并完成基础、巩固、迁移练习",
        })
        for review_day in (first_day + 2, first_day + 4):
            if review_day <= 7:
                days[review_day - 1]["tasks"].append({
                    "knowledgePoint": point,
                    "type": "spaced_review",
                    "action": "遮住答案重做并核对错因是否消除",
                })
    return days


def _build_handout_docx(
    snapshot: dict,
    sections: list[dict],
    citations: list[dict],
    plan: list[dict],
    degraded_reasons: list[str],
) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    for name, size in (("Normal", 10.5), ("Heading 1", 16), ("Heading 2", 12)):
        style = document.styles[name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("个性化复习讲义")
    title_run.bold = True
    title_run.font.size = Pt(22)
    document.add_paragraph(f"学生：{snapshot['studentName']}　范围：{snapshot['scope']}　薄弱点：{'、'.join(str(item.get('knowledgePoint') or '') for item in sections)}")
    if degraded_reasons:
        document.add_paragraph("说明：部分 AI 增强或资料引用不可用，未核验内容已明确标注；原题事实与原始解析保持不变。")
    questions_by_point = _snapshot_questions_by_point(snapshot)
    citations_by_point = _citations_by_point(citations)
    for section_data in sections:
        point = str(section_data.get("knowledgePoint") or "未分类")
        document.add_heading(point, level=1)
        document.add_paragraph(f"证据状态：{_evidence_label(section_data.get('evidenceStatus'))}")
        _docx_section(document, "学习诊断", section_data.get("learningDiagnosis"))
        _docx_section(document, "核心概念", section_data.get("coreConcept"))
        _docx_section(document, "先修知识", _joined(section_data.get("prerequisites")))
        _docx_section(document, "常见错误", _joined(section_data.get("commonErrorPatterns")))
        _docx_section(document, "解题方法", _numbered(section_data.get("method")))
        _docx_worked_example(document, section_data.get("workedExample"))
        document.add_heading("关联错题", level=2)
        for row in questions_by_point.get(point, []):
            document.add_paragraph(str(row.get("stem") or ""))
        document.add_heading("分层练习", level=2)
        practices = section_data.get("tieredPractice") if isinstance(section_data.get("tieredPractice"), list) else []
        if practices:
            for practice in practices:
                document.add_paragraph(f"{practice.get('level', '')}练习：{practice.get('prompt', '')}")
        else:
            document.add_paragraph("AI 分层练习暂不可用，请先重练关联错题。")
        _docx_section(document, "复习行动", _joined(section_data.get("reviewActions")))
        document.add_heading("资料引用", level=2)
        point_citations = citations_by_point.get(point, [])
        if point_citations:
            for item in point_citations:
                document.add_paragraph(f"{item['title']} · {item['location']} · {item['snippet']}")
        else:
            document.add_paragraph("未检索到可核验的本地资料引用。")
    document.add_heading("7 天复习计划", level=1)
    for day in plan:
        tasks = day.get("tasks") if isinstance(day.get("tasks"), list) else []
        content = "；".join(f"{item.get('knowledgePoint')}：{item.get('action')}" for item in tasks) or "完成当日常规学习并整理疑问。"
        document.add_paragraph(f"第 {day.get('day')} 天：{content}")
    document.add_heading("答案与解析", level=1)
    for section_data in sections:
        point = str(section_data.get("knowledgePoint") or "未分类")
        document.add_heading(point, level=2)
        for row in questions_by_point.get(point, []):
            document.add_paragraph(f"原题：{row.get('stem', '')}\n正确答案：{row.get('correctAnswer', '')}\n原始解析：{row.get('explanation', '')}")
        example = section_data.get("workedExample") if isinstance(section_data.get("workedExample"), dict) else {}
        document.add_paragraph(f"例题：{example.get('prompt', '')}\n答案：{example.get('answer', '')}\n步骤：{_joined(example.get('steps'))}")
        for practice in section_data.get("tieredPractice") or []:
            document.add_paragraph(f"{practice.get('level', '')}练习：{practice.get('prompt', '')}\n答案：{practice.get('answer', '')}\n解析：{practice.get('explanation', '')}")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _docx_section(document, title: str, content) -> None:
    document.add_heading(title, level=2)
    document.add_paragraph(content or "暂无")


def _build_handout_pdf(
    snapshot: dict,
    sections: list[dict],
    citations: list[dict],
    plan: list[dict],
    degraded_reasons: list[str],
) -> bytes:
    pdf_font = register_pdf_cjk_font()
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=A4, rightMargin=18 * mm, leftMargin=18 * mm, topMargin=16 * mm, bottomMargin=18 * mm, title="智学引擎个性化复习讲义", author="智学引擎")
    styles = getSampleStyleSheet()
    title = ParagraphStyle("HTitle", parent=styles["Title"], fontName=pdf_font, fontSize=20, leading=26, alignment=TA_CENTER, textColor=colors.HexColor("#0F172A"))
    h1 = ParagraphStyle("HH1", parent=styles["Heading1"], fontName=pdf_font, fontSize=15, leading=21, textColor=colors.HexColor("#1F4E78"), spaceBefore=10)
    h2 = ParagraphStyle("HH2", parent=styles["Heading2"], fontName=pdf_font, fontSize=11, leading=16, textColor=colors.HexColor("#334155"), spaceBefore=6)
    body = ParagraphStyle("HBody", parent=styles["BodyText"], fontName=pdf_font, fontSize=9.5, leading=15, spaceAfter=4)
    story = [Paragraph("个性化复习讲义", title), Paragraph(_xml(f"学生：{snapshot['studentName']}　薄弱点：{'、'.join(str(item.get('knowledgePoint') or '') for item in sections)}"), body)]
    if degraded_reasons:
        story.append(Paragraph("部分 AI 增强或资料引用不可用，未核验内容已明确标注；原题事实与原始解析保持不变。", body))
    questions_by_point = _snapshot_questions_by_point(snapshot)
    citations_by_point = _citations_by_point(citations)
    for section_data in sections:
        point = str(section_data.get("knowledgePoint") or "未分类")
        story.append(Paragraph(_xml(point), h1))
        story.append(Paragraph(_xml(f"证据状态：{_evidence_label(section_data.get('evidenceStatus'))}"), body))
        for label, content in (
            ("学习诊断", section_data.get("learningDiagnosis")),
            ("核心概念", section_data.get("coreConcept")),
            ("先修知识", _joined(section_data.get("prerequisites"))),
            ("常见错误", _joined(section_data.get("commonErrorPatterns"))),
            ("解题方法", _numbered(section_data.get("method"))),
        ):
            story.extend([Paragraph(label, h2), Paragraph(_xml(content), body)])
        example = section_data.get("workedExample") if isinstance(section_data.get("workedExample"), dict) else {}
        story.extend([Paragraph("例题精讲", h2), Paragraph(_xml(f"{example.get('prompt', '')}；步骤：{_joined(example.get('steps'))}"), body)])
        story.append(Paragraph("关联错题", h2))
        for row in questions_by_point.get(point, []):
            story.append(Paragraph(_xml(row.get("stem")), body))
        story.append(Paragraph("分层练习", h2))
        practices = section_data.get("tieredPractice") if isinstance(section_data.get("tieredPractice"), list) else []
        if practices:
            for practice in practices:
                story.append(Paragraph(_xml(f"{practice.get('level', '')}练习：{practice.get('prompt', '')}"), body))
        else:
            story.append(Paragraph("AI 分层练习暂不可用，请先重练关联错题。", body))
        story.extend([Paragraph("复习行动", h2), Paragraph(_xml(_joined(section_data.get("reviewActions"))), body)])
        story.append(Paragraph("资料引用", h2))
        point_citations = citations_by_point.get(point, [])
        if point_citations:
            for item in point_citations:
                story.append(Paragraph(_xml(f"{item['title']} · {item['location']} · {item['snippet']}"), body))
        else:
            story.append(Paragraph("未检索到可核验的本地资料引用。", body))
    story.append(Paragraph("7 天复习计划", h1))
    for day in plan:
        tasks = day.get("tasks") if isinstance(day.get("tasks"), list) else []
        content = "；".join(f"{item.get('knowledgePoint')}：{item.get('action')}" for item in tasks) or "完成当日常规学习并整理疑问。"
        story.append(Paragraph(_xml(f"第 {day.get('day')} 天：{content}"), body))
    story.append(Paragraph("答案与解析", h1))
    for section_data in sections:
        point = str(section_data.get("knowledgePoint") or "未分类")
        story.append(Paragraph(_xml(point), h2))
        for row in questions_by_point.get(point, []):
            story.append(Paragraph(_xml(f"原题：{row.get('stem', '')}　正确答案：{row.get('correctAnswer', '')}　原始解析：{row.get('explanation', '')}"), body))
        example = section_data.get("workedExample") if isinstance(section_data.get("workedExample"), dict) else {}
        story.append(Paragraph(_xml(f"例题：{example.get('prompt', '')}　答案：{example.get('answer', '')}　步骤：{_joined(example.get('steps'))}"), body))
        for practice in section_data.get("tieredPractice") or []:
            story.append(Paragraph(_xml(f"{practice.get('level', '')}练习：{practice.get('prompt', '')}　答案：{practice.get('answer', '')}　解析：{practice.get('explanation', '')}"), body))
    doc.build(story)
    return output.getvalue()


def _save_job(job: ReviewHandoutJob, **changes) -> None:
    with transaction.atomic():
        current = ReviewHandoutJob.objects.select_for_update().get(pk=job.pk)
        if current.status == "cancelled" and changes.get("status") != "cancelled":
            raise ReviewHandoutCancelled
        for field, value in changes.items():
            setattr(current, field, value)
            setattr(job, field, value)
        current.save(update_fields=[*changes.keys(), "updated_at"])


def _job_cancelled(job_id) -> bool:
    return ReviewHandoutJob.objects.filter(pk=job_id, status="cancelled").exists()


def _safe_degraded_reasons(reasons) -> list[str]:
    safe = []
    for reason in reasons if isinstance(reasons, (list, tuple)) else []:
        normalized = str(reason or "").strip().lower()
        if normalized in SAFE_DEGRADED_REASONS and normalized not in safe:
            safe.append(normalized)
    return safe


def _snapshot_questions_by_point(snapshot: dict) -> dict[str, list[dict]]:
    grouped: dict[str, list[dict]] = defaultdict(list)
    for row in snapshot.get("questions") or []:
        if isinstance(row, dict):
            grouped[str(row.get("knowledgePoint") or "未分类")].append(row)
    return grouped


def _citations_by_point(citations: list[dict]) -> dict[str, list[dict]]:
    grouped: dict[str, list[dict]] = defaultdict(list)
    for citation in citations if isinstance(citations, list) else []:
        if isinstance(citation, dict):
            grouped[str(citation.get("knowledgePoint") or "未分类")].append(citation)
    return grouped


def _joined(value) -> str:
    if isinstance(value, list):
        return "；".join(str(item) for item in value if str(item).strip())
    return str(value or "")


def _numbered(value) -> str:
    if not isinstance(value, list):
        return str(value or "")
    return "\n".join(f"{index}. {item}" for index, item in enumerate(value, start=1))


def _evidence_label(value) -> str:
    return {"verified": "已核验", "partial": "部分核验", "unverified": "未核验"}.get(str(value), "未核验")


def _docx_worked_example(document, value) -> None:
    example = value if isinstance(value, dict) else {}
    document.add_heading("例题精讲", level=2)
    document.add_paragraph(str(example.get("prompt") or "暂无"))
    document.add_paragraph(f"步骤：{_numbered(example.get('steps')) or '暂无'}")


def _artifact_directory(job: ReviewHandoutJob) -> Path:
    root = Path(getattr(settings, "A3_REVIEW_HANDOUT_ROOT", settings.BASE_DIR / "tmp" / "review-handouts"))
    return root / str(job.user_id) / str(job.id)


def resolve_handout_artifact(job: ReviewHandoutJob, export_format: str) -> Path:
    normalized = str(export_format or "").lower()
    if normalized not in {"docx", "pdf"}:
        raise WrongExportRequestError("讲义下载格式必须是 docx 或 pdf")
    stored = job.docx_path if normalized == "docx" else job.pdf_path
    if job.status != "succeeded" or not stored:
        raise WrongExportRequestError("讲义文件尚未生成")
    root = Path(getattr(settings, "A3_REVIEW_HANDOUT_ROOT", settings.BASE_DIR / "tmp" / "review-handouts")).resolve()
    path = Path(stored).resolve()
    try:
        path.relative_to(root)
    except ValueError as exc:
        raise WrongExportRequestError("讲义文件路径无效") from exc
    if not path.is_file() or path.suffix.lower() != f".{normalized}":
        raise WrongExportRequestError("讲义文件不存在，请重新生成")
    return path


def _atomic_write(path: Path, content: bytes) -> None:
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_bytes(content)
    os.replace(temporary, path)


def _citation_location(metadata: dict) -> str:
    page = metadata.get("page_number") or metadata.get("pageNumber")
    block = metadata.get("block_id") or metadata.get("blockId")
    if page:
        return f"第 {page} 页"
    if block:
        return f"位置 {str(block)[:80]}"
    return "资料正文"


def _xml(value) -> str:
    return str(value or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _is_due(value, now) -> bool:
    if not value:
        return True
    try:
        return timezone.datetime.fromisoformat(str(value)) <= now
    except ValueError:
        return True


def _recency_score(value, now) -> float:
    if not value:
        return 0
    try:
        occurred = timezone.datetime.fromisoformat(str(value))
    except ValueError:
        return 0
    days = max(0.0, (now - occurred).total_seconds() / 86400)
    return max(0.0, 10.0 - min(days, 10.0))
