from __future__ import annotations

import json
import os
from pathlib import Path

from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from backend.apps.assessments.review_handouts import (
    create_review_handout_job,
    generate_review_handout,
    resolve_handout_artifact,
)
from backend.apps.assessments.wrong_export_jobs import (
    create_wrong_export_job,
    generate_wrong_export_job,
    resolve_wrong_export_artifact,
)


class Command(BaseCommand):
    help = "安全验证 AI 错题导出和复习讲义，不输出 Prompt、资料内容或服务器路径。"

    def add_arguments(self, parser):
        parser.add_argument("--user-id", required=True)
        parser.add_argument(
            "--real-model",
            action="store_true",
            help="显式允许调用已配置的 OpenAI-compatible 模型服务。",
        )

    def handle(self, *args, **options):
        user = self._get_user(options["user_id"])
        real_model = bool(options["real_model"])
        previous_enabled = os.environ.get("A3_ASSESSMENT_EXPORT_LLM_ENABLED")
        os.environ["A3_ASSESSMENT_EXPORT_LLM_ENABLED"] = "true" if real_model else "false"
        try:
            summaries = self._run_smoke(user)
        finally:
            if previous_enabled is None:
                os.environ.pop("A3_ASSESSMENT_EXPORT_LLM_ENABLED", None)
            else:
                os.environ["A3_ASSESSMENT_EXPORT_LLM_ENABLED"] = previous_enabled

        if real_model and not any(item["model"] not in {"", "rules"} for item in summaries):
            raise CommandError("真实模型未成功参与本次 smoke，请检查服务配置和审计结果。")
        for summary in summaries:
            self.stdout.write(json.dumps(summary, ensure_ascii=False, separators=(",", ":")))

    def _get_user(self, user_id):
        try:
            return get_user_model().objects.get(pk=user_id)
        except (ValueError, TypeError, get_user_model().DoesNotExist) as exc:
            raise CommandError("指定用户不存在。") from exc

    def _run_smoke(self, user) -> list[dict]:
        summaries = []
        for export_format in ("docx", "xlsx", "pdf"):
            job, _ = create_wrong_export_job(user, "all", export_format)
            generate_wrong_export_job(str(job.id))
            job.refresh_from_db()
            if job.status != "succeeded":
                raise CommandError("错题导出 smoke 失败，请检查任务审计状态。")
            path = resolve_wrong_export_artifact(job)
            _parse_artifact(path, export_format)
            summaries.append(_summary(job, path))

        handout, _ = create_review_handout_job(user, "all")
        generate_review_handout(str(handout.id))
        handout.refresh_from_db()
        if handout.status != "succeeded":
            raise CommandError("复习讲义 smoke 失败，请检查任务审计状态。")
        for export_format in ("docx", "pdf"):
            path = resolve_handout_artifact(handout, export_format)
            _parse_artifact(path, export_format)
            summaries.append(_summary(handout, path))
        return summaries


def _parse_artifact(path: Path, export_format: str) -> None:
    try:
        if export_format == "docx":
            document = Document(path)
            if not document.paragraphs:
                raise ValueError("empty document")
        elif export_format == "xlsx":
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                if not workbook.sheetnames:
                    raise ValueError("empty workbook")
            finally:
                workbook.close()
        elif export_format == "pdf":
            if not PdfReader(path).pages:
                raise ValueError("empty pdf")
        else:
            raise ValueError("unsupported artifact")
    except Exception as exc:
        raise CommandError("生成文件无法解析，请检查渲染器。") from exc


def _summary(job, path: Path) -> dict:
    audit = job.llm_audit if isinstance(job.llm_audit, dict) else {}
    citations = job.citations if isinstance(job.citations, list) else []
    return {
        "jobId": str(job.id),
        "status": str(job.status),
        "model": str(audit.get("model") or "rules"),
        "verified": bool(job.verified),
        "degraded": bool(job.degraded),
        "citationCount": len(citations),
        "fileSize": path.stat().st_size,
    }
