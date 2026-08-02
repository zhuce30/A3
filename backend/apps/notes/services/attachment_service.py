from __future__ import annotations

import hashlib
import html.parser
import io
import ipaddress
import os
import re
import socket
import stat
import tempfile
import unicodedata
import urllib.parse
import uuid
import zipfile
from pathlib import Path, PurePosixPath
from typing import Any

import requests
from requests.auth import AuthBase
from django.conf import settings
from django.contrib.auth import get_user_model
from django.db import transaction
from django.db.models import Sum
from django.utils import timezone

from backend.apps.notes.models import Note, NoteAttachment


MAX_ATTACHMENTS_PER_NOTE = 30
MAX_LINK_BYTES = 2 * 1024 * 1024
MAX_LINK_REDIRECTS = 3
LINK_TIMEOUT_SECONDS = 5
MAX_ARCHIVE_ENTRIES = 10_000
MAX_ARCHIVE_RATIO = 200

OFFICE_MIMES = {
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
ALLOWED_TYPES = {
    ".pdf": {"application/pdf"},
    ".docx": {OFFICE_MIMES[".docx"]},
    ".pptx": {OFFICE_MIMES[".pptx"]},
    ".xlsx": {OFFICE_MIMES[".xlsx"]},
    ".txt": {"text/plain"},
    ".md": {"text/markdown", "text/plain", "text/x-markdown"},
    ".markdown": {"text/markdown", "text/plain", "text/x-markdown"},
    ".png": {"image/png"},
    ".jpg": {"image/jpeg"},
    ".jpeg": {"image/jpeg"},
    ".webp": {"image/webp"},
}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}
OFFICE_ROOTS = {".docx": "word/document.xml", ".pptx": "ppt/presentation.xml", ".xlsx": "xl/workbook.xml"}
ACTIVE_OFFICE_PARTS = ("vbaproject", "activex/", "embeddings/", "oleobject")
EXECUTABLE_SIGNATURES = (b"MZ", b"\x7fELF", b"\xfe\xed\xfa", b"\xcf\xfa\xed\xfe")


class _NoAuth(AuthBase):
    def __call__(self, request):
        request.headers.pop("Authorization", None)
        request.headers.pop("Cookie", None)
        return request


NO_AUTH = _NoAuth()


class AttachmentServiceError(Exception):
    def __init__(self, message: str, *, code: str = "invalid_attachment", status: int = 400) -> None:
        super().__init__(message)
        self.code = code
        self.status = status


class AttachmentNotFound(AttachmentServiceError):
    def __init__(self, message: str = "附件不存在") -> None:
        super().__init__(message, code="not_found", status=404)


def serialize_attachment(attachment: NoteAttachment) -> dict[str, Any]:
    return {
        "id": str(attachment.id),
        "noteId": str(attachment.note_id),
        "kind": attachment.kind,
        "displayName": attachment.display_name,
        "mimeType": attachment.mime_type,
        "sizeBytes": attachment.size_bytes,
        "parseStatus": attachment.parse_status,
        "errorCode": attachment.error_code,
        "errorMessage": attachment.error_message,
        "createdAt": attachment.created_at.isoformat() if attachment.created_at else None,
        "updatedAt": attachment.updated_at.isoformat() if attachment.updated_at else None,
    }


def get_owned_note(user, note_id) -> Note:
    try:
        return Note.objects.get(user=user, pk=note_id, deleted_at__isnull=True)
    except (Note.DoesNotExist, ValueError) as exc:
        raise AttachmentNotFound("笔记不存在") from exc


def get_owned_attachment(user, attachment_id, *, lock: bool = False) -> NoteAttachment:
    queryset = NoteAttachment.objects.select_for_update() if lock else NoteAttachment.objects
    try:
        return queryset.select_related("note").get(
            user=user,
            pk=attachment_id,
            deleted_at__isnull=True,
            note__deleted_at__isnull=True,
        )
    except (NoteAttachment.DoesNotExist, ValueError) as exc:
        raise AttachmentNotFound() from exc


def _safe_filename(value: Any) -> tuple[str, str]:
    if not isinstance(value, str):
        raise AttachmentServiceError("附件名称无效", code="invalid_attachment_name")
    normalized = unicodedata.normalize("NFKC", value).strip()
    if (
        not normalized
        or len(normalized) > 255
        or normalized in {".", ".."}
        or "/" in normalized
        or "\\" in normalized
        or any(ord(char) < 32 or ord(char) == 127 for char in normalized)
    ):
        raise AttachmentServiceError("附件名称无效", code="invalid_attachment_name")
    suffix = Path(normalized).suffix.lower()
    if suffix not in ALLOWED_TYPES:
        raise AttachmentServiceError("不支持该附件类型", code="unsupported_attachment_type")
    stem = Path(normalized).stem.strip().strip(".") or "attachment"
    stem = re.sub(r"[^\w .()-]", "_", stem, flags=re.UNICODE).strip(" .") or "attachment"
    safe_name = f"{stem[:220]}{suffix}"
    return safe_name, suffix


def validate_attachment_filename(value: Any) -> str:
    """Validate an unsanitized client filename at service boundaries."""
    return _safe_filename(value)[0]


def _read_upload(upload) -> bytes:
    maximum = int(settings.A3_NOTES_ATTACHMENT_MAX_BYTES)
    declared_size = getattr(upload, "size", None)
    if declared_size is not None and declared_size > maximum:
        raise AttachmentServiceError("附件超过大小限制", code="attachment_too_large", status=413)
    chunks: list[bytes] = []
    total = 0
    for chunk in upload.chunks():
        total += len(chunk)
        if total > maximum:
            raise AttachmentServiceError("附件超过大小限制", code="attachment_too_large", status=413)
        chunks.append(chunk)
    return b"".join(chunks)


def _validate_signature(suffix: str, content: bytes) -> None:
    if any(content.startswith(signature) for signature in EXECUTABLE_SIGNATURES):
        raise AttachmentServiceError("附件包含不安全内容", code="unsafe_attachment_content")
    signatures = {
        ".pdf": content.startswith(b"%PDF-"),
        ".png": content.startswith(b"\x89PNG\r\n\x1a\n"),
        ".jpg": content.startswith(b"\xff\xd8\xff"),
        ".jpeg": content.startswith(b"\xff\xd8\xff"),
        ".webp": len(content) >= 12 and content[:4] == b"RIFF" and content[8:12] == b"WEBP",
    }
    if suffix in signatures and not signatures[suffix]:
        raise AttachmentServiceError("附件内容与扩展名不一致", code="attachment_signature_mismatch")
    if suffix in OFFICE_MIMES:
        _validate_office_container(suffix, content)
    elif suffix in {".txt", ".md", ".markdown"}:
        if b"\x00" in content[:8192]:
            raise AttachmentServiceError("文本附件包含无效二进制内容", code="attachment_signature_mismatch")
        try:
            content[:8192].decode("utf-8")
        except UnicodeDecodeError as exc:
            raise AttachmentServiceError("文本附件必须使用 UTF-8 编码", code="attachment_signature_mismatch") from exc


def _validate_office_container(suffix: str, content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
            names = {info.filename.replace("\\", "/").lstrip("/") for info in infos}
            if len(infos) > MAX_ARCHIVE_ENTRIES or OFFICE_ROOTS[suffix] not in names or "[Content_Types].xml" not in names:
                raise ValueError("invalid office structure")
            total_uncompressed = 0
            archive_limit = max(int(settings.A3_NOTES_ATTACHMENT_MAX_BYTES) * 10, 2 * 1024 * 1024)
            for info in infos:
                path = PurePosixPath(info.filename.replace("\\", "/"))
                lowered = str(path).lower().lstrip("/")
                if path.is_absolute() or ".." in path.parts or info.flag_bits & 0x1:
                    raise ValueError("unsafe archive entry")
                if stat.S_ISLNK(info.external_attr >> 16) or any(marker in lowered for marker in ACTIVE_OFFICE_PARTS):
                    raise ValueError("active office content")
                total_uncompressed += info.file_size
                if total_uncompressed > archive_limit:
                    raise ValueError("decompression limit")
                if info.compress_size == 0 and info.file_size > 0:
                    raise ValueError("invalid compression ratio")
                if info.compress_size and info.file_size / info.compress_size > MAX_ARCHIVE_RATIO:
                    raise ValueError("compression ratio")
            content_types = archive.read("[Content_Types].xml").lower()
            if OFFICE_MIMES[suffix].encode("ascii") not in content_types:
                raise ValueError("office content type mismatch")
            if b"macroenabled" in content_types or b"vbaproject" in content_types:
                raise ValueError("macro content type")
            for info in infos:
                if info.filename.lower().endswith(".rels") and info.file_size <= 1024 * 1024:
                    relationships = archive.read(info).lower()
                    if b'targetmode="external"' in relationships and b"/hyperlink" not in relationships:
                        raise ValueError("external active relationship")
    except (zipfile.BadZipFile, KeyError, ValueError, RuntimeError) as exc:
        raise AttachmentServiceError("Office 附件结构不安全或已损坏", code="unsafe_office_document") from exc


def _validate_declared_mime(suffix: str, declared_mime: str) -> str:
    normalized = str(declared_mime or "").split(";", 1)[0].strip().lower()
    if normalized not in ALLOWED_TYPES[suffix]:
        raise AttachmentServiceError("附件 MIME 类型与扩展名不一致", code="attachment_mime_mismatch")
    return OFFICE_MIMES.get(suffix, "text/markdown" if suffix in {".md", ".markdown"} else normalized)


def storage_root() -> Path:
    root = Path(settings.A3_NOTES_STORAGE_ROOT).expanduser().resolve()
    root.mkdir(parents=True, exist_ok=True)
    return root


def _user_storage_uuid(user) -> uuid.UUID:
    return uuid.uuid5(uuid.NAMESPACE_URL, f"urn:a3:personal-notes:user:{user.pk}")


def resolve_relative_path(relative_path: str, *, require_file: bool = False) -> Path:
    pure = PurePosixPath(str(relative_path))
    if pure.is_absolute() or not pure.parts or any(part in {"", ".", ".."} for part in pure.parts):
        raise AttachmentNotFound()
    root = storage_root()
    candidate = root.joinpath(*pure.parts).resolve()
    try:
        candidate.relative_to(root)
    except ValueError as exc:
        raise AttachmentNotFound() from exc
    if require_file and (not candidate.is_file() or candidate.is_symlink()):
        raise AttachmentNotFound()
    return candidate


def _check_capacity(user, note: Note, incoming_bytes: int) -> None:
    get_user_model().objects.select_for_update().get(pk=user.pk)
    active = NoteAttachment.objects.filter(user=user, deleted_at__isnull=True)
    if active.filter(note=note).count() >= MAX_ATTACHMENTS_PER_NOTE:
        raise AttachmentServiceError("每条笔记最多可保留 30 个附件", code="attachment_count_exceeded", status=409)
    used = active.aggregate(total=Sum("size_bytes"))["total"] or 0
    if used + incoming_bytes > int(settings.A3_NOTES_USER_QUOTA_BYTES):
        raise AttachmentServiceError("个人附件空间不足", code="attachment_quota_exceeded", status=409)


def create_file_attachment(user, note_id, upload) -> NoteAttachment:
    note = get_owned_note(user, note_id)
    safe_name, suffix = _safe_filename(getattr(upload, "name", ""))
    content = _read_upload(upload)
    mime_type = _validate_declared_mime(suffix, getattr(upload, "content_type", ""))
    _validate_signature(suffix, content)
    attachment_id = uuid.uuid4()
    relative_path = PurePosixPath(
        str(_user_storage_uuid(user)), str(note.pk), str(attachment_id), safe_name
    ).as_posix()
    final_path = resolve_relative_path(relative_path)
    final_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(dir=final_path.parent, prefix=".upload-", delete=False) as temporary:
            temporary.write(content)
            temporary.flush()
            os.fsync(temporary.fileno())
            temporary_path = Path(temporary.name)
        try:
            with transaction.atomic():
                _check_capacity(user, note, len(content))
                os.replace(temporary_path, final_path)
                temporary_path = None
                return NoteAttachment.objects.create(
                    id=attachment_id,
                    user=user,
                    note=note,
                    kind=NoteAttachment.Kind.IMAGE if suffix in IMAGE_EXTENSIONS else NoteAttachment.Kind.FILE,
                    display_name=safe_name,
                    relative_path=relative_path,
                    mime_type=mime_type,
                    size_bytes=len(content),
                    sha256=hashlib.sha256(content).hexdigest(),
                )
        except Exception:
            final_path.unlink(missing_ok=True)
            raise
    finally:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)


def _validate_public_url(url: str) -> tuple[str, set[str]]:
    if not isinstance(url, str) or len(url) > 2048:
        raise AttachmentServiceError("链接格式无效", code="unsafe_note_link")
    parsed = urllib.parse.urlsplit(url)
    if parsed.scheme.lower() not in {"http", "https"} or not parsed.hostname or parsed.username or parsed.password:
        raise AttachmentServiceError("仅支持安全的 HTTP(S) 链接", code="unsafe_note_link")
    if parsed.fragment:
        parsed = parsed._replace(fragment="")
        url = urllib.parse.urlunsplit(parsed)
    try:
        port = parsed.port or (443 if parsed.scheme.lower() == "https" else 80)
        records = socket.getaddrinfo(parsed.hostname, port, type=socket.SOCK_STREAM)
    except (OSError, ValueError, UnicodeError) as exc:
        raise AttachmentServiceError("链接地址无法安全解析", code="unsafe_note_link") from exc
    addresses = {str(record[4][0]).split("%", 1)[0] for record in records}
    if not addresses:
        raise AttachmentServiceError("链接地址无法安全解析", code="unsafe_note_link")
    for address in addresses:
        try:
            ip = ipaddress.ip_address(address)
        except ValueError as exc:
            raise AttachmentServiceError("链接地址无法安全解析", code="unsafe_note_link") from exc
        if not ip.is_global:
            raise AttachmentServiceError("链接目标不是公网地址", code="unsafe_note_link")
    return url, addresses


def _response_peer_ip(response) -> str | None:
    candidates = (
        getattr(getattr(getattr(response, "raw", None), "_connection", None), "sock", None),
        getattr(getattr(getattr(response, "raw", None), "connection", None), "sock", None),
    )
    for sock in candidates:
        if sock is not None:
            try:
                return str(sock.getpeername()[0]).split("%", 1)[0]
            except (OSError, AttributeError, TypeError):
                continue
    return None


def fetch_safe_link(url: str) -> tuple[str, bytes, str]:
    current = url
    for redirect_count in range(MAX_LINK_REDIRECTS + 1):
        current, resolved_addresses = _validate_public_url(current)
        try:
            response = requests.get(
                current,
                allow_redirects=False,
                stream=True,
                timeout=LINK_TIMEOUT_SECONDS,
                headers={"User-Agent": "A3-Notes-LinkValidator/1.0", "Accept": "text/plain,text/html,application/xhtml+xml"},
                auth=NO_AUTH,
                cookies={},
            )
        except requests.RequestException as exc:
            raise AttachmentServiceError("链接暂时无法访问", code="note_link_unavailable") from exc
        peer_ip = _response_peer_ip(response)
        if peer_ip is not None and peer_ip not in resolved_addresses:
            response.close()
            raise AttachmentServiceError("链接解析结果在请求期间发生变化", code="unsafe_note_link")
        if response.status_code in {301, 302, 303, 307, 308}:
            location = response.headers.get("Location", "")
            response.close()
            if redirect_count >= MAX_LINK_REDIRECTS or not location:
                raise AttachmentServiceError("链接重定向次数过多", code="unsafe_note_link")
            current = urllib.parse.urljoin(current, location)
            continue
        if response.status_code < 200 or response.status_code >= 300:
            response.close()
            raise AttachmentServiceError("链接暂时无法访问", code="note_link_unavailable")
        try:
            declared_length = int(response.headers.get("Content-Length", "0") or 0)
        except ValueError:
            declared_length = 0
        if declared_length > MAX_LINK_BYTES:
            response.close()
            raise AttachmentServiceError("链接内容超过大小限制", code="note_link_too_large")
        chunks: list[bytes] = []
        total = 0
        try:
            try:
                for chunk in response.iter_content(chunk_size=64 * 1024):
                    if not chunk:
                        continue
                    total += len(chunk)
                    if total > MAX_LINK_BYTES:
                        raise AttachmentServiceError("链接内容超过大小限制", code="note_link_too_large")
                    chunks.append(chunk)
            except requests.RequestException as exc:
                raise AttachmentServiceError("链接暂时无法访问", code="note_link_unavailable") from exc
        finally:
            response.close()
        mime_type = str(response.headers.get("Content-Type", "")).split(";", 1)[0].strip().lower()
        return current, b"".join(chunks), mime_type
    raise AttachmentServiceError("链接重定向次数过多", code="unsafe_note_link")


@transaction.atomic
def create_link_attachment(user, note_id, data: dict[str, Any]) -> NoteAttachment:
    note = get_owned_note(user, note_id)
    url = data.get("url")
    final_url, body, mime_type = fetch_safe_link(url)
    display_name = data.get("displayName") or urllib.parse.urlsplit(final_url).hostname or "网页链接"
    if not isinstance(display_name, str) or not display_name.strip() or len(display_name.strip()) > 255:
        raise AttachmentServiceError("链接名称无效", code="invalid_attachment_name")
    _check_capacity(user, note, 0)
    return NoteAttachment.objects.create(
        user=user,
        note=note,
        kind=NoteAttachment.Kind.LINK,
        display_name=unicodedata.normalize("NFKC", display_name).strip(),
        source_url=final_url,
        mime_type=mime_type,
        size_bytes=0,
        sha256=hashlib.sha256(body).hexdigest(),
    )


def soft_delete_attachment(user, attachment_id) -> NoteAttachment:
    with transaction.atomic():
        attachment = get_owned_attachment(user, attachment_id, lock=True)
        attachment.deleted_at = timezone.now()
        attachment.save(update_fields={"deleted_at", "updated_at"})
        return attachment


def reset_attachment_for_reparse(user, attachment_id) -> NoteAttachment:
    with transaction.atomic():
        attachment = get_owned_attachment(user, attachment_id, lock=True)
        attachment.parse_status = NoteAttachment.ParseStatus.PENDING
        attachment.extracted_text = ""
        attachment.error_code = ""
        attachment.error_message = ""
        attachment.preview_metadata = {}
        attachment.save(update_fields={
            "parse_status", "extracted_text", "error_code", "error_message", "preview_metadata", "updated_at"
        })
        return attachment


class _TextExtractor(html.parser.HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.hidden_depth = 0

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag in {"script", "style", "noscript"}:
            self.hidden_depth += 1

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style", "noscript"} and self.hidden_depth:
            self.hidden_depth -= 1

    def handle_data(self, data: str) -> None:
        if not self.hidden_depth and data.strip():
            self.parts.append(data.strip())


def _decode_web_body(body: bytes, mime_type: str) -> str:
    text = body.decode("utf-8", errors="replace")
    if mime_type in {"text/html", "application/xhtml+xml"}:
        parser = _TextExtractor()
        parser.feed(text)
        return "\n".join(parser.parts)
    return text


def extract_attachment_text(attachment: NoteAttachment) -> tuple[str, dict[str, Any]]:
    maximum = int(settings.A3_NOTES_EXTRACTED_TEXT_MAX_CHARS)
    if attachment.kind == NoteAttachment.Kind.LINK:
        _, body, mime_type = fetch_safe_link(attachment.source_url)
        text = _decode_web_body(body, mime_type)
        return text[:maximum], {"truncated": len(text) > maximum, "parser": "safe_link"}
    path = resolve_relative_path(attachment.relative_path, require_file=True)
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".markdown"}:
        text = path.read_text(encoding="utf-8")
        parser_name = "utf8_text"
    elif suffix == ".docx":
        from backend.domains.knowledge_base.docx_parser import parse_docx_blocks

        text = "\n\n".join(block.markdown for block in parse_docx_blocks(path))
        parser_name = "knowledge_base_docx"
    elif suffix in {".pdf", ".pptx"}:
        from backend.domains.knowledge_base.document_loader import load_documents

        documents, skipped = load_documents([path])
        if skipped and not documents:
            raise ValueError("document parser failed")
        text = "\n\n".join(document.get_content() for document in documents)
        parser_name = "knowledge_base_document_loader"
    elif suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        rows: list[str] = []
        try:
            for sheet in workbook.worksheets:
                rows.append(sheet.title)
                for row in sheet.iter_rows():
                    line = "\t".join(str(cell.value) for cell in row if cell.value is not None).strip()
                    if line:
                        rows.append(line)
                    if sum(len(item) + 1 for item in rows) > maximum:
                        break
                if sum(len(item) + 1 for item in rows) > maximum:
                    break
        finally:
            workbook.close()
        text = "\n".join(rows)
        parser_name = "openpyxl"
    elif suffix in IMAGE_EXTENSIONS:
        from backend.domains.knowledge_base.multimodal_parser import parse_resource

        result = parse_resource(path)
        text = "\n\n".join(
            block.text for page in result.pages for block in page.blocks if block.text.strip()
        )
        parser_name = result.parser
    else:
        raise ValueError("unsupported attachment parser")
    return text[:maximum], {"truncated": len(text) > maximum, "parser": parser_name}


def parse_and_persist_attachment(attachment: NoteAttachment) -> dict[str, Any]:
    text, metadata = extract_attachment_text(attachment)
    updated = NoteAttachment.objects.filter(
        pk=attachment.pk,
        deleted_at__isnull=True,
        parse_status=NoteAttachment.ParseStatus.PROCESSING,
    ).update(
        parse_status=NoteAttachment.ParseStatus.SUCCEEDED,
        extracted_text=text,
        preview_metadata=metadata,
        error_code="",
        error_message="",
        updated_at=timezone.now(),
    )
    return {"status": "succeeded" if updated else "skipped", "attachmentId": str(attachment.pk)}
