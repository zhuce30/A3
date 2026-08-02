from __future__ import annotations

import socket
import threading
import time
import uuid
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from backend.apps.notes.models import NoteAttachment
from backend.apps.notes.tests.test_attachments import AttachmentTestCase


PUBLIC_IPV4 = "93.184.216.34"


class _SocketPeerOverride:
    def __init__(self, wrapped: socket.socket, peer: str) -> None:
        self._wrapped = wrapped
        self._peer = peer

    def getpeername(self):
        return self._peer, self._wrapped.getpeername()[1]

    def __getattr__(self, name):
        return getattr(self._wrapped, name)


class _CountingHandler(BaseHTTPRequestHandler):
    requests_seen = 0
    drip = False

    def do_GET(self):
        type(self).requests_seen += 1
        self.send_response(200)
        self.send_header("Content-Type", "text/plain")
        self.end_headers()
        if not type(self).drip:
            self.wfile.write(b"ok")
            return
        for _ in range(8):
            try:
                self.wfile.write(b"x")
                self.wfile.flush()
            except (BrokenPipeError, ConnectionResetError):
                break
            time.sleep(0.9)

    def log_message(self, _format, *args):
        return


class _LocalServer:
    def __init__(self, *, drip: bool = False) -> None:
        handler = type("ProbeHandler", (_CountingHandler,), {"requests_seen": 0, "drip": drip})
        self.handler = handler
        self.server = ThreadingHTTPServer(("127.0.0.1", 0), handler)
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)

    def __enter__(self):
        self.thread.start()
        return self

    def __exit__(self, exc_type, exc, traceback):
        self.server.shutdown()
        self.server.server_close()
        self.thread.join(timeout=2)

    @property
    def port(self) -> int:
        return self.server.server_address[1]


def _dns_record(address: str, port: int):
    family = socket.AF_INET6 if ":" in address else socket.AF_INET
    return [(family, socket.SOCK_STREAM, socket.IPPROTO_TCP, "", (address, port))]


class PinnedTransportProbeTests(AttachmentTestCase):
    def test_dns_rebinding_private_listener_receives_no_request(self) -> None:
        from backend.apps.notes.services.link_fetcher import SafeLinkError, fetch_safe_link

        with _LocalServer() as probe:
            resolutions = iter(
                [_dns_record(PUBLIC_IPV4, probe.port), _dns_record("127.0.0.1", probe.port)]
            )

            def refuse_public_target(address, timeout=None, source_address=None):
                self.assertEqual(address, (PUBLIC_IPV4, probe.port))
                raise socket.timeout("public target unavailable")

            with patch(
                "backend.apps.notes.services.link_fetcher.socket.getaddrinfo",
                side_effect=lambda *args, **kwargs: next(resolutions),
            ), patch(
                "backend.apps.notes.services.link_fetcher.socket.create_connection",
                side_effect=refuse_public_target,
            ):
                with self.assertRaises(SafeLinkError):
                    fetch_safe_link(f"http://rebind.test:{probe.port}/resource")

            self.assertEqual(probe.handler.requests_seen, 0)

    def test_slow_drip_obeys_single_five_second_deadline(self) -> None:
        from backend.apps.notes.services.link_fetcher import SafeLinkError, fetch_safe_link

        with _LocalServer(drip=True) as probe:
            real_create_connection = socket.create_connection

            def connect_probe(address, timeout=None, source_address=None):
                self.assertEqual(address, (PUBLIC_IPV4, probe.port))
                connected = real_create_connection(("127.0.0.1", probe.port), timeout, source_address)
                return _SocketPeerOverride(connected, PUBLIC_IPV4)

            with patch(
                "backend.apps.notes.services.link_fetcher.socket.getaddrinfo",
                return_value=_dns_record(PUBLIC_IPV4, probe.port),
            ), patch(
                "backend.apps.notes.services.link_fetcher.socket.create_connection",
                side_effect=connect_probe,
            ):
                started = time.monotonic()
                with self.assertRaises(SafeLinkError) as raised:
                    fetch_safe_link(f"http://slow.test:{probe.port}/resource")
                elapsed = time.monotonic() - started

            self.assertEqual(raised.exception.code, "note_link_timeout")
            self.assertLessEqual(elapsed, 5.5)


class ParseLeasePolicyTests(AttachmentTestCase):
    def test_attachment_lease_fields_have_durable_defaults(self) -> None:
        attachment = self.create_attachment()

        self.assertEqual(attachment.parse_attempts, 0)
        self.assertIsNone(attachment.parse_lease_token)

    @patch("backend.apps.notes.services.attachment_service.extract_attachment_text")
    def test_duplicate_fresh_delivery_for_active_lease_skips(self, extract) -> None:
        from backend.apps.notes.tasks import parse_attachment_task

        lease = uuid.uuid4()
        attachment = self.create_attachment(
            parse_status=NoteAttachment.ParseStatus.PROCESSING,
            parse_generation=2,
            parse_attempts=1,
            parse_lease_token=lease,
        )

        result = parse_attachment_task(str(attachment.id), 2, str(lease))

        self.assertTrue(result["skipped"])
        extract.assert_not_called()

    @patch("backend.apps.notes.tasks._task_was_redelivered", return_value=True)
    @patch("backend.apps.notes.services.attachment_service.extract_attachment_text", return_value=("recovered", {}))
    def test_redelivery_reclaims_same_lease_and_increments_attempt(self, _extract, _redelivered) -> None:
        from backend.apps.notes.tasks import parse_attachment_task

        lease = uuid.uuid4()
        attachment = self.create_attachment(
            parse_status=NoteAttachment.ParseStatus.PROCESSING,
            parse_generation=3,
            parse_attempts=1,
            parse_lease_token=lease,
        )

        result = parse_attachment_task(str(attachment.id), 3, str(lease))

        attachment.refresh_from_db()
        self.assertEqual(result["status"], "succeeded")
        self.assertEqual(attachment.parse_attempts, 2)

    @patch("backend.apps.notes.tasks._task_was_redelivered", return_value=True)
    @patch("backend.apps.notes.services.attachment_service.extract_attachment_text")
    def test_redelivery_after_attempt_limit_terminally_fails_without_parsing(self, extract, _redelivered) -> None:
        from backend.apps.notes.tasks import MAX_PARSE_ATTEMPTS, parse_attachment_task

        lease = uuid.uuid4()
        attachment = self.create_attachment(
            parse_status=NoteAttachment.ParseStatus.PROCESSING,
            parse_generation=4,
            parse_attempts=MAX_PARSE_ATTEMPTS,
            parse_lease_token=lease,
        )

        result = parse_attachment_task(str(attachment.id), 4, str(lease))

        attachment.refresh_from_db()
        self.assertEqual(result["status"], "failed")
        self.assertEqual(attachment.parse_status, NoteAttachment.ParseStatus.FAILED)
        self.assertEqual(attachment.error_code, "attachment_parse_attempts_exhausted")
        extract.assert_not_called()

    @patch("backend.apps.notes.services.attachment_service.extract_attachment_text")
    def test_old_lease_cannot_overwrite_current_generation(self, extract) -> None:
        from backend.apps.notes.tasks import parse_attachment_task

        current_lease = uuid.uuid4()
        attachment = self.create_attachment(
            parse_generation=5,
            parse_lease_token=current_lease,
        )

        result = parse_attachment_task(str(attachment.id), 4, str(uuid.uuid4()))

        self.assertTrue(result["skipped"])
        extract.assert_not_called()


class ParserIterationBudgetTests(AttachmentTestCase):
    def test_docx_stops_reading_paragraphs_when_character_budget_is_full(self) -> None:
        from docx import Document

        from backend.domains.knowledge_base.docx_parser import parse_docx_blocks

        path = Path(self.storage.name) / "many-paragraphs.docx"
        document = Document()
        for _ in range(100):
            document.add_paragraph("x")
        document.save(path)

        with patch(
            "backend.domains.knowledge_base.docx_parser._paragraph_markdown",
            side_effect=lambda _paragraph, text: text,
        ) as render:
            blocks = parse_docx_blocks(path, max_blocks=10, max_chars=5)

        self.assertLessEqual(sum(len(block.markdown) for block in blocks), 5)
        self.assertLessEqual(render.call_count, 3)

    def test_pptx_shape_budget_does_not_materialize_entire_collection(self) -> None:
        from backend.domains.knowledge_base.attachment_text_parser import parse_attachment_text

        class ShapeStream:
            def __init__(self):
                self.iterated = 0

            def __iter__(self):
                for _ in range(1_000):
                    self.iterated += 1
                    yield SimpleNamespace(has_text_frame=True, text="x", has_table=False)

        shapes = ShapeStream()
        presentation = SimpleNamespace(
            slides=[SimpleNamespace(shapes=shapes)],
        )
        path = Path(self.storage.name) / "slides.pptx"
        path.write_bytes(b"placeholder")

        with patch("pptx.Presentation", return_value=presentation):
            result = parse_attachment_text(
                path,
                "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                max_chars=1,
            )

        self.assertEqual(result.text, "x")
        self.assertEqual(shapes.iterated, 1)

    def test_xlsx_iteration_passes_column_budget_to_openpyxl(self) -> None:
        from backend.domains.knowledge_base.attachment_text_parser import (
            MAX_XLSX_CELLS_PER_ROW,
            MAX_XLSX_ROWS_PER_SHEET,
            parse_attachment_text,
        )

        worksheet = Mock(title="Sheet1")
        worksheet.iter_rows.return_value = [("value",)]
        workbook = Mock(worksheets=[worksheet])
        path = Path(self.storage.name) / "sheet.xlsx"
        path.write_bytes(b"placeholder")

        with patch("openpyxl.load_workbook", return_value=workbook):
            parse_attachment_text(
                path,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                max_chars=100,
            )

        worksheet.iter_rows.assert_called_once_with(
            max_row=MAX_XLSX_ROWS_PER_SHEET,
            max_col=MAX_XLSX_CELLS_PER_ROW,
            values_only=True,
        )
