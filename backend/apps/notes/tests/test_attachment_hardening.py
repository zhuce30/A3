from __future__ import annotations

import json
import socket
import tempfile
import uuid
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from django.db import connection
from django.test import override_settings
from django.urls import reverse

from backend.apps.notes.models import NoteAttachment
from backend.apps.notes.tests.test_attachments import AttachmentTestCase


PUBLIC_IPV4 = "93.184.216.34"


def dns_result(address: str = PUBLIC_IPV4):
    family = socket.AF_INET6 if ":" in address else socket.AF_INET
    return [(family, socket.SOCK_STREAM, 6, "", (address, 443))]


def http_response(
    *,
    peer: str | None = PUBLIC_IPV4,
    status: int = 200,
    headers: dict[str, str] | None = None,
    chunks: list[bytes] | None = None,
):
    response = Mock(status_code=status, headers=headers or {"Content-Type": "text/html"})
    response.iter_content.return_value = chunks if chunks is not None else [b"safe body"]
    if peer is None:
        response.raw = SimpleNamespace()
    else:
        sock = Mock()
        sock.getpeername.return_value = (peer, 443)
        response.raw = SimpleNamespace(_connection=SimpleNamespace(sock=sock))
    return response


class SsrfTransportHardeningTests(AttachmentTestCase):
    @patch("backend.apps.notes.services.link_fetcher.requests.Session")
    @patch("backend.apps.notes.services.link_fetcher.socket.getaddrinfo", return_value=dns_result())
    def test_session_disables_environment_transport_state(self, _dns, session_factory) -> None:
        from backend.apps.notes.services.link_fetcher import fetch_safe_link

        session = session_factory.return_value
        session.get.return_value = http_response()

        fetch_safe_link("https://example.test/reference")

        self.assertFalse(session.trust_env)
        self.assertIsNone(session.auth)
        session.cookies.clear.assert_called()
        session.proxies.clear.assert_called()
        self.assertNotIn("auth", session.get.call_args.kwargs)
        self.assertNotIn("cookies", session.get.call_args.kwargs)
        self.assertIsInstance(session.get.call_args.kwargs["timeout"], tuple)
        session.close.assert_called_once()

    @patch("backend.apps.notes.services.link_fetcher.requests.Session")
    @patch("backend.apps.notes.services.link_fetcher.socket.getaddrinfo", return_value=dns_result())
    def test_missing_peer_fails_closed_and_closes_response(self, _dns, session_factory) -> None:
        from backend.apps.notes.services.link_fetcher import SafeLinkError, fetch_safe_link

        response = http_response(peer=None)
        session_factory.return_value.get.return_value = response

        with self.assertRaises(SafeLinkError) as raised:
            fetch_safe_link("https://example.test/reference")

        self.assertEqual(raised.exception.code, "unsafe_note_link")
        response.close.assert_called_once()

    @patch("backend.apps.notes.services.link_fetcher.requests.Session")
    @patch("backend.apps.notes.services.link_fetcher.socket.getaddrinfo", return_value=dns_result())
    def test_peer_outside_approved_dns_set_is_rejected(self, _dns, session_factory) -> None:
        from backend.apps.notes.services.link_fetcher import SafeLinkError, fetch_safe_link

        session_factory.return_value.get.return_value = http_response(peer="93.184.216.35")

        with self.assertRaises(SafeLinkError) as raised:
            fetch_safe_link("https://example.test/reference")

        self.assertEqual(raised.exception.code, "unsafe_note_link")

    @patch("backend.apps.notes.services.link_fetcher.requests.Session")
    @patch(
        "backend.apps.notes.services.link_fetcher.socket.getaddrinfo",
        return_value=dns_result("::ffff:93.184.216.34"),
    )
    def test_ipv4_mapped_ipv6_matches_ipv4_peer(self, _dns, session_factory) -> None:
        from backend.apps.notes.services.link_fetcher import fetch_safe_link

        session_factory.return_value.get.return_value = http_response(peer=PUBLIC_IPV4)

        result = fetch_safe_link("https://example.test/reference")

        self.assertEqual(result.body, b"safe body")

    @patch("backend.apps.notes.services.link_fetcher.time.monotonic")
    @patch("backend.apps.notes.services.link_fetcher.requests.Session")
    @patch("backend.apps.notes.services.link_fetcher.socket.getaddrinfo", return_value=dns_result())
    def test_total_deadline_applies_during_streaming(self, _dns, session_factory, monotonic) -> None:
        from backend.apps.notes.services.link_fetcher import SafeLinkError, fetch_safe_link

        monotonic.side_effect = [0.0, 0.1, 0.2, 0.3, 0.4, 0.5, 1.0, 1.1, 4.9, 5.1]
        response = http_response(chunks=[b"one", b"two", b"three"])
        session_factory.return_value.get.return_value = response

        with self.assertRaises(SafeLinkError) as raised:
            fetch_safe_link("https://example.test/reference")

        self.assertEqual(raised.exception.code, "note_link_timeout")
        response.close.assert_called_once()


class LinkCreationHardeningTests(AttachmentTestCase):
    @patch("backend.apps.notes.services.attachment_service.fetch_safe_link")
    def test_network_fetch_happens_outside_database_transaction(self, fetch) -> None:
        from backend.apps.notes.services.link_fetcher import SafeLinkResponse
        from backend.apps.notes.services.attachment_service import create_link_attachment

        baseline_atomic_depth = len(connection.atomic_blocks)

        def assert_short_transaction(_url):
            self.assertEqual(len(connection.atomic_blocks), baseline_atomic_depth)
            return SafeLinkResponse(
                final_url="https://example.test/reference",
                body=b"<p>validated text</p>",
                mime_type="text/html",
            )

        fetch.side_effect = assert_short_transaction
        attachment = create_link_attachment(
            self.user,
            self.note.id,
            {"url": "https://example.test/reference", "displayName": "reference"},
        )

        self.assertEqual(attachment.parse_status, NoteAttachment.ParseStatus.SUCCEEDED)
        self.assertEqual(attachment.extracted_text, "validated text")

    @patch("backend.apps.notes.views_attachments.enqueue_attachment_parse")
    @patch("backend.apps.notes.services.attachment_service.fetch_safe_link")
    def test_initial_link_creation_fetches_once_and_does_not_enqueue(self, fetch, enqueue) -> None:
        from backend.apps.notes.services.link_fetcher import SafeLinkResponse

        fetch.return_value = SafeLinkResponse(
            final_url="https://example.test/reference",
            body=b"<p>single fetch</p>",
            mime_type="text/html",
        )
        response = self.client.post(
            reverse("notes:attachment-link", args=[self.note.id]),
            data=json.dumps({"url": "https://example.test/reference", "displayName": "reference"}),
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 201)
        attachment = NoteAttachment.objects.get()
        self.assertEqual(attachment.parse_status, NoteAttachment.ParseStatus.SUCCEEDED)
        self.assertEqual(attachment.extracted_text, "single fetch")
        fetch.assert_called_once()
        enqueue.assert_not_called()


class QuotaHardeningTests(AttachmentTestCase):
    @override_settings(A3_NOTES_USER_QUOTA_BYTES=8)
    def test_soft_deleted_physical_file_still_counts_toward_quota(self) -> None:
        self.create_attachment(size_bytes=5, deleted_at="2026-08-02T00:00:00Z")

        response = self.upload("more.txt", b"four", "text/plain")

        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.json()["code"], "attachment_quota_exceeded")


class ParseGenerationHardeningTests(AttachmentTestCase):
    def test_attachment_generation_fields_have_durable_defaults(self) -> None:
        attachment = self.create_attachment()

        self.assertEqual(attachment.parse_generation, 1)
        self.assertIsNone(attachment.parse_started_at)

    @patch("backend.apps.notes.views_attachments.enqueue_attachment_parse")
    def test_each_reparse_increments_generation_under_lock(self, enqueue) -> None:
        attachment = self.create_attachment(parse_status=NoteAttachment.ParseStatus.SUCCEEDED)
        enqueue.side_effect = lambda item: item
        url = reverse("notes:attachment-reparse", args=[attachment.id])

        first = self.client.post(url)
        second = self.client.post(url)

        self.assertEqual(first.status_code, 202)
        self.assertEqual(second.status_code, 202)
        attachment.refresh_from_db()
        self.assertEqual(attachment.parse_generation, 3)
        self.assertEqual([call.args[0].parse_generation for call in enqueue.call_args_list], [2, 3])

    @patch("backend.apps.notes.services.attachment_service.extract_attachment_text")
    def test_old_generation_cannot_overwrite_reparse(self, extract) -> None:
        from backend.apps.notes.tasks import parse_attachment_task

        attachment = self.create_attachment(parse_generation=1)

        def finish_after_reparse(_attachment):
            NoteAttachment.objects.filter(pk=attachment.pk).update(
                parse_generation=2,
                parse_status=NoteAttachment.ParseStatus.PENDING,
                extracted_text="new generation",
            )
            return "old generation", {"parser": "test"}

        extract.side_effect = finish_after_reparse
        result = parse_attachment_task(str(attachment.id), 1)

        attachment.refresh_from_db()
        self.assertTrue(result["skipped"])
        self.assertEqual(attachment.parse_generation, 2)
        self.assertEqual(attachment.extracted_text, "new generation")

    @patch("backend.apps.notes.services.attachment_service.extract_attachment_text")
    def test_old_generation_failure_is_reported_as_skipped(self, extract) -> None:
        from backend.apps.notes.tasks import parse_attachment_task

        attachment = self.create_attachment(parse_generation=1)

        def fail_after_reparse(_attachment):
            NoteAttachment.objects.filter(pk=attachment.pk).update(
                parse_generation=2,
                parse_status=NoteAttachment.ParseStatus.PENDING,
                error_code="",
            )
            raise ValueError("old task failure")

        extract.side_effect = fail_after_reparse

        result = parse_attachment_task(str(attachment.id), 1)

        attachment.refresh_from_db()
        self.assertTrue(result["skipped"])
        self.assertEqual(attachment.parse_generation, 2)
        self.assertEqual(attachment.error_code, "")

    @patch("backend.apps.notes.services.attachment_service.extract_attachment_text", return_value=("recovered", {}))
    def test_same_generation_redelivery_recovers_processing_state(self, _extract) -> None:
        from backend.apps.notes.tasks import parse_attachment_task

        attachment = self.create_attachment(
            parse_generation=4,
            parse_status=NoteAttachment.ParseStatus.PROCESSING,
        )

        result = parse_attachment_task(str(attachment.id), 4)

        attachment.refresh_from_db()
        self.assertEqual(result["status"], "succeeded")
        self.assertEqual(attachment.extracted_text, "recovered")

    def test_parse_task_uses_worker_loss_safe_delivery_options(self) -> None:
        from backend.apps.notes.tasks import parse_attachment_task

        self.assertTrue(parse_attachment_task.acks_late)
        self.assertTrue(parse_attachment_task.reject_on_worker_lost)


class BrokerFailureHardeningTests(AttachmentTestCase):
    @patch("backend.apps.notes.tasks.parse_attachment_task.delay", side_effect=RuntimeError("broker secret"))
    def test_enqueue_failure_is_persisted_without_raising(self, _delay) -> None:
        from backend.apps.notes.tasks import enqueue_attachment_parse

        attachment = self.create_attachment(parse_generation=2)

        returned = enqueue_attachment_parse(attachment)

        returned.refresh_from_db()
        self.assertEqual(returned.error_code, "attachment_enqueue_failed")
        self.assertEqual(returned.error_message, "附件解析任务暂时无法提交，请重试")
        self.assertEqual(returned.parse_status, NoteAttachment.ParseStatus.FAILED)
        self.assertNotIn("secret", returned.error_message)

    @patch("backend.apps.notes.tasks.parse_attachment_task.delay", side_effect=RuntimeError("broker unavailable"))
    def test_upload_returns_created_attachment_when_broker_fails(self, _delay) -> None:
        response = self.upload("notes.txt", b"safe", "text/plain")

        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.json()["attachment"]["errorCode"], "attachment_enqueue_failed")
        self.assertEqual(NoteAttachment.objects.get().parse_status, NoteAttachment.ParseStatus.FAILED)

    def test_reparse_broker_failure_is_visible_and_retryable(self) -> None:
        attachment = self.create_attachment(parse_status=NoteAttachment.ParseStatus.SUCCEEDED)
        url = reverse("notes:attachment-reparse", args=[attachment.id])

        with patch(
            "backend.apps.notes.tasks.parse_attachment_task.delay",
            side_effect=RuntimeError("broker unavailable"),
        ):
            failed = self.client.post(url)

        self.assertEqual(failed.status_code, 202)
        self.assertEqual(failed.json()["attachment"]["errorCode"], "attachment_enqueue_failed")
        attachment.refresh_from_db()
        self.assertEqual(attachment.parse_generation, 2)
        self.assertEqual(attachment.parse_status, NoteAttachment.ParseStatus.FAILED)

        with patch("backend.apps.notes.tasks.parse_attachment_task.delay") as delay:
            retried = self.client.post(url)

        attachment.refresh_from_db()
        self.assertEqual(retried.status_code, 202)
        self.assertEqual(attachment.parse_generation, 3)
        self.assertEqual(attachment.parse_status, NoteAttachment.ParseStatus.PENDING)
        delay.assert_called_once_with(str(attachment.id), 3)


class AttachmentSafeParserTests(AttachmentTestCase):
    @patch("backend.domains.knowledge_base.docx_parser.parse_docx_blocks")
    def test_docx_reuse_passes_hard_work_and_character_budgets(self, parse_blocks) -> None:
        from backend.domains.knowledge_base.attachment_text_parser import (
            MAX_DOCX_BLOCKS,
            parse_attachment_text,
        )

        parse_blocks.return_value = [SimpleNamespace(markdown="bounded docx")]
        path = Path(self.storage.name) / "document.docx"
        path.write_bytes(b"placeholder")

        result = parse_attachment_text(
            path,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            max_chars=20,
        )

        self.assertEqual(result.text, "bounded docx")
        parse_blocks.assert_called_once_with(path, max_blocks=MAX_DOCX_BLOCKS, max_chars=20)

    @patch("backend.domains.knowledge_base.multimodal_parser._parse_pdf", side_effect=AssertionError("private parser called"))
    @patch("backend.domains.knowledge_base.multimodal_parser._render_pdf_pages", side_effect=AssertionError("render called"))
    @patch("backend.domains.knowledge_base.multimodal_parser._understand_image", side_effect=AssertionError("provider called"))
    @patch("pypdf.PdfReader")
    def test_pdf_parser_is_text_only_bounded_and_never_calls_multimodal_pipeline(
        self,
        reader,
        _provider,
        _render,
        _private_parser,
    ) -> None:
        from backend.domains.knowledge_base.attachment_text_parser import parse_attachment_text

        first_page = Mock()
        first_page.extract_text.return_value = "x" * 100
        second_page = Mock()
        second_page.extract_text.return_value = "should not be read"
        reader.return_value.pages = [first_page, second_page]
        path = Path(self.storage.name) / "misleading.txt"
        path.write_bytes(b"%PDF-1.7")

        result = parse_attachment_text(path, "application/pdf", max_chars=10)

        self.assertEqual(result.text, "x" * 10)
        first_page.extract_text.assert_called_once()
        second_page.extract_text.assert_not_called()
        self.assertFalse(list(Path(self.storage.name).rglob("*.a3_rendered*")))

    @patch("backend.domains.knowledge_base.multimodal_parser._render_pptx_slides", side_effect=AssertionError("render called"))
    @patch("backend.domains.knowledge_base.multimodal_parser._understand_image", side_effect=AssertionError("provider called"))
    def test_pptx_parser_creates_no_render_side_effects(self, _provider, _render) -> None:
        from pptx import Presentation

        from backend.domains.knowledge_base.attachment_text_parser import parse_attachment_text

        path = Path(self.storage.name) / "slides.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = "bounded slide text"
        presentation.save(path)

        result = parse_attachment_text(
            path,
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            max_chars=100,
        )

        self.assertIn("bounded slide text", result.text)
        self.assertFalse(list(Path(self.storage.name).rglob("*.a3_rendered*")))

    @patch("backend.domains.knowledge_base.attachment_text_parser.parse_attachment_text")
    def test_notes_service_reuses_public_attachment_safe_parser(self, parser) -> None:
        from backend.apps.notes.services.attachment_service import extract_attachment_text

        parser.return_value = SimpleNamespace(text="safe parser", metadata={"parser": "attachment_safe_pdf"})
        attachment = self.create_attachment(mime_type="application/pdf", display_name="wrong.txt")

        text, metadata = extract_attachment_text(attachment)

        self.assertEqual(text, "safe parser")
        self.assertEqual(metadata["parser"], "attachment_safe_pdf")
        parser.assert_called_once()
