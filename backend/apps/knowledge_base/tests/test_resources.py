from __future__ import annotations

import tempfile
import json
import os
from datetime import timedelta
from io import StringIO
from pathlib import Path
from unittest.mock import MagicMock, patch

from django.core.management import call_command
from django.core.files.uploadedfile import SimpleUploadedFile
from PIL import Image
from django.test import Client, SimpleTestCase, TestCase, override_settings
from django.utils import timezone

from backend.apps.accounts.models import User, UserRole
from backend.apps.knowledge_base.models import (
    IndexJobStatus,
    IndexJobType,
    KnowledgeAsset,
    KnowledgeGraphClaim,
    KnowledgeGraphEntity,
    KnowledgeIndexJob,
    KnowledgeBlock,
    KnowledgeConcept,
    KnowledgeDefinition,
    KnowledgeIndexEntry,
    KnowledgePage,
    KnowledgeParseRun,
    KnowledgeResource,
    ResourceScope,
    ResourceStatus,
)
from backend.apps.knowledge_base.display import format_source_type_label, source_origin_label
from backend.apps.knowledge_base.services import clear_text_index_storage, create_parse_job, create_upload_index_job, create_wiki_rebuild_job, delete_resource, knowledge_base_stats, reparse_resource, retry_index_job, save_and_index_upload, serialize_index_job, serialize_resource
from backend.apps.knowledge_base.tasks import run_index_job
from backend.domains.knowledge_base.document_parse_v2 import finalize_parse_result
from backend.domains.knowledge_base.multimodal_parser import BaiduDocumentParserProvider, MultimodalParseResult, ParsedBlock, ParsedPage, QwenOcrDocumentParserProvider, VisionUnderstanding
from backend.domains.knowledge_base.document_loader import iter_source_files


MINIMAL_TEXT_PDF = b"""%PDF-1.4
1 0 obj
<< /Type /Catalog /Pages 2 0 R >>
endobj
2 0 obj
<< /Type /Pages /Kids [3 0 R] /Count 1 >>
endobj
3 0 obj
<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>
endobj
4 0 obj
<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>
endobj
5 0 obj
<< /Length 55 >>
stream
BT /F1 18 Tf 72 720 Td (Root locus asymptote rule) Tj ET
endstream
endobj
xref
0 6
0000000000 65535 f 
0000000009 00000 n 
0000000058 00000 n 
0000000115 00000 n 
0000000241 00000 n 
0000000311 00000 n 
trailer
<< /Size 6 /Root 1 0 R >>
startxref
416
%%EOF
"""


class ParsedBlockMarkdownMetadataTests(SimpleTestCase):
    def test_baidu_document_parser_provider_degrades_without_token(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            image_path = Path(temp_dir) / "page.png"
            image_path.write_bytes(b"fake-image")
            with patch.dict(os.environ, {"A3_BAIDU_PADDLEOCR_TOKEN": "", "BAIDU_PADDLEOCR_TOKEN": ""}, clear=False):
                result = BaiduDocumentParserProvider().understand_image(image_path)

        self.assertEqual(result.to_search_text(), "")
        self.assertIn("A3_BAIDU_PADDLEOCR_TOKEN", result.warnings[0])

    @patch("backend.domains.knowledge_base.multimodal_parser._poll_baidu_paddleocr_job")
    @patch("backend.domains.knowledge_base.multimodal_parser._submit_baidu_paddleocr_job")
    def test_baidu_document_parser_provider_returns_structured_result(self, submit_job, poll_job):
        submit_job.return_value = {"data": {"jobId": "job-1"}}
        poll_job.return_value = {
            "data": {
                "status": "succeeded",
                "markdown": "# 进程调度\n| 算法 | 特点 |\n|---|---|\n| RR | 时间片 |",
            }
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            image_path = Path(temp_dir) / "page.png"
            image_path.write_bytes(b"fake-image")
            with patch.dict(os.environ, {"A3_BAIDU_PADDLEOCR_TOKEN": "token"}, clear=False):
                result = BaiduDocumentParserProvider().understand_image(image_path)

        self.assertIn("进程调度", result.to_search_text())
        self.assertEqual(result.raw["provider"], "baidu_paddleocr_vl")
        self.assertEqual(result.raw["jobId"], "job-1")
        self.assertIsInstance(result.raw["latencyMs"], int)

    @patch("backend.domains.knowledge_base.multimodal_parser._poll_baidu_paddleocr_job")
    @patch("backend.domains.knowledge_base.multimodal_parser._submit_baidu_paddleocr_job")
    def test_baidu_document_parser_provider_retries_empty_result_with_enhanced_options(self, submit_job, poll_job):
        submit_job.side_effect = [{"data": {"jobId": "job-empty"}}, {"data": {"jobId": "job-retry"}}]
        poll_job.side_effect = [
            {"data": {"state": "done"}},
            {"data": {"state": "done", "markdown": "Enhanced OCR text"}},
        ]
        with tempfile.TemporaryDirectory() as temp_dir:
            image_path = Path(temp_dir) / "page.png"
            image_path.write_bytes(b"fake-image")
            with patch.dict(os.environ, {"A3_BAIDU_PADDLEOCR_TOKEN": "token", "A3_BAIDU_PADDLEOCR_RETRY_EMPTY": "1"}, clear=False):
                result = BaiduDocumentParserProvider().understand_image(image_path)

        self.assertEqual(result.to_search_text(), "Enhanced OCR text")
        self.assertEqual(submit_job.call_count, 2)
        self.assertFalse(submit_job.call_args_list[0].kwargs["optional_payload"]["useDocUnwarping"])
        self.assertTrue(submit_job.call_args_list[1].kwargs["optional_payload"]["useDocUnwarping"])

    @patch(
        "backend.domains.knowledge_base.multimodal_parser._download_text",
        return_value='{"result":{"layoutParsingResults":[{"markdown":{"text":"# Deadlock\\n| Condition | Meaning |\\n|---|---|\\n| Hold and wait | holds resources |"}}]}}\n',
    )
    def test_baidu_document_parser_provider_reads_jsonl_result_url(self, _download_text):
        from backend.domains.knowledge_base.multimodal_parser import _baidu_result_to_vision_understanding

        result = _baidu_result_to_vision_understanding(
            {"data": {"state": "done", "resultUrl": {"jsonUrl": "https://example.test/result.jsonl"}}},
            job_id="job-jsonl",
        )

        self.assertIn("Deadlock", result.to_search_text())
        self.assertIn("Hold and wait", result.to_search_text())
        self.assertEqual(result.raw["jobId"], "job-jsonl")

    @patch("backend.domains.knowledge_base.multimodal_parser._render_pdf_to_assets", return_value=({}, []))
    def test_pdf_render_pages_uses_configured_page_limit(self, render_pdf_to_assets):
        from backend.domains.knowledge_base.multimodal_parser import _render_pdf_pages

        with tempfile.TemporaryDirectory() as temp_dir:
            pdf_path = Path(temp_dir) / "lesson.pdf"
            pdf_path.write_bytes(MINIMAL_TEXT_PDF)
            with patch.dict(os.environ, {"A3_PDF_RENDER_MAX_PAGES": "72"}, clear=False):
                _render_pdf_pages(pdf_path)

        self.assertEqual(render_pdf_to_assets.call_args.kwargs["max_pages"], 72)

    def test_parsed_block_adds_markdown_metadata_without_changing_raw_text(self):
        block = ParsedBlock(
            modality="text",
            text="系统方程为：s² + 2s + Kg = 0",
            metadata={"sourceType": "pptx", "blockType": "text_box"},
        )

        self.assertEqual(block.text, "系统方程为：s² + 2s + Kg = 0")
        self.assertEqual(block.metadata["contentFormat"], "markdown")
        self.assertEqual(block.metadata["markdownText"], "系统方程为：$s^2 + 2s + Kg = 0$")
        self.assertTrue(block.metadata["rawTextPreserved"])
        self.assertEqual(block.metadata["formatFeatures"], ["math"])

    def test_document_parse_v2_adds_provider_order_and_quality_metrics(self):
        result = finalize_parse_result(
            MultimodalParseResult(
                parser="native",
                status="completed",
                pages=[
                    ParsedPage(
                        page_number=1,
                        blocks=[
                            ParsedBlock(
                                modality="text",
                                text="操作系统负责资源管理和抽象接口。",
                                confidence=0.92,
                                metadata={"sourceType": "pdf"},
                            ),
                            ParsedBlock(
                                modality="table",
                                text="| 调度算法 | 特点 |\n|---|---|\n| RR | 时间片 |",
                                confidence=0.86,
                                metadata={"sourceType": "pdf", "blockType": "table"},
                            ),
                        ],
                    )
                ],
            )
        )

        document_parse = result.result["documentParse"]
        first_block = result.pages[0].blocks[0]
        self.assertEqual(document_parse["schemaVersion"], "a3.document-parse.v2")
        self.assertTrue(document_parse["citationReadiness"])
        self.assertEqual(document_parse["providerChain"], ["native_pdf"])
        self.assertEqual(first_block.metadata["documentParseSchemaVersion"], "a3.document-parse.v2")
        self.assertEqual(first_block.metadata["readingOrderIndex"], 1)
        self.assertEqual(first_block.metadata["provider"], "native_pdf")


class KnowledgeBaseDocumentDiscoveryTests(SimpleTestCase):
    def test_source_discovery_ignores_legacy_ppt_files(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            (root / "legacy.ppt").write_bytes(b"legacy")
            (root / "lesson.pptx").write_bytes(b"pptx")
            (root / "slides.pdf").write_bytes(b"pdf")
            (root / "qa.xlsx").write_bytes(b"xlsx")

            files = iter_source_files(root)

        self.assertEqual({path.suffix.lower() for path in files}, {".pptx", ".pdf", ".xlsx"})
        self.assertNotIn("legacy.ppt", {path.name for path in files})

    def test_source_display_metadata_labels_full_index_source_types(self):
        self.assertEqual(format_source_type_label({"source_type": "qa_xlsx"}), "问答题库")
        self.assertEqual(format_source_type_label({"resource_type": "courseware"}), "PPTX 课件")
        self.assertEqual(format_source_type_label({"source_path": "lesson.pdf"}), "PDF 课件")
        self.assertEqual(source_origin_label({"source_type": "pdf"}), "全量知识库索引")
        self.assertEqual(source_origin_label({"resource_id": "upload-1"}), "教师上传索引")


class KnowledgeResourceTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(phone="13800139000", password="secure-pass-123")
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.visual_embed_patch = patch("backend.domains.knowledge_base.visual_retriever._embed_image", return_value=[1.0, 0.0])
        self.visual_embed_patch.start()
        self.addCleanup(self.visual_embed_patch.stop)

    @patch("backend.domains.knowledge_base.index_manager.append_source_file", return_value=(3, 8, []))
    def test_upload_is_saved_and_indexed(self, append_source_file):
        upload = SimpleUploadedFile("lesson.pdf", b"fake-pdf", content_type="application/pdf")

        with override_settings(KNOWLEDGE_UPLOAD_ROOT=self.temp_dir.name):
            resource = save_and_index_upload(self.user, upload)

        self.assertEqual(resource.status, ResourceStatus.INDEXED)
        self.assertEqual(resource.document_count, 3)
        self.assertEqual(resource.chunk_count, 8)
        self.assertTrue(Path(resource.stored_path).exists())
        append_source_file.assert_called_once()

    def test_resource_scope_defaults_to_teacher_for_existing_upload_helpers(self):
        resource = KnowledgeResource.objects.create(
            owner=self.user,
            original_name="lesson.pdf",
            stored_path=str((Path(self.temp_dir.name) / "lesson.pdf").resolve()),
            file_type=".pdf",
            file_size=10,
        )

        self.assertEqual(resource.resource_scope, ResourceScope.TEACHER)
        self.assertEqual(resource.visibility, "owner")

    def test_clear_text_index_storage_removes_stale_build_report(self):
        import backend.apps.knowledge_base.services as services

        faiss_dir = Path(self.temp_dir.name) / "faiss_index"
        metadata_dir = Path(self.temp_dir.name) / "metadata"
        faiss_dir.mkdir()
        metadata_dir.mkdir()
        (faiss_dir / "default__vector_store.json").write_text("{}", encoding="utf-8")
        build_report = metadata_dir / "build_report.json"
        build_report.write_text('{"source_files_seen": 105, "documents_loaded": 7381}', encoding="utf-8")

        with patch.object(services, "FAISS_INDEX_DIR", faiss_dir), patch.object(services, "METADATA_DIR", metadata_dir):
            clear_text_index_storage()

        self.assertFalse(faiss_dir.exists())
        self.assertFalse(build_report.exists())

    def test_legacy_ppt_is_rejected(self):
        upload = SimpleUploadedFile("legacy.ppt", b"legacy", content_type="application/octet-stream")

        with override_settings(KNOWLEDGE_UPLOAD_ROOT=self.temp_dir.name):
            with self.assertRaisesRegex(ValueError, "旧版 .ppt 不能直接入库，请先转换为 .pptx"):
                save_and_index_upload(self.user, upload)

        self.assertFalse(KnowledgeResource.objects.exists())

    def test_unsupported_upload_type_returns_readable_chinese_error(self):
        upload = SimpleUploadedFile("lesson.txt", b"text", content_type="text/plain")

        with override_settings(KNOWLEDGE_UPLOAD_ROOT=self.temp_dir.name):
            with self.assertRaisesRegex(ValueError, "仅支持 DOCX、PDF、PPTX、XLSX、PNG、JPG、JPEG、WEBP 文件"):
                save_and_index_upload(self.user, upload)

        self.assertFalse(KnowledgeResource.objects.exists())

    def test_oversized_upload_returns_readable_chinese_error(self):
        upload = SimpleUploadedFile("lesson.pdf", b"fake-pdf", content_type="application/pdf")

        with override_settings(KNOWLEDGE_UPLOAD_ROOT=self.temp_dir.name, KNOWLEDGE_UPLOAD_MAX_BYTES=1):
            with self.assertRaisesRegex(ValueError, "单个文件不能超过 200MB"):
                save_and_index_upload(self.user, upload)

        self.assertFalse(KnowledgeResource.objects.exists())

    def test_reparse_recovers_resource_stuck_in_indexing_status(self):
        from pptx import Presentation

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = "Frequency response"
        pptx_path = Path(self.temp_dir.name) / "frequency-response.pptx"
        presentation.save(pptx_path)
        resource = KnowledgeResource.objects.create(
            owner=self.user,
            original_name="frequency-response.pptx",
            stored_path=str(pptx_path),
            file_type=".pptx",
            file_size=pptx_path.stat().st_size,
            status=ResourceStatus.INDEXING,
        )

        reparse_resource(resource)
        resource.refresh_from_db()

        self.assertIn(resource.status, {ResourceStatus.INDEXED, ResourceStatus.INDEXED_WITH_WARNINGS})
        self.assertGreater(resource.pages.count(), 0)
        self.assertGreater(resource.blocks.count(), 0)

    @patch("backend.domains.knowledge_base.multimodal_parser._understand_image", return_value=VisionUnderstanding(warnings=["OCR engine is not installed"]))
    def test_image_file_type_is_accepted_for_upload_record(self, _understand_image):
        upload = SimpleUploadedFile("diagram.webp", b"fake-webp", content_type="image/webp")

        with override_settings(KNOWLEDGE_UPLOAD_ROOT=self.temp_dir.name):
            resource = save_and_index_upload(self.user, upload)

        self.assertEqual(resource.file_type, ".webp")
        self.assertEqual(resource.status, ResourceStatus.INDEXED_WITH_WARNINGS)

    @patch("backend.domains.knowledge_base.index_manager.append_source_file", return_value=(1, 2, []))
    def test_xlsx_upload_creates_table_block(self, _append_source_file):
        from openpyxl import Workbook

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "根轨迹"
        worksheet.append(["问题", "答案"])
        worksheet.append(["渐近线怎么求", "按极点零点数量计算"])
        xlsx_path = Path(self.temp_dir.name) / "qa.xlsx"
        workbook.save(xlsx_path)
        upload = SimpleUploadedFile(
            "qa.xlsx",
            xlsx_path.read_bytes(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        with override_settings(KNOWLEDGE_UPLOAD_ROOT=self.temp_dir.name):
            resource = save_and_index_upload(self.user, upload)

        self.assertEqual(resource.status, ResourceStatus.INDEXED)
        self.assertEqual(resource.pages.count(), 1)
        self.assertEqual(resource.blocks.filter(modality="table", text__contains="渐近线").count(), 1)
        self.assertGreaterEqual(serialize_resource(resource)["multimodal"]["textBlocks"], 0)
        self.assertEqual(serialize_resource(resource)["multimodal"]["tableBlocks"], 1)

    @patch("backend.domains.knowledge_base.multimodal_parser._understand_image", return_value=VisionUnderstanding())
    @patch("backend.domains.knowledge_base.index_manager.append_source_file", return_value=(1, 2, []))
    def test_pdf_upload_creates_text_page_block(self, _append_source_file, _understand_image):
        upload = SimpleUploadedFile("lesson.pdf", MINIMAL_TEXT_PDF, content_type="application/pdf")

        with override_settings(KNOWLEDGE_UPLOAD_ROOT=self.temp_dir.name):
            resource = save_and_index_upload(self.user, upload)

        self.assertEqual(resource.status, ResourceStatus.INDEXED)
        self.assertEqual(resource.pages.count(), 1)
        self.assertGreaterEqual(resource.blocks.filter(modality="text", text__contains="Root locus").count(), 1)
        self.assertGreaterEqual(serialize_resource(resource)["multimodal"]["textBlocks"], 1)

    @patch("backend.domains.knowledge_base.multimodal_parser._understand_image", return_value=VisionUnderstanding())
    @patch("backend.domains.knowledge_base.index_manager.append_source_file", return_value=(1, 2, []))
    def test_pdf_upload_renders_page_asset_for_visual_evidence(self, _append_source_file, _understand_image):
        upload = SimpleUploadedFile("lesson.pdf", MINIMAL_TEXT_PDF, content_type="application/pdf")

        with override_settings(KNOWLEDGE_UPLOAD_ROOT=self.temp_dir.name):
            resource = save_and_index_upload(self.user, upload)

        rendered_asset = resource.assets.filter(kind="rendered_page", mime_type="image/png").first()
        self.assertIsNotNone(rendered_asset)
        self.assertTrue(Path(rendered_asset.storage_path).exists())
        self.assertGreater(rendered_asset.width, 0)
        self.assertGreater(rendered_asset.height, 0)
        self.assertEqual(resource.blocks.filter(modality="page_image", asset=rendered_asset).count(), 1)
        self.assertEqual(resource.blocks.get(modality="page_image").index_entries.filter(index_type="visual_vector").count(), 1)

    def test_stats_expose_full_rebuild_index_separately_from_uploaded_resources(self):
        metadata_dir = Path(self.temp_dir.name) / "metadata"
        data_root = Path(self.temp_dir.name) / "data"
        metadata_dir.mkdir()
        data_root.mkdir()
        (metadata_dir / "build_report.json").write_text(
            '{"source_files_seen": 105, "documents_loaded": 7381, "skipped": [{"file": "bad.xlsx"}, {"file": "empty.xlsx"}, {"file": "legacy.ppt"}], "embed_model": "demo-embed"}',
            encoding="utf-8",
        )
        KnowledgeResource.objects.create(
            owner=self.user,
            original_name="teacher-upload.pdf",
            stored_path=str((data_root / "teacher-upload.pdf").resolve()),
            file_type=".pdf",
            file_size=10,
            status=ResourceStatus.INDEXED,
            document_count=4,
        )

        with patch("backend.apps.knowledge_base.services.METADATA_DIR", metadata_dir), patch("backend.apps.knowledge_base.services.DATA_ROOT", data_root):
            stats = knowledge_base_stats(self.user)

        self.assertEqual(stats["globalIndexedFiles"], 105)
        self.assertEqual(stats["globalIndexedDocuments"], 7381)
        self.assertEqual(stats["skippedFiles"], 3)
        self.assertEqual(stats["managedFiles"], 1)
        self.assertEqual(stats["indexedFiles"], 1)
        self.assertEqual(stats["indexedDocuments"], 7381)
        self.assertEqual(stats["uploadIndexedDocuments"], 4)
        self.assertEqual(stats["embedModel"], "demo-embed")

    def test_resource_serialization_includes_multimodal_quality(self):
        path = Path(self.temp_dir.name) / "slide.png"
        path.write_bytes(b"fake-image")
        resource = KnowledgeResource.objects.create(
            owner=self.user,
            original_name="slide.png",
            stored_path=str(path),
            file_type=".png",
            file_size=path.stat().st_size,
            status=ResourceStatus.INDEXED_WITH_WARNINGS,
            document_count=1,
            chunk_count=1,
        )
        parse_run = KnowledgeParseRun.objects.create(
            resource=resource,
            status="completed_with_warnings",
            parser="native",
            warnings=["OCR not configured"],
            quality_score=0.5,
        )
        page = KnowledgePage.objects.create(
            resource=resource,
            parse_run=parse_run,
            page_number=1,
            summary="uploaded image",
            quality_score=0.5,
            warnings=["OCR not configured"],
        )
        asset = KnowledgeAsset.objects.create(
            resource=resource,
            page=page,
            kind="original",
            mime_type="image/png",
            storage_path=str(path),
            sha256="abc123",
        )
        block = KnowledgeBlock.objects.create(
            resource=resource,
            page=page,
            asset=asset,
            modality="image",
            text="uploaded image",
            confidence=0.4,
            warnings=["OCR not configured"],
        )
        KnowledgeIndexEntry.objects.create(
            block=block,
            index_type="visual_vector",
            index_backend="native",
            external_id=str(block.id),
            status="indexed",
        )

        payload = serialize_resource(resource)

        self.assertEqual(payload["status"], ResourceStatus.INDEXED_WITH_WARNINGS)
        self.assertEqual(payload["multimodal"]["parseStatus"], "completed_with_warnings")
        self.assertEqual(payload["multimodal"]["pages"], 1)
        self.assertEqual(payload["multimodal"]["imageBlocks"], 1)
        self.assertEqual(payload["multimodal"]["visualIndexed"], 1)
        self.assertEqual(payload["multimodal"]["qualityScore"], 0.5)
        self.assertEqual(payload["multimodal"]["warnings"], ["OCR not configured"])

    def test_evaluate_multimodal_rag_command_outputs_smoke_metrics(self):
        path = Path(self.temp_dir.name) / "question.png"
        path.write_bytes(b"fake-image")
        resource = KnowledgeResource.objects.create(
            owner=self.user,
            original_name="question.png",
            stored_path=str(path),
            file_type=".png",
            file_size=path.stat().st_size,
            status=ResourceStatus.INDEXED_WITH_WARNINGS,
        )
        parse_run = KnowledgeParseRun.objects.create(
            resource=resource,
            status="completed_with_warnings",
            parser="native",
            warnings=["OCR not configured"],
            quality_score=0.45,
        )
        page = KnowledgePage.objects.create(resource=resource, parse_run=parse_run, page_number=1)
        asset = KnowledgeAsset.objects.create(resource=resource, page=page, kind="original", storage_path=str(path))
        block = KnowledgeBlock.objects.create(resource=resource, page=page, asset=asset, modality="image", text="question")
        KnowledgeIndexEntry.objects.create(block=block, index_type="visual_vector", index_backend="native", status="indexed")
        output = StringIO()

        call_command("evaluate_multimodal_rag", "--top-k", "3", stdout=output)

        text = output.getvalue()
        self.assertIn("parseSuccessRate", text)
        self.assertIn("visualCoverage", text)
        self.assertIn("hybridReadyRate", text)

    @patch("backend.domains.knowledge_base.multimodal_parser._understand_image", return_value=VisionUnderstanding())
    def test_backfill_multimodal_rag_command_parses_existing_resources(self, _understand_image):
        path = Path(self.temp_dir.name) / "legacy-upload.pdf"
        path.write_bytes(MINIMAL_TEXT_PDF)
        resource = KnowledgeResource.objects.create(
            owner=self.user,
            original_name="legacy-upload.pdf",
            stored_path=str(path),
            file_type=".pdf",
            file_size=path.stat().st_size,
            status=ResourceStatus.INDEXED,
        )
        output = StringIO()

        call_command("backfill_multimodal_rag", "--limit", "10", stdout=output)

        resource.refresh_from_db()
        self.assertEqual(resource.parse_runs.count(), 1)
        self.assertIn("parsedResources: 1", output.getvalue())

    @patch("backend.domains.knowledge_base.index_manager.append_source_file", return_value=(1, 2, []))
    @patch("backend.apps.knowledge_base.services.remove_visual_resource_entries")
    def test_delete_removes_database_record_file_assets_and_visual_index(self, remove_visual_resource_entries, _append_source_file):
        upload = SimpleUploadedFile("lesson.xlsx", b"xlsx", content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        with override_settings(KNOWLEDGE_UPLOAD_ROOT=self.temp_dir.name):
            resource = save_and_index_upload(self.user, upload)
        path = Path(resource.stored_path)
        asset_path = Path(self.temp_dir.name) / "lesson-preview.png"
        Image.new("RGB", (4, 4), color=(255, 255, 255)).save(asset_path)
        KnowledgeAsset.objects.create(
            resource=resource,
            kind="rendered_page",
            storage_path=str(asset_path),
            mime_type="image/png",
            width=4,
            height=4,
        )
        resource_id = str(resource.id)

        delete_resource(resource)

        self.assertFalse(path.exists())
        self.assertFalse(asset_path.exists())
        self.assertFalse(KnowledgeResource.objects.exists())
        remove_visual_resource_entries.assert_called_once_with(resource_id, path)

    @patch("backend.domains.knowledge_base.multimodal_parser._understand_image", return_value=VisionUnderstanding(warnings=["OCR engine is not installed"]))
    def test_image_upload_creates_page_asset_block_and_warning_without_ocr(self, _understand_image):
        image_path = Path(self.temp_dir.name) / "question.png"
        Image.new("RGB", (6, 5), color=(0, 255, 0)).save(image_path)
        upload = SimpleUploadedFile("question.png", image_path.read_bytes(), content_type="image/png")

        with override_settings(KNOWLEDGE_UPLOAD_ROOT=self.temp_dir.name):
            resource = save_and_index_upload(self.user, upload)

        self.assertEqual(resource.status, ResourceStatus.INDEXED_WITH_WARNINGS)
        self.assertEqual(resource.pages.count(), 1)
        self.assertEqual(resource.assets.count(), 1)
        self.assertEqual(resource.assets.first().width, 6)
        self.assertEqual(resource.assets.first().height, 5)
        self.assertEqual(resource.blocks.filter(modality="image").count(), 1)
        self.assertEqual(resource.blocks.first().index_entries.filter(index_type="visual_vector").count(), 1)
        self.assertIn("OCR", serialize_resource(resource)["multimodal"]["warnings"][0])

    @patch("backend.domains.knowledge_base.multimodal_parser._understand_image", return_value=VisionUnderstanding(text="Root locus OCR text"))
    def test_image_upload_uses_ocr_text_as_searchable_block(self, _understand_image):
        image_path = Path(self.temp_dir.name) / "question-ocr.png"
        Image.new("RGB", (8, 7), color=(255, 255, 255)).save(image_path)
        upload = SimpleUploadedFile("question-ocr.png", image_path.read_bytes(), content_type="image/png")

        with override_settings(KNOWLEDGE_UPLOAD_ROOT=self.temp_dir.name):
            resource = save_and_index_upload(self.user, upload)

        page = resource.pages.get()
        self.assertEqual(page.ocr_text, "Root locus OCR text")
        text_block = resource.blocks.get(modality="text", text="Root locus OCR text")
        self.assertEqual(text_block.index_entries.filter(index_type="text_vector", status="indexed").count(), 1)
        self.assertEqual(text_block.index_entries.filter(index_type="lexical", status="indexed").count(), 1)
        self.assertEqual(serialize_resource(resource)["multimodal"]["ocrChars"], len("Root locus OCR text"))

    def test_ocr_image_uses_openai_compatible_vision_provider_when_configured(self):
        from backend.domains.knowledge_base.multimodal_parser import _ocr_image

        image_path = Path(self.temp_dir.name) / "cloud-ocr.png"
        Image.new("RGB", (8, 7), color=(255, 255, 255)).save(image_path)
        response = MagicMock()
        response.__enter__.return_value.read.return_value = json.dumps({
            "choices": [{"message": {"content": "Cloud OCR text"}}],
        }).encode("utf-8")
        response.__enter__.return_value.status = 200

        with patch.dict(os.environ, {
            "A3_OCR_PROVIDER": "openai_vision",
            "A3_OCR_API_KEY": "test-key",
            "A3_OCR_BASE_URL": "https://vision.example.test/v1",
            "A3_OCR_MODEL": "vision-ocr-test",
        }), patch("urllib.request.urlopen", return_value=response) as urlopen:
            text, warnings = _ocr_image(image_path)

        self.assertEqual(text, "Cloud OCR text")
        self.assertEqual(warnings, [])
        request = urlopen.call_args.args[0]
        self.assertEqual(request.full_url, "https://vision.example.test/v1/chat/completions")
        self.assertEqual(request.headers["Authorization"], "Bearer test-key")
        payload = json.loads(request.data.decode("utf-8"))
        self.assertEqual(payload["model"], "vision-ocr-test")
        self.assertEqual(payload["messages"][0]["content"][1]["type"], "image_url")

    def test_qwen_ocr_provider_uses_dashscope_document_ocr_model(self):
        image_path = Path(self.temp_dir.name) / "qwen-ocr.png"
        Image.new("RGB", (8, 7), color=(255, 255, 255)).save(image_path)
        response = MagicMock()
        response.__enter__.return_value.read.return_value = json.dumps({
            "choices": [{"message": {"content": "进程调度\n| 算法 | 特点 |\n|---|---|\n| RR | 时间片 |"}}],
        }).encode("utf-8")

        with patch.dict(os.environ, {
            "A3_OCR_PROVIDER": "qwen_ocr",
            "A3_LLM_API_KEY": "dashscope-key",
            "A3_OCR_API_KEY": "${A3_LLM_API_KEY}",
            "A3_OCR_BASE_URL": "https://dashscope.aliyuncs.com/compatible-mode/v1",
            "A3_OCR_MODEL": "qwen-vl-ocr-latest",
        }), patch("urllib.request.urlopen", return_value=response) as urlopen:
            understanding = QwenOcrDocumentParserProvider().understand_image(image_path)

        self.assertIn("进程调度", understanding.to_search_text())
        self.assertEqual(understanding.raw["provider"], "qwen_ocr")
        self.assertEqual(understanding.raw["model"], "qwen-vl-ocr-latest")
        request = urlopen.call_args.args[0]
        self.assertEqual(request.full_url, "https://dashscope.aliyuncs.com/compatible-mode/v1/chat/completions")
        self.assertEqual(request.headers["Authorization"], "Bearer dashscope-key")
        payload = json.loads(request.data.decode("utf-8"))
        self.assertEqual(payload["model"], "qwen-vl-ocr-latest")
        self.assertIn("use Markdown tables for tables", payload["messages"][1]["content"][0]["text"])

    def test_ocr_image_falls_back_to_local_when_vision_provider_fails(self):
        from backend.domains.knowledge_base.multimodal_parser import _ocr_image

        image_path = Path(self.temp_dir.name) / "fallback-ocr.png"
        Image.new("RGB", (8, 7), color=(255, 255, 255)).save(image_path)

        with patch.dict(os.environ, {
            "A3_OCR_PROVIDER": "openai_vision",
            "A3_OCR_API_KEY": "test-key",
            "A3_OCR_BASE_URL": "https://vision.example.test/v1",
            "A3_OCR_MODEL": "vision-ocr-test",
        }), patch("urllib.request.urlopen", side_effect=RuntimeError("provider down")), patch(
            "backend.domains.knowledge_base.multimodal_parser._ocr_image_local",
            return_value=("Local OCR text", []),
        ):
            text, warnings = _ocr_image(image_path)

        self.assertEqual(text, "Local OCR text")
        self.assertIn("OCR API failed", warnings[0])

    def test_vision_provider_json_response_creates_structured_understanding(self):
        from backend.domains.knowledge_base.multimodal_parser import _understand_image

        image_path = Path(self.temp_dir.name) / "structured-vlm.png"
        Image.new("RGB", (8, 7), color=(255, 255, 255)).save(image_path)
        vision_payload = {
            "text": "Root locus stability",
            "formulas": ["zeta = -sigma / sqrt(sigma^2 + omega_d^2)"],
            "tables": [{"caption": "Parameter table", "markdown": "| K | Overshoot |"}],
            "figures": [{"caption": "Root locus", "description": "Blue curve shows closed-loop pole trajectory"}],
            "keywords": ["root locus", "damping ratio"],
        }
        response = MagicMock()
        response.__enter__.return_value.read.return_value = json.dumps({
            "choices": [{"message": {"content": json.dumps(vision_payload)}}],
        }).encode("utf-8")

        with patch.dict(os.environ, {
            "A3_OCR_PROVIDER": "openai_vision",
            "A3_OCR_API_KEY": "test-key",
            "A3_OCR_BASE_URL": "https://vision.example.test/v1",
            "A3_OCR_MODEL": "vision-doc-test",
        }), patch("urllib.request.urlopen", return_value=response):
            understanding = _understand_image(image_path)

        self.assertEqual(understanding.text, "Root locus stability")
        self.assertEqual(understanding.formulas, ["zeta = -sigma / sqrt(sigma^2 + omega_d^2)"])
        self.assertEqual(understanding.tables[0]["markdown"], "| K | Overshoot |")
        self.assertEqual(understanding.figures[0]["description"], "Blue curve shows closed-loop pole trajectory")
        self.assertEqual(understanding.keywords, ["root locus", "damping ratio"])
        self.assertIn("Root locus stability", understanding.to_search_text())
        self.assertIn("| K | Overshoot |", understanding.to_search_text())

    def test_vision_provider_plain_text_response_extracts_formula_and_table_fallbacks(self):
        from backend.domains.knowledge_base.multimodal_parser import _understand_image

        image_path = Path(self.temp_dir.name) / "plain-vlm.png"
        Image.new("RGB", (8, 7), color=(255, 255, 255)).save(image_path)
        response = MagicMock()
        response.__enter__.return_value.read.return_value = json.dumps({
            "choices": [
                {
                    "message": {
                        "content": (
                            "Root locus damping ratio\n"
                            "Formula: zeta = -sigma / sqrt(sigma^2 + omega_d^2)\n"
                            "Table: K | Overshoot\n"
                            "1 | 8%"
                        )
                    }
                }
            ],
        }).encode("utf-8")

        with patch.dict(os.environ, {
            "A3_OCR_PROVIDER": "openai_vision",
            "A3_OCR_API_KEY": "test-key",
            "A3_OCR_BASE_URL": "https://vision.example.test/v1",
            "A3_OCR_MODEL": "vision-doc-test",
        }), patch("urllib.request.urlopen", return_value=response):
            understanding = _understand_image(image_path)

        self.assertIn("Root locus damping ratio", understanding.text)
        self.assertEqual(understanding.formulas, ["zeta = -sigma / sqrt(sigma^2 + omega_d^2)"])
        self.assertEqual(understanding.tables[0]["markdown"], "K | Overshoot\n1 | 8%")

    def test_vision_provider_json_text_field_extracts_missing_formula_and_table_fallbacks(self):
        from backend.domains.knowledge_base.multimodal_parser import _understand_image

        image_path = Path(self.temp_dir.name) / "json-text-only-vlm.png"
        Image.new("RGB", (8, 7), color=(255, 255, 255)).save(image_path)
        response = MagicMock()
        response.__enter__.return_value.read.return_value = json.dumps({
            "choices": [
                {
                    "message": {
                        "content": json.dumps({
                            "text": (
                                "Root locus damping ratio\n"
                                "Formula: zeta = -sigma / sqrt(sigma^2 + omega_d^2)\n"
                                "Table: K | Overshoot\n"
                                "1 | 8%"
                            ),
                            "formulas": [],
                            "tables": [],
                            "figures": [],
                            "keywords": [],
                        })
                    }
                }
            ],
        }).encode("utf-8")

        with patch.dict(os.environ, {
            "A3_OCR_PROVIDER": "openai_vision",
            "A3_OCR_API_KEY": "test-key",
            "A3_OCR_BASE_URL": "https://vision.example.test/v1",
            "A3_OCR_MODEL": "vision-doc-test",
        }), patch("urllib.request.urlopen", return_value=response):
            understanding = _understand_image(image_path)

        self.assertEqual(understanding.formulas, ["zeta = -sigma / sqrt(sigma^2 + omega_d^2)"])
        self.assertEqual(understanding.tables[0]["markdown"], "K | Overshoot\n1 | 8%")

    def test_vision_provider_repairs_truncated_json_object_response(self):
        from backend.domains.knowledge_base.multimodal_parser import _vision_understanding_from_text

        content = (
            '{\n'
            '  "text": "Beam section",\n'
            '  "formulas": [],\n'
            '  "tables": [],\n'
            '  "figures": [{"caption": "section", "description": "cross section diagram"}],\n'
            '  "keywords": ["beam", "section"]\n'
        )

        understanding = _vision_understanding_from_text(content, provider="qwen_vl")

        self.assertEqual(understanding.text, "Beam section")
        self.assertEqual(understanding.figures[0]["description"], "cross section diagram")
        self.assertEqual(understanding.keywords, ["beam", "section"])

    @patch("backend.domains.knowledge_base.multimodal_parser._understand_image")
    def test_image_upload_creates_structured_vlm_blocks_for_text_table_formula_and_figure(self, understand_image):
        from backend.domains.knowledge_base.multimodal_parser import VisionUnderstanding

        understand_image.return_value = VisionUnderstanding(
            text="Root locus stability",
            formulas=["zeta = -sigma / sqrt(sigma^2 + omega_d^2)"],
            tables=[{"caption": "Parameter table", "markdown": "| K | Overshoot |"}],
            figures=[{"caption": "Root locus", "description": "Blue curve shows closed-loop pole trajectory"}],
            keywords=["root locus", "damping ratio"],
        )
        image_path = Path(self.temp_dir.name) / "structured-upload.png"
        Image.new("RGB", (8, 7), color=(255, 255, 255)).save(image_path)
        upload = SimpleUploadedFile("structured-upload.png", image_path.read_bytes(), content_type="image/png")

        with override_settings(KNOWLEDGE_UPLOAD_ROOT=self.temp_dir.name):
            resource = save_and_index_upload(self.user, upload)

        self.assertEqual(resource.blocks.filter(modality="text", text="Root locus stability").count(), 1)
        self.assertEqual(resource.blocks.filter(modality="formula", text__contains="zeta =").count(), 1)
        self.assertEqual(resource.blocks.filter(modality="table", text="| K | Overshoot |").count(), 1)
        image_block = resource.blocks.get(modality="image")
        self.assertIn("Blue curve shows closed-loop pole trajectory", image_block.text)
        self.assertEqual(image_block.metadata["keywords"], ["root locus", "damping ratio"])
        self.assertEqual(resource.blocks.get(modality="formula").index_entries.filter(index_type="text_vector", status="indexed").count(), 1)
        self.assertEqual(resource.blocks.get(modality="table").index_entries.filter(index_type="lexical", status="indexed").count(), 1)
        self.assertEqual(serialize_resource(resource)["multimodal"]["formulaBlocks"], 1)
        self.assertEqual(serialize_resource(resource)["multimodal"]["tableBlocks"], 1)

    @patch("backend.domains.knowledge_base.multimodal_parser._understand_image", return_value=VisionUnderstanding(text="Rendered page OCR"))
    @patch("backend.domains.knowledge_base.index_manager.append_source_file", return_value=(1, 2, []))
    def test_pdf_rendered_page_ocr_creates_searchable_text_block(self, _append_source_file, _understand_image):
        upload = SimpleUploadedFile("lesson-ocr.pdf", MINIMAL_TEXT_PDF, content_type="application/pdf")

        with override_settings(KNOWLEDGE_UPLOAD_ROOT=self.temp_dir.name):
            resource = save_and_index_upload(self.user, upload)

        page = resource.pages.get()
        self.assertEqual(page.ocr_text, "Rendered page OCR")
        self.assertEqual(resource.blocks.filter(modality="text", text="Rendered page OCR").count(), 1)
        self.assertEqual(serialize_resource(resource)["multimodal"]["ocrChars"], len("Rendered page OCR"))

    @patch("backend.domains.knowledge_base.index_manager.append_source_file", return_value=(1, 2, []))
    def test_pptx_text_slide_creates_slide_page_and_text_block(self, _append_source_file):
        from pptx import Presentation

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = "根轨迹规则"
        pptx_path = Path(self.temp_dir.name) / "root-locus.pptx"
        presentation.save(pptx_path)
        upload = SimpleUploadedFile(
            "root-locus.pptx",
            pptx_path.read_bytes(),
            content_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )

        with override_settings(KNOWLEDGE_UPLOAD_ROOT=self.temp_dir.name):
            resource = save_and_index_upload(self.user, upload)

        self.assertEqual(resource.status, ResourceStatus.INDEXED)
        self.assertEqual(resource.pages.count(), 1)
        self.assertEqual(resource.pages.first().slide_number, 1)
        self.assertGreaterEqual(resource.blocks.filter(modality="text", text__contains="根轨迹").count(), 1)
        self.assertGreaterEqual(serialize_resource(resource)["multimodal"]["qualityScore"], 0.7)

    @patch("backend.domains.knowledge_base.multimodal_parser._understand_image", return_value=VisionUnderstanding(warnings=["OCR engine is not installed"]))
    @patch("backend.domains.knowledge_base.index_manager.append_source_file", side_effect=ValueError("no extractable text"))
    def test_pptx_without_extractable_text_is_indexed_with_warnings_not_plain_indexed(self, _append_source_file, _understand_image):
        from pptx import Presentation

        presentation = Presentation()
        presentation.slides.add_slide(presentation.slide_layouts[6])
        pptx_path = Path(self.temp_dir.name) / "image-only.pptx"
        presentation.save(pptx_path)
        upload = SimpleUploadedFile(
            "image-only.pptx",
            pptx_path.read_bytes(),
            content_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )

        with override_settings(KNOWLEDGE_UPLOAD_ROOT=self.temp_dir.name):
            resource = save_and_index_upload(self.user, upload)

        self.assertEqual(resource.status, ResourceStatus.INDEXED_WITH_WARNINGS)
        self.assertEqual(resource.pages.count(), 1)
        self.assertEqual(resource.blocks.filter(modality="slide_image").count(), 1)
        self.assertIn("OCR", " ".join(serialize_resource(resource)["multimodal"]["warnings"]))

    @patch("backend.domains.knowledge_base.multimodal_parser._understand_image", return_value=VisionUnderstanding(text="Root locus slide OCR"))
    @patch("backend.domains.knowledge_base.multimodal_parser._render_pptx_slides")
    @patch("backend.domains.knowledge_base.index_manager.append_source_file", side_effect=ValueError("no extractable text"))
    def test_image_only_pptx_rendered_slide_ocr_creates_searchable_text_block(self, _append_source_file, render_slides, _understand_image):
        from pptx import Presentation
        from backend.domains.knowledge_base.multimodal_parser import ParsedAsset

        presentation = Presentation()
        presentation.slides.add_slide(presentation.slide_layouts[6])
        pptx_path = Path(self.temp_dir.name) / "image-only-ocr.pptx"
        presentation.save(pptx_path)
        rendered_path = Path(self.temp_dir.name) / "slide-1.png"
        Image.new("RGB", (10, 8), color=(255, 255, 255)).save(rendered_path)
        render_slides.return_value = (
            {
                1: ParsedAsset(
                    kind="rendered_slide",
                    storage_path=str(rendered_path),
                    mime_type="image/png",
                    width=10,
                    height=8,
                )
            },
            [],
        )
        upload = SimpleUploadedFile(
            "image-only-ocr.pptx",
            pptx_path.read_bytes(),
            content_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )

        with override_settings(KNOWLEDGE_UPLOAD_ROOT=self.temp_dir.name):
            resource = save_and_index_upload(self.user, upload)

        page = resource.pages.get()
        self.assertEqual(resource.status, ResourceStatus.INDEXED_WITH_WARNINGS)
        self.assertEqual(page.ocr_text, "Root locus slide OCR")
        self.assertEqual(resource.assets.filter(kind="rendered_slide", width=10, height=8).count(), 1)
        self.assertEqual(resource.blocks.filter(modality="slide_image", text="Root locus slide OCR").count(), 1)
        text_block = resource.blocks.get(modality="text", text="Root locus slide OCR")
        self.assertEqual(text_block.index_entries.filter(index_type="text_vector", status="indexed").count(), 1)
        self.assertEqual(text_block.index_entries.filter(index_type="lexical", status="indexed").count(), 1)
        self.assertEqual(serialize_resource(resource)["multimodal"]["ocrChars"], len("Root locus slide OCR"))

    @patch("backend.apps.knowledge_base.tasks.close_old_connections")
    @patch("backend.domains.knowledge_base.index_manager.append_source_file", return_value=(2, 6, []))
    def test_incremental_job_updates_resource_and_progress(self, _append_source_file, _close_connections):
        path = Path(self.temp_dir.name) / "lesson.pdf"
        path.write_bytes(b"fake-pdf")
        job = KnowledgeIndexJob.objects.create(
            owner=self.user,
            job_type=IndexJobType.UPLOAD,
            total_files=1,
        )
        resource = KnowledgeResource.objects.create(
            owner=self.user,
            index_job=job,
            original_name="lesson.pdf",
            stored_path=str(path),
            file_type=".pdf",
            file_size=path.stat().st_size,
        )

        run_index_job.run(str(job.id))

        job.refresh_from_db()
        resource.refresh_from_db()
        self.assertEqual(job.status, IndexJobStatus.SUCCEEDED)
        self.assertEqual(job.progress, 100)
        self.assertEqual(job.processed_files, 1)
        self.assertEqual(resource.status, ResourceStatus.INDEXED)
        self.assertEqual(resource.document_count, 2)
        self.assertEqual(resource.chunk_count, 6)
        payload = serialize_index_job(job)
        self.assertEqual(payload["resources"][0]["name"], "lesson.pdf")
        self.assertEqual(payload["resources"][0]["status"], ResourceStatus.INDEXED)
        self.assertEqual(payload["summary"], {"indexed": 1, "failed": 0})

    @override_settings(DEBUG=True)
    @patch("backend.apps.knowledge_base.tasks.close_old_connections")
    @patch("backend.domains.knowledge_base.index_manager.append_source_file", return_value=(1, 3, []))
    @patch("backend.apps.knowledge_base.tasks.run_index_job.delay", side_effect=ModuleNotFoundError("pywin32_system32"))
    def test_upload_job_falls_back_to_sync_indexing_when_queue_is_unavailable(
        self,
        _delay,
        _append_source_file,
        _close_connections,
    ):
        path = Path(self.temp_dir.name) / "lesson.xlsx"
        path.write_bytes(b"fake-xlsx")
        resource = KnowledgeResource.objects.create(
            owner=self.user,
            original_name="lesson.xlsx",
            stored_path=str(path),
            file_type=".xlsx",
            file_size=path.stat().st_size,
        )

        job = create_upload_index_job(self.user, [resource])

        job.refresh_from_db()
        resource.refresh_from_db()
        self.assertEqual(job.status, IndexJobStatus.SUCCEEDED)
        self.assertEqual(job.progress, 100)
        self.assertEqual(resource.status, ResourceStatus.INDEXED)
        self.assertEqual(resource.document_count, 1)
        self.assertEqual(resource.chunk_count, 3)

    @patch("backend.apps.knowledge_base.tasks.run_index_job.delay")
    def test_parse_job_is_queued_and_attached_to_resources(self, delay):
        delay.return_value.id = "parse-task-id"
        path = Path(self.temp_dir.name) / "lesson.pdf"
        path.write_bytes(b"fake-pdf")
        resource = KnowledgeResource.objects.create(
            owner=self.user,
            original_name="lesson.pdf",
            stored_path=str(path),
            file_type=".pdf",
            file_size=path.stat().st_size,
        )

        with patch.dict(os.environ, {"A3_KB_SYNC_PARSE": ""}):
            job = create_parse_job(self.user, [resource])

        resource.refresh_from_db()
        self.assertEqual(job.job_type, IndexJobType.PARSE)
        self.assertEqual(job.status, IndexJobStatus.QUEUED)
        self.assertEqual(job.celery_task_id, "parse-task-id")
        self.assertEqual(resource.index_job_id, job.id)
        payload = serialize_index_job(job)
        self.assertEqual(payload["type"], "parse")
        self.assertEqual(payload["summary"], {"parsed": 0, "indexed": 0, "failed": 0})

    @patch("backend.apps.knowledge_base.tasks.run_index_job.delay")
    def test_parse_job_can_defer_dispatch(self, delay):
        resource = KnowledgeResource.objects.create(
            owner=self.user,
            original_name="deferred-lesson.pdf",
            stored_path=str((Path(self.temp_dir.name) / "deferred-lesson.pdf").resolve()),
            file_type=".pdf",
            file_size=10,
        )

        job = create_parse_job(self.user, [resource], dispatch=False)

        delay.assert_not_called()
        resource.refresh_from_db()
        self.assertEqual(job.status, IndexJobStatus.QUEUED)
        self.assertEqual(job.celery_task_id, "")
        self.assertEqual(resource.index_job_id, job.id)
        self.assertEqual(resource.status, ResourceStatus.INDEXING)

    @patch("backend.apps.knowledge_base.tasks.close_old_connections")
    @patch("backend.apps.knowledge_base.services.reparse_resource")
    def test_parse_job_updates_progress_status_and_summary(self, reparse_resource_mock, _close_connections):
        path = Path(self.temp_dir.name) / "lesson.pdf"
        path.write_bytes(b"fake-pdf")
        job = KnowledgeIndexJob.objects.create(
            owner=self.user,
            job_type=IndexJobType.PARSE,
            total_files=1,
        )
        resource = KnowledgeResource.objects.create(
            owner=self.user,
            index_job=job,
            original_name="lesson.pdf",
            stored_path=str(path),
            file_type=".pdf",
            file_size=path.stat().st_size,
        )

        def fake_reparse(target, progress_callback=None):
            self.assertIsNotNone(progress_callback)
            progress_callback({
                "phase": "native_parse",
                "phaseLabel": "正在解析 PDF",
                "phaseProgress": 50,
                "event": "已解析第 1 页",
                "page": {"current": 1, "total": 1},
            })
            KnowledgeParseRun.objects.create(
                resource=target,
                status="completed",
                parser="native_pdf",
                pages=1,
                text_blocks=1,
                quality_score=0.91,
                result={"documentParse": {"schemaVersion": "a3.document-parse.v2", "citationReadiness": True}},
            )
            KnowledgeBlock.objects.create(resource=target, modality="text", text="进程调度", confidence=0.9)
            return MultimodalParseResult(
                parser="native_pdf",
                status="completed",
                pages=[ParsedPage(page_number=1, blocks=[ParsedBlock(modality="text", text="进程调度")])],
            )

        reparse_resource_mock.side_effect = fake_reparse

        run_index_job.run(str(job.id))

        job.refresh_from_db()
        resource.refresh_from_db()
        self.assertEqual(job.status, IndexJobStatus.SUCCEEDED)
        self.assertEqual(job.progress, 100)
        self.assertEqual(job.processed_files, 1)
        self.assertEqual(resource.status, ResourceStatus.INDEXED)
        self.assertEqual(resource.document_count, 1)
        self.assertEqual(resource.chunk_count, 1)
        self.assertTrue(any(event["phase"] == "native_parse" for event in job.progress_detail["recentEvents"]))
        self.assertEqual(serialize_index_job(job)["summary"], {"parsed": 1, "indexed": 1, "failed": 0})

    @patch("backend.apps.knowledge_base.tasks.close_old_connections")
    @patch("backend.apps.knowledge_base.wiki_builder.safe_rebuild_resource_wiki")
    @patch("backend.domains.knowledge_base.index_manager.append_source_file", return_value=(1, 2, []))
    def test_incremental_job_persists_multimodal_progress_detail(self, _append_source_file, _wiki_builder, _close_connections):
        path = Path(self.temp_dir.name) / "lesson.pdf"
        path.write_bytes(b"fake-pdf")
        job = KnowledgeIndexJob.objects.create(
            owner=self.user,
            job_type=IndexJobType.UPLOAD,
            total_files=1,
        )
        resource = KnowledgeResource.objects.create(
            owner=self.user,
            index_job=job,
            original_name="lesson.pdf",
            stored_path=str(path),
            file_type=".pdf",
            file_size=path.stat().st_size,
        )

        def fake_parse(_source_path, progress_callback=None):
            self.assertIsNotNone(progress_callback)
            progress_callback({
                "phase": "ocr",
                "phaseLabel": "正在 OCR 第 1/2 页",
                "page": {"current": 1, "total": 2},
                "metrics": {"pagesParsed": 1, "warnings": 0},
                "event": "第 1 页 OCR 完成",
                "phaseProgress": 40,
            })
            progress_callback({
                "phase": "chunk",
                "phaseLabel": "正在写入知识块 2/2",
                "block": {"current": 2, "total": 2},
                "metrics": {"blocksCreated": 2, "pagesParsed": 2},
                "event": "已生成 2 个知识块",
                "phaseProgress": 80,
            })
            return MultimodalParseResult(
                parser="native_pdf",
                status="completed",
                pages=[
                    ParsedPage(page_number=1, blocks=[ParsedBlock(modality="text", text="死锁必要条件")]),
                    ParsedPage(page_number=2, blocks=[ParsedBlock(modality="text", text="银行家算法")]),
                ],
                result={"pages": 2, "textPages": 2, "ocrPages": 1},
            )

        with patch("backend.apps.knowledge_base.services.parse_resource", side_effect=fake_parse):
            run_index_job.run(str(job.id))

        job.refresh_from_db()
        payload = serialize_index_job(job)
        self.assertEqual(payload["progressDetail"]["phase"], "done")
        self.assertEqual(payload["progressDetail"]["resourceName"], "lesson.pdf")
        self.assertEqual(payload["progressDetail"]["page"], {"current": 2, "total": 2})
        self.assertEqual(payload["progressDetail"]["block"], {"current": 2, "total": 2})
        self.assertEqual(payload["progressDetail"]["metrics"]["blocksCreated"], 2)
        self.assertTrue(any("2" in event["message"] for event in payload["progressDetail"]["recentEvents"]))
        self.assertEqual(payload["progressDetail"]["recentEvents"][-1]["phase"], "done")

    @patch("backend.apps.knowledge_base.tasks.run_index_job.delay")
    def test_wiki_rebuild_job_is_queued_and_attached_to_indexed_resources(self, delay):
        delay.return_value.id = "wiki-task-id"
        resource = KnowledgeResource.objects.create(
            owner=self.user,
            original_name="wiki-job.pdf",
            stored_path=str((Path(self.temp_dir.name) / "wiki-job.pdf").resolve()),
            file_type=".pdf",
            file_size=10,
            status=ResourceStatus.INDEXED,
        )

        job = create_wiki_rebuild_job(self.user, [resource])

        resource.refresh_from_db()
        self.assertEqual(job.job_type, IndexJobType.WIKI)
        self.assertEqual(job.status, IndexJobStatus.QUEUED)
        self.assertEqual(job.celery_task_id, "wiki-task-id")
        self.assertEqual(resource.index_job_id, job.id)
        payload = serialize_index_job(job)
        self.assertEqual(payload["type"], "wiki")
        self.assertEqual(payload["summary"], {"rebuilt": 0, "indexed": 1, "failed": 0})

    @patch("backend.apps.knowledge_base.tasks.close_old_connections")
    def test_wiki_rebuild_job_updates_progress_status_and_summary(self, _close_old_connections):
        job = KnowledgeIndexJob.objects.create(owner=self.user, job_type=IndexJobType.WIKI, total_files=1)
        path = Path(self.temp_dir.name) / "wiki-job-run.pdf"
        path.write_text("Root Locus is a graphical method for analyzing closed-loop pole movement.", encoding="utf-8")
        resource = KnowledgeResource.objects.create(
            owner=self.user,
            index_job=job,
            original_name="wiki-job-run.pdf",
            stored_path=str(path.resolve()),
            file_type=".pdf",
            file_size=path.stat().st_size,
            status=ResourceStatus.INDEXED,
        )
        page = KnowledgePage.objects.create(resource=resource, page_number=1)
        KnowledgeBlock.objects.create(resource=resource, page=page, modality="text", text=path.read_text(encoding="utf-8"), confidence=0.9)

        run_index_job.run(str(job.id))

        job.refresh_from_db()
        self.assertEqual(job.status, IndexJobStatus.SUCCEEDED)
        self.assertEqual(job.progress, 100)
        self.assertEqual(job.processed_files, 1)
        self.assertTrue(KnowledgeConcept.objects.filter(resource=resource, normalized_label="root locus").exists())
        self.assertEqual(serialize_index_job(job)["summary"], {"rebuilt": 1, "indexed": 1, "failed": 0})

    def test_evaluate_document_parsing_reports_quality_metrics(self):
        path = Path(self.temp_dir.name) / "os.pdf"
        path.write_bytes(b"fake-pdf")
        resource = KnowledgeResource.objects.create(
            owner=self.user,
            original_name="第1章 操作系统概述.pdf",
            stored_path=str(path),
            file_type=".pdf",
            file_size=path.stat().st_size,
            course_key="OS-001",
            status=ResourceStatus.INDEXED,
        )
        KnowledgeParseRun.objects.create(
            resource=resource,
            status="completed",
            parser="native_pdf",
            pages=1,
            text_blocks=1,
            table_blocks=1,
            formula_blocks=1,
            quality_score=0.88,
            result={
                "documentParse": {
                    "schemaVersion": "a3.document-parse.v2",
                    "textCoverage": 1,
                    "garbledRatio": 0,
                    "citationReadiness": True,
                    "providerChain": ["native_pdf"],
                }
            },
        )
        KnowledgeBlock.objects.create(resource=resource, modality="table", text="| 算法 | 特点 |\n|---|---|")
        KnowledgeBlock.objects.create(resource=resource, modality="formula", text="$T=(n-1)q$")

        output = StringIO()
        call_command("evaluate_document_parsing", "--course-key", "OS-001", "--json", stdout=output)
        payload = json.loads(output.getvalue())

        self.assertEqual(payload["schemaVersion"], "a3.document-parse-eval.v1")
        self.assertEqual(payload["summary"]["resources"], 1)
        self.assertEqual(payload["summary"]["parsedResources"], 1)
        self.assertEqual(payload["summary"]["averageGarbledRatio"], 0)
        self.assertEqual(payload["summary"]["citationReadyRate"], 1.0)
        self.assertEqual(payload["summary"]["tableBlocks"], 1)
        self.assertEqual(payload["summary"]["formulaBlocks"], 1)

    @patch("backend.apps.knowledge_base.tasks.PersonalMaterialAgent")
    @patch("backend.apps.knowledge_base.tasks.close_old_connections")
    @patch("backend.domains.knowledge_base.index_manager.append_source_file", return_value=(1, 3, []))
    def test_personal_upload_job_runs_personal_material_agent(self, _append_source_file, _close_connections, agent_cls):
        path = Path(self.temp_dir.name) / "student-note.pdf"
        path.write_bytes(b"fake-pdf")
        job = KnowledgeIndexJob.objects.create(
            owner=self.user,
            job_type=IndexJobType.UPLOAD,
            total_files=1,
        )
        resource = KnowledgeResource.objects.create(
            owner=self.user,
            index_job=job,
            original_name="student-note.pdf",
            stored_path=str(path),
            file_type=".pdf",
            file_size=path.stat().st_size,
            resource_scope=ResourceScope.PERSONAL,
        )

        run_index_job.run(str(job.id))

        agent_cls.return_value.analyze_and_writeback.assert_called_once_with(resource)

    @patch("backend.apps.knowledge_base.tasks.run_index_job.delay")
    def test_failed_job_can_be_requeued(self, delay):
        delay.return_value.id = "retry-task-id"
        job = KnowledgeIndexJob.objects.create(
            owner=self.user,
            job_type=IndexJobType.UPLOAD,
            status=IndexJobStatus.FAILED,
            progress=45,
            error_message="temporary failure",
        )

        with patch.dict(os.environ, {"A3_KB_SYNC_INDEX_ON_UPLOAD": ""}):
            retry_index_job(job)

        job.refresh_from_db()
        self.assertEqual(job.status, IndexJobStatus.QUEUED)
        self.assertEqual(job.progress, 0)
        self.assertEqual(job.error_message, "")
        self.assertEqual(job.celery_task_id, "retry-task-id")

    @patch("backend.apps.knowledge_base.tasks.run_index_job.delay")
    def test_failed_job_can_be_reset_without_dispatch(self, delay):
        job = KnowledgeIndexJob.objects.create(
            owner=self.user,
            job_type=IndexJobType.UPLOAD,
            status=IndexJobStatus.FAILED,
            progress=45,
            error_message="temporary failure",
        )

        retry_index_job(job, dispatch=False)

        delay.assert_not_called()
        job.refresh_from_db()
        self.assertEqual(job.status, IndexJobStatus.QUEUED)
        self.assertEqual(job.progress, 0)
        self.assertEqual(job.error_message, "")
        self.assertEqual(job.celery_task_id, "")

    @override_settings(KNOWLEDGE_INDEX_JOB_STALE_MINUTES=15)
    @patch("backend.config.celery.app.control.inspect")
    def test_resources_api_keeps_stale_heartbeat_job_running_when_celery_task_is_active(self, inspect):
        inspect.return_value.active.return_value = {
            "celery@worker": [{"id": "active-index-task"}],
        }
        job = KnowledgeIndexJob.objects.create(
            owner=self.user,
            job_type=IndexJobType.PARSE,
            status=IndexJobStatus.RUNNING,
            progress=81,
            celery_task_id="active-index-task",
            total_files=1,
        )
        resource = KnowledgeResource.objects.create(
            owner=self.user,
            index_job=job,
            original_name="long-running-graph.pdf",
            stored_path=str(Path(self.temp_dir.name) / "long-running-graph.pdf"),
            file_type=".pdf",
            file_size=100,
            status=ResourceStatus.INDEXING,
        )
        stale_at = timezone.now() - timedelta(minutes=30)
        KnowledgeIndexJob.objects.filter(id=job.id).update(updated_at=stale_at)
        KnowledgeResource.objects.filter(id=resource.id).update(updated_at=stale_at)

        client = Client()
        client.force_login(self.user)
        response = client.get("/api/kb/resources/")

        self.assertEqual(response.status_code, 200)
        job.refresh_from_db()
        resource.refresh_from_db()
        self.assertEqual(job.status, IndexJobStatus.RUNNING)
        self.assertEqual(resource.status, ResourceStatus.INDEXING)

    @override_settings(KNOWLEDGE_INDEX_JOB_STALE_MINUTES=15)
    def test_resources_api_marks_stale_running_index_job_failed_after_restart(self):
        path = Path(self.temp_dir.name) / "stale-running.pdf"
        path.write_bytes(b"fake-pdf")
        job = KnowledgeIndexJob.objects.create(
            owner=self.user,
            job_type=IndexJobType.PARSE,
            status=IndexJobStatus.RUNNING,
            progress=81,
            total_files=1,
            current_file="stale-running.pdf",
            progress_detail={
                "schemaVersion": "a3.kb-progress.v1",
                "phase": "persist",
                "phaseLabel": "正在写入解析结果 72/72 页",
            },
        )
        resource = KnowledgeResource.objects.create(
            owner=self.user,
            index_job=job,
            original_name="stale-running.pdf",
            stored_path=str(path),
            file_type=".pdf",
            file_size=path.stat().st_size,
            visibility="owner",
            status=ResourceStatus.INDEXING,
        )
        stale_at = timezone.now() - timedelta(minutes=30)
        KnowledgeIndexJob.objects.filter(id=job.id).update(started_at=stale_at, updated_at=stale_at)
        KnowledgeResource.objects.filter(id=resource.id).update(updated_at=stale_at)

        client = Client()
        client.force_login(self.user)
        response = client.get("/api/kb/resources/")

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload["success"])
        resource_payload = next(item for item in payload["resources"] if item["id"] == str(resource.id))
        self.assertEqual(resource_payload["status"], ResourceStatus.FAILED)
        self.assertIn("服务重启", resource_payload["error"])
        job.refresh_from_db()
        self.assertEqual(job.status, IndexJobStatus.FAILED)
        self.assertEqual(job.progress, 81)
        self.assertIn("服务重启", job.error_message)
        self.assertEqual(job.progress_detail["phase"], "interrupted")
        self.assertTrue(job.result["interrupted"])
        self.assertEqual(job.result["failed"][0]["resourceId"], str(resource.id))

    def test_delete_active_resource_cancels_stuck_index_job(self):
        path = Path(self.temp_dir.name) / "delete-active.pdf"
        path.write_bytes(b"fake-pdf")
        job = KnowledgeIndexJob.objects.create(
            owner=self.user,
            job_type=IndexJobType.PARSE,
            status=IndexJobStatus.RUNNING,
            progress=42,
            total_files=1,
            current_file="delete-active.pdf",
        )
        resource = KnowledgeResource.objects.create(
            owner=self.user,
            index_job=job,
            original_name="delete-active.pdf",
            stored_path=str(path),
            file_type=".pdf",
            file_size=path.stat().st_size,
            status=ResourceStatus.INDEXING,
        )

        client = Client()
        client.force_login(self.user)
        response = client.delete(f"/api/kb/resources/{resource.id}/")

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()["success"])
        self.assertFalse(path.exists())
        self.assertFalse(KnowledgeResource.objects.filter(id=resource.id).exists())
        job.refresh_from_db()
        self.assertEqual(job.status, IndexJobStatus.FAILED)
        self.assertIn("用户删除资料", job.error_message)
        self.assertEqual(job.progress_detail["phase"], "cancelled")
        self.assertTrue(job.result["cancelledByDelete"])


class KnowledgeBaseApiPermissionTests(TestCase):
    def setUp(self):
        self.student = User.objects.create_user(phone="13800139010", password="secure-pass-123")
        self.teacher = User.objects.create_user(
            phone="13800139011",
            password="secure-pass-123",
            role=UserRole.TEACHER,
            display_name="知识库教师",
        )
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.client = Client()

    def test_student_can_open_personal_knowledge_base_api(self):
        self.client.force_login(self.student)
        response = self.client.get("/api/kb/resources/")
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()["success"])
        self.assertEqual(response.json()["resources"], [])

    def test_teacher_can_open_knowledge_base_api(self):
        self.client.force_login(self.teacher)
        response = self.client.get("/api/kb/resources/")
        self.assertEqual(response.status_code, 200)

    @patch("backend.apps.knowledge_base.services.clear_text_index_storage")
    @patch("backend.apps.knowledge_base.tasks.run_index_job.delay")
    def test_teacher_can_reset_derived_knowledge_base_without_queueing_reparse(self, delay, clear_text_index_storage):
        source_path = Path(self.temp_dir.name) / "reset-source.pdf"
        source_path.write_bytes(b"source")
        pending_path = Path(self.temp_dir.name) / "pending-source.pdf"
        pending_path.write_bytes(b"pending")
        asset_path = Path(self.temp_dir.name) / "reset-asset.png"
        asset_path.write_bytes(b"asset")
        resource = KnowledgeResource.objects.create(
            owner=self.teacher,
            original_name="reset-source.pdf",
            stored_path=str(source_path),
            file_type=".pdf",
            file_size=source_path.stat().st_size,
            resource_scope=ResourceScope.TEACHER,
            status=ResourceStatus.INDEXED,
            document_count=2,
            chunk_count=4,
        )
        pending_resource = KnowledgeResource.objects.create(
            owner=self.teacher,
            original_name="pending-source.pdf",
            stored_path=str(pending_path),
            file_type=".pdf",
            file_size=pending_path.stat().st_size,
            resource_scope=ResourceScope.TEACHER,
            status=ResourceStatus.UPLOADED,
            document_count=0,
            chunk_count=0,
        )
        parse_run = KnowledgeParseRun.objects.create(resource=resource, status="completed", parser="native", pages=1)
        page = KnowledgePage.objects.create(resource=resource, parse_run=parse_run, page_number=1)
        asset = KnowledgeAsset.objects.create(resource=resource, page=page, kind="page", storage_path=str(asset_path))
        block = KnowledgeBlock.objects.create(resource=resource, page=page, asset=asset, modality="text", text="根轨迹", confidence=0.9)
        KnowledgeIndexEntry.objects.create(block=block, index_type="text_vector", index_backend="faiss", status="indexed")
        entity = KnowledgeGraphEntity.objects.create(
            owner=self.teacher,
            resource=resource,
            label="根轨迹",
            normalized_label="根轨迹",
            entity_type="concept",
            summary="根轨迹用于控制系统分析。",
        )
        KnowledgeGraphClaim.objects.create(
            owner=self.teacher,
            resource=resource,
            text="根轨迹用于分析闭环极点。",
            entity_ids=[str(entity.id)],
            verified=True,
        )
        concept = KnowledgeConcept.objects.create(
            owner=self.teacher,
            resource=resource,
            preferred_label="根轨迹",
            normalized_label="根轨迹",
        )
        KnowledgeDefinition.objects.create(
            owner=self.teacher,
            resource=resource,
            concept=concept,
            source_block=block,
            definition="根轨迹是控制系统分析方法。",
        )
        self.client.force_login(self.teacher)

        response = self.client.post(
            "/api/kb/reset/",
            data=json.dumps({"action": "reset_derived"}),
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload["success"])
        self.assertEqual(payload["action"], "reset_derived")
        self.assertEqual(payload["deleted"]["resources"], 0)
        self.assertEqual(payload["deleted"]["blocks"], 1)
        self.assertNotIn("job", payload)
        delay.assert_not_called()
        clear_text_index_storage.assert_called_once()
        resource.refresh_from_db()
        pending_resource.refresh_from_db()
        self.assertEqual(resource.status, ResourceStatus.UPLOADED)
        self.assertIsNone(resource.index_job_id)
        self.assertEqual(pending_resource.status, ResourceStatus.UPLOADED)
        self.assertIsNone(pending_resource.index_job_id)
        self.assertTrue(source_path.exists())
        self.assertTrue(pending_path.exists())
        self.assertFalse(asset_path.exists())
        self.assertFalse(resource.blocks.exists())
        self.assertFalse(resource.graph_entities.exists())
        self.assertFalse(resource.concepts.exists())

    @patch("backend.apps.knowledge_base.services.clear_text_index_storage")
    def test_teacher_can_clear_knowledge_base_after_confirmation(self, clear_text_index_storage):
        source_path = Path(self.temp_dir.name) / "clear-source.pdf"
        source_path.write_bytes(b"source")
        resource = KnowledgeResource.objects.create(
            owner=self.teacher,
            original_name="clear-source.pdf",
            stored_path=str(source_path),
            file_type=".pdf",
            file_size=source_path.stat().st_size,
            resource_scope=ResourceScope.TEACHER,
            status=ResourceStatus.INDEXED,
        )
        KnowledgeBlock.objects.create(resource=resource, modality="text", text="调度", confidence=0.9)
        self.client.force_login(self.teacher)

        rejected = self.client.post(
            "/api/kb/reset/",
            data=json.dumps({"action": "clear_all", "confirmText": "清空"}),
            content_type="application/json",
        )
        self.assertEqual(rejected.status_code, 400)
        self.assertTrue(KnowledgeResource.objects.filter(id=resource.id).exists())

        response = self.client.post(
            "/api/kb/reset/",
            data=json.dumps({"action": "clear_all", "confirmText": "清空知识库"}),
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload["success"])
        self.assertEqual(payload["action"], "clear_all")
        self.assertEqual(payload["deleted"]["resources"], 1)
        clear_text_index_storage.assert_called_once()
        self.assertFalse(KnowledgeResource.objects.filter(id=resource.id).exists())
        self.assertFalse(source_path.exists())

    @patch("backend.apps.knowledge_base.services.clear_text_index_storage")
    def test_admin_clear_knowledge_base_removes_student_and_teacher_resources(self, clear_text_index_storage):
        admin = User.objects.create_superuser(
            phone="13800139012",
            password="secure-pass-123",
            display_name="管理员",
        )
        student_path = Path(self.temp_dir.name) / "student-personal.pdf"
        teacher_path = Path(self.temp_dir.name) / "teacher-owned.pdf"
        student_path.write_bytes(b"student")
        teacher_path.write_bytes(b"teacher")
        student_resource = KnowledgeResource.objects.create(
            owner=self.student,
            original_name="student-personal.pdf",
            stored_path=str(student_path),
            file_type=".pdf",
            file_size=student_path.stat().st_size,
            resource_scope=ResourceScope.PERSONAL,
            status=ResourceStatus.INDEXED,
        )
        teacher_resource = KnowledgeResource.objects.create(
            owner=self.teacher,
            original_name="teacher-owned.pdf",
            stored_path=str(teacher_path),
            file_type=".pdf",
            file_size=teacher_path.stat().st_size,
            resource_scope=ResourceScope.TEACHER,
            status=ResourceStatus.INDEXED,
        )
        self.client.force_login(admin)

        response = self.client.post(
            "/api/kb/reset/",
            data=json.dumps({"action": "clear_all", "confirmText": "清空知识库"}),
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload["success"])
        self.assertEqual(payload["deleted"]["resources"], 2)
        clear_text_index_storage.assert_called_once()
        self.assertFalse(KnowledgeResource.objects.filter(id__in=[student_resource.id, teacher_resource.id]).exists())
        self.assertFalse(student_path.exists())
        self.assertFalse(teacher_path.exists())

    @patch("backend.apps.knowledge_base.tasks.close_old_connections")
    @patch("backend.domains.knowledge_base.index_manager.append_source_file", return_value=(1, 2, []))
    def test_student_upload_creates_personal_resource(self, _append_source_file, _close_connections):
        upload = SimpleUploadedFile("student-note.pdf", b"fake-pdf", content_type="application/pdf")
        self.client.force_login(self.student)

        with override_settings(KNOWLEDGE_UPLOAD_ROOT=self.temp_dir.name):
            response = self.client.post("/api/kb/resources/upload/", {"files": [upload]})

        self.assertEqual(response.status_code, 202)
        payload = response.json()
        self.assertTrue(payload["success"])
        resource = KnowledgeResource.objects.get(owner=self.student)
        self.assertEqual(resource.resource_scope, ResourceScope.PERSONAL)
        self.assertEqual(payload["resources"][0]["scope"], ResourceScope.PERSONAL)
        self.assertEqual(payload["resources"][0]["scopeLabel"], "我的资料")

    def test_teacher_resource_list_excludes_student_personal_resources(self):
        KnowledgeResource.objects.create(
            owner=self.student,
            original_name="student-private.pdf",
            stored_path=str((Path(self.temp_dir.name) / "student-private.pdf").resolve()),
            file_type=".pdf",
            file_size=10,
            resource_scope=ResourceScope.PERSONAL,
            status=ResourceStatus.INDEXED,
        )

        self.client.force_login(self.teacher)
        response = self.client.get("/api/kb/resources/")

        self.assertEqual(response.status_code, 200)
        names = {item["name"] for item in response.json()["resources"]}
        self.assertNotIn("student-private.pdf", names)

    def test_teacher_can_view_global_resources_owned_by_another_teacher(self):
        other_teacher = User.objects.create_user(phone="13800139012", password="secure-pass-123", role=UserRole.TEACHER)
        KnowledgeResource.objects.create(
            owner=other_teacher,
            original_name="os-global-courseware.pdf",
            stored_path=str((Path(self.temp_dir.name) / "os-global-courseware.pdf").resolve()),
            file_type=".pdf",
            file_size=10,
            resource_scope=ResourceScope.GLOBAL,
            course_key="OS-001",
            visibility="shared",
            status=ResourceStatus.INDEXED,
        )
        KnowledgeResource.objects.create(
            owner=other_teacher,
            original_name="other-teacher-private.pdf",
            stored_path=str((Path(self.temp_dir.name) / "other-teacher-private.pdf").resolve()),
            file_type=".pdf",
            file_size=10,
            resource_scope=ResourceScope.TEACHER,
            visibility="owner",
            status=ResourceStatus.INDEXED,
        )

        self.client.force_login(self.teacher)
        response = self.client.get("/api/kb/resources/")

        self.assertEqual(response.status_code, 200)
        names = {item["name"] for item in response.json()["resources"]}
        self.assertIn("os-global-courseware.pdf", names)
        self.assertNotIn("other-teacher-private.pdf", names)

    def test_owner_keeps_non_offline_resource_access_regardless_of_scope(self):
        owned = KnowledgeResource.objects.create(
            owner=self.student,
            original_name="student-owned-teacher-scope.pdf",
            stored_path=str((Path(self.temp_dir.name) / "student-owned-teacher-scope.pdf").resolve()),
            file_type=".pdf",
            file_size=10,
            resource_scope=ResourceScope.TEACHER,
            visibility="owner",
            status=ResourceStatus.INDEXING,
        )
        other_student = User.objects.create_user(phone="13800139019", password="secure-pass-123")

        self.client.force_login(self.student)
        owner_response = self.client.get("/api/kb/resources/")
        self.client.force_login(other_student)
        other_response = self.client.get("/api/kb/resources/")

        self.assertIn(str(owned.pk), {item["id"] for item in owner_response.json()["resources"]})
        self.assertNotIn(str(owned.pk), {item["id"] for item in other_response.json()["resources"]})

    def test_resource_list_includes_wiki_build_status(self):
        resource = KnowledgeResource.objects.create(
            owner=self.teacher,
            original_name="wiki-ready.pdf",
            stored_path=str((Path(self.temp_dir.name) / "wiki-ready.pdf").resolve()),
            file_type=".pdf",
            file_size=10,
            resource_scope=ResourceScope.TEACHER,
            status=ResourceStatus.INDEXED,
        )
        concept = KnowledgeConcept.objects.create(
            owner=self.teacher,
            resource=resource,
            preferred_label="Root Locus",
            normalized_label="root locus",
            summary="A graph method.",
            source_count=1,
            confidence=0.8,
        )
        KnowledgeDefinition.objects.create(
            owner=self.teacher,
            concept=concept,
            resource=resource,
            definition="Root Locus is a graph method.",
            confidence=0.8,
        )

        self.client.force_login(self.teacher)
        response = self.client.get("/api/kb/resources/")

        self.assertEqual(response.status_code, 200)
        wiki = response.json()["resources"][0]["wiki"]
        self.assertEqual(wiki["status"], "ready")
        self.assertEqual(wiki["conceptCount"], 1)
        self.assertEqual(wiki["definitionCount"], 1)

    def test_teacher_can_rebuild_resource_wiki_from_knowledge_base(self):
        path = Path(self.temp_dir.name) / "wiki-rebuild.pdf"
        path.write_text("Root Locus is a graphical method for analyzing closed-loop pole movement.", encoding="utf-8")
        resource = KnowledgeResource.objects.create(
            owner=self.teacher,
            original_name="wiki-rebuild.pdf",
            stored_path=str(path.resolve()),
            file_type=".pdf",
            file_size=path.stat().st_size,
            resource_scope=ResourceScope.TEACHER,
            status=ResourceStatus.INDEXED,
        )
        page = KnowledgePage.objects.create(resource=resource, page_number=1)
        KnowledgeBlock.objects.create(resource=resource, page=page, modality="text", text=path.read_text(encoding="utf-8"), confidence=0.9)

        self.client.force_login(self.teacher)
        response = self.client.post(f"/api/kb/resources/{resource.id}/wiki/rebuild/")

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload["success"])
        self.assertEqual(payload["wiki"]["status"], "ready")
        self.assertTrue(KnowledgeConcept.objects.filter(resource=resource, normalized_label="root locus").exists())
        self.assertEqual(payload["resource"]["wiki"]["conceptCount"], 1)

    @patch("backend.apps.knowledge_base.tasks.run_index_job.delay")
    def test_teacher_can_queue_batch_wiki_rebuild_from_knowledge_base(self, delay):
        delay.return_value.id = "wiki-batch-task-id"
        KnowledgeResource.objects.create(
            owner=self.teacher,
            original_name="teacher-wiki.pdf",
            stored_path=str((Path(self.temp_dir.name) / "teacher-wiki.pdf").resolve()),
            file_type=".pdf",
            file_size=10,
            resource_scope=ResourceScope.TEACHER,
            status=ResourceStatus.INDEXED,
        )
        KnowledgeResource.objects.create(
            owner=self.student,
            original_name="student-private-wiki.pdf",
            stored_path=str((Path(self.temp_dir.name) / "student-private-wiki.pdf").resolve()),
            file_type=".pdf",
            file_size=10,
            resource_scope=ResourceScope.PERSONAL,
            status=ResourceStatus.INDEXED,
        )

        self.client.force_login(self.teacher)
        response = self.client.post("/api/kb/wiki/rebuild/")

        self.assertEqual(response.status_code, 202)
        payload = response.json()
        self.assertTrue(payload["success"])
        self.assertEqual(payload["job"]["type"], "wiki")
        self.assertEqual(payload["job"]["totalFiles"], 1)
        self.assertEqual(payload["job"]["status"], "queued")

    def test_student_can_view_own_index_jobs_only(self):
        own_job = KnowledgeIndexJob.objects.create(
            owner=self.student,
            job_type=IndexJobType.UPLOAD,
            total_files=1,
        )
        KnowledgeIndexJob.objects.create(
            owner=self.teacher,
            job_type=IndexJobType.UPLOAD,
            total_files=1,
        )

        self.client.force_login(self.student)
        response = self.client.get("/api/kb/jobs/")

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload["success"])
        self.assertEqual([item["id"] for item in payload["jobs"]], [str(own_job.id)])

    @patch("backend.apps.knowledge_base.views.search_knowledge_base", return_value=[])
    def test_search_is_limited_to_current_teacher_indexed_resources(self, search_knowledge_base):
        other_teacher = User.objects.create_user(
            phone="13800139012",
            password="secure-pass-123",
            role=UserRole.TEACHER,
        )
        own_path = str((Path.cwd() / "own.pdf").resolve())
        other_path = str((Path.cwd() / "other.pdf").resolve())
        KnowledgeResource.objects.create(
            owner=self.teacher,
            original_name="own.pdf",
            stored_path=own_path,
            file_type=".pdf",
            file_size=10,
            status=ResourceStatus.INDEXED,
        )
        KnowledgeResource.objects.create(
            owner=other_teacher,
            original_name="other.pdf",
            stored_path=other_path,
            file_type=".pdf",
            file_size=10,
            status=ResourceStatus.INDEXED,
        )

        self.client.force_login(self.teacher)
        response = self.client.get("/api/kb/search/?q=root&course=自动控制原理")

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()["success"])
        search_knowledge_base.assert_called_once()
        kwargs = search_knowledge_base.call_args.kwargs
        self.assertEqual(kwargs["allowed_source_paths"], {own_path})
        self.assertEqual(kwargs["course"], "自动控制原理")

    @patch("backend.apps.knowledge_base.views.search_knowledge_base", return_value=[])
    def test_search_passes_second_stage_config_to_real_retriever(self, search_knowledge_base):
        own_path = str((Path.cwd() / "stage-two.pdf").resolve())
        KnowledgeResource.objects.create(
            owner=self.teacher,
            original_name="stage-two.pdf",
            stored_path=own_path,
            file_type=".pdf",
            file_size=10,
            status=ResourceStatus.INDEXED,
        )

        self.client.force_login(self.teacher)
        response = self.client.get(
            "/api/kb/search/"
            "?q=root"
            "&top_k=6"
            "&include_assets=1"
            "&fast=1"
            "&mode=bm25"
            "&similarity_threshold=0.42"
            "&bm25_candidates=37"
            "&vector_weight=0.65"
            "&bm25_weight=0.35"
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload["success"])
        self.assertEqual(
            payload["searchConfig"],
            {
                "mode": "bm25",
                "fast": True,
                "topK": 6,
                "includeAssets": True,
                "similarityThreshold": 0.42,
                "bm25Candidates": 37,
                "vectorWeight": 0.65,
                "bm25Weight": 0.35,
            },
        )
        kwargs = search_knowledge_base.call_args.kwargs
        self.assertEqual(kwargs["retrieval_mode"], "bm25")
        self.assertTrue(kwargs["force_lexical"])
        self.assertEqual(kwargs["similarity_threshold"], 0.42)
        self.assertEqual(kwargs["bm25_candidate_count"], 37)
        self.assertEqual(kwargs["vector_weight"], 0.65)
        self.assertEqual(kwargs["bm25_weight"], 0.35)

    @patch("backend.apps.knowledge_base.views.search_knowledge_base")
    def test_search_without_indexed_resources_returns_empty_results(self, search_knowledge_base):
        self.client.force_login(self.teacher)
        response = self.client.get("/api/kb/search/?q=root")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["results"], [])
        search_knowledge_base.assert_not_called()

    @patch("backend.apps.knowledge_base.views.search_knowledge_base")
    def test_search_results_include_original_upload_name(self, search_knowledge_base):
        from backend.domains.knowledge_base.schemas import KnowledgeBaseSearchResult

        stored_path = str((Path.cwd() / "stored" / "uuid_uploaded.xlsx").resolve())
        KnowledgeResource.objects.create(
            owner=self.teacher,
            original_name="uploaded.xlsx",
            stored_path=stored_path,
            file_type=".xlsx",
            file_size=10,
            status=ResourceStatus.INDEXED,
        )
        search_knowledge_base.return_value = [
            KnowledgeBaseSearchResult(
                score=1.0,
                text="标准问题：根轨迹",
                metadata={"source_path": stored_path, "title": "根轨迹"},
            )
        ]

        self.client.force_login(self.teacher)
        response = self.client.get("/api/kb/search/?q=root")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["results"][0]["metadata"]["original_name"], "uploaded.xlsx")

    @patch("backend.apps.knowledge_base.views.search_knowledge_base")
    def test_search_results_include_multimodal_block_location_and_preview(self, search_knowledge_base):
        from backend.domains.knowledge_base.schemas import KnowledgeBaseSearchResult

        stored_path = str((Path.cwd() / "stored" / "root-locus.pptx").resolve())
        resource = KnowledgeResource.objects.create(
            owner=self.teacher,
            original_name="root-locus.pptx",
            stored_path=stored_path,
            file_type=".pptx",
            file_size=10,
            status=ResourceStatus.INDEXED_WITH_WARNINGS,
        )
        parse_run = KnowledgeParseRun.objects.create(resource=resource, status="completed_with_warnings", parser="native")
        page = KnowledgePage.objects.create(resource=resource, parse_run=parse_run, page_number=3, slide_number=3)
        asset = KnowledgeAsset.objects.create(
            resource=resource,
            page=page,
            kind="rendered_slide",
            storage_path="internal-slide.png",
            mime_type="image/png",
            width=800,
            height=450,
        )
        KnowledgeBlock.objects.create(
            resource=resource,
            page=page,
            asset=asset,
            modality="slide_image",
            text="Root locus slide OCR",
            confidence=0.7,
        )
        text_block = KnowledgeBlock.objects.create(
            resource=resource,
            page=page,
            modality="text",
            text="Root locus equation: s² + 2s + Kg = 0",
            confidence=0.68,
            metadata={
                "ocr": True,
                "markdownText": "Root locus equation: $s^2 + 2s + Kg = 0$",
                "contentFormat": "markdown",
                "formatFeatures": ["math"],
            },
        )
        neighbor_block = KnowledgeBlock.objects.create(
            resource=resource,
            page=page,
            modality="text",
            text="The context explains how deadlock prevention maps to process scheduling constraints.",
            confidence=0.72,
        )
        search_knowledge_base.return_value = [
            KnowledgeBaseSearchResult(
                score=0.91,
                text="Root locus equation: s² + 2s + Kg = 0",
                metadata={"source_path": stored_path, "title": "root locus"},
            )
        ]

        self.client.force_login(self.teacher)
        response = self.client.get("/api/kb/search/?q=root&include_assets=1")

        self.assertEqual(response.status_code, 200)
        result = response.json()["results"][0]
        self.assertEqual(result["metadata"]["original_name"], "root-locus.pptx")
        self.assertEqual(result["blockId"], str(text_block.id))
        self.assertEqual(result["modality"], "text")
        self.assertEqual(result["location"], {"page": 3, "slide": 3, "bbox": []})
        self.assertEqual(result["asset"]["assetId"], str(asset.id))
        self.assertEqual(result["asset"]["kind"], "rendered_slide")
        self.assertEqual(result["asset"]["previewUrl"], f"/api/kb/assets/{asset.id}/preview/")
        self.assertEqual(result["previewUrl"], f"/api/kb/assets/{asset.id}/preview/")
        self.assertEqual(result["markdownSnippet"], "Root locus equation: $s^2 + 2s + Kg = 0$")
        self.assertEqual(result["contentFormat"], "markdown")
        self.assertEqual(result["formatFeatures"], ["math"])
        self.assertIn("Root locus equation: $s^2 + 2s + Kg = 0$", result["contextSnippet"])
        self.assertIn("deadlock prevention", result["contextSnippet"])
        self.assertEqual([item["blockId"] for item in result["contextBlocks"]], [str(text_block.id), str(neighbor_block.id)])
        self.assertNotIn("internal-slide.png", str(result))

    def test_parse_detail_returns_multimodal_pages_and_blocks(self):
        resource = KnowledgeResource.objects.create(
            owner=self.teacher,
            original_name="question.png",
            stored_path=str((Path.cwd() / "question.png").resolve()),
            file_type=".png",
            file_size=10,
            status=ResourceStatus.INDEXED_WITH_WARNINGS,
        )
        parse_run = KnowledgeParseRun.objects.create(
            resource=resource,
            status="completed_with_warnings",
            parser="native",
            warnings=["OCR not configured"],
            quality_score=0.45,
        )
        page = KnowledgePage.objects.create(resource=resource, parse_run=parse_run, page_number=1, summary="image")
        asset = KnowledgeAsset.objects.create(resource=resource, page=page, kind="original", storage_path="internal.png")
        KnowledgeBlock.objects.create(resource=resource, page=page, asset=asset, modality="image", text="question")

        self.client.force_login(self.teacher)
        response = self.client.get(f"/api/kb/resources/{resource.id}/parse/")

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload["success"])
        self.assertEqual(payload["multimodal"]["parseStatus"], "completed_with_warnings")
        self.assertEqual(payload["pages"][0]["pageNumber"], 1)
        self.assertEqual(payload["blocks"][0]["modality"], "image")

    def test_teacher_can_preview_own_asset_without_exposing_storage_path(self):
        path = Path(self.temp_dir.name) / "asset-preview.png"
        image = Image.new("RGB", (4, 3), color=(255, 0, 0))
        image.save(path)
        resource = KnowledgeResource.objects.create(
            owner=self.teacher,
            original_name="asset-preview.png",
            stored_path=str(path),
            file_type=".png",
            file_size=path.stat().st_size,
            status=ResourceStatus.INDEXED_WITH_WARNINGS,
        )
        parse_run = KnowledgeParseRun.objects.create(resource=resource, status="completed_with_warnings", parser="native")
        page = KnowledgePage.objects.create(resource=resource, parse_run=parse_run, page_number=1)
        asset = KnowledgeAsset.objects.create(resource=resource, page=page, kind="original", storage_path=str(path), mime_type="image/png")
        KnowledgeBlock.objects.create(resource=resource, page=page, asset=asset, modality="image", text="asset preview")

        self.client.force_login(self.teacher)
        detail = self.client.get(f"/api/kb/resources/{resource.id}/parse/").json()
        response = self.client.get(f"/api/kb/assets/{asset.id}/preview/")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "image/png")
        self.assertIn("previewUrl", detail["blocks"][0])
        self.assertNotIn(str(path), str(detail))
        for closer in getattr(response, "_resource_closers", []):
            closer()

    def test_teacher_cannot_preview_other_teacher_asset(self):
        other_teacher = User.objects.create_user(phone="13800139013", password="secure-pass-123", role=UserRole.TEACHER)
        path = Path(self.temp_dir.name) / "other-asset.png"
        Image.new("RGB", (2, 2), color=(0, 0, 255)).save(path)
        resource = KnowledgeResource.objects.create(
            owner=other_teacher,
            original_name="other-asset.png",
            stored_path=str(path),
            file_type=".png",
            file_size=path.stat().st_size,
            status=ResourceStatus.INDEXED_WITH_WARNINGS,
        )
        asset = KnowledgeAsset.objects.create(resource=resource, kind="original", storage_path=str(path), mime_type="image/png")

        self.client.force_login(self.teacher)
        response = self.client.get(f"/api/kb/assets/{asset.id}/preview/")

        self.assertEqual(response.status_code, 404)

    @patch("backend.apps.knowledge_base.tasks.run_index_job.delay")
    def test_teacher_can_reparse_single_resource(self, delay):
        delay.return_value.id = "single-parse-task-id"
        path = Path(self.temp_dir.name) / "question.png"
        path.write_bytes(b"fake-image")
        resource = KnowledgeResource.objects.create(
            owner=self.teacher,
            original_name="question.png",
            stored_path=str(path),
            file_type=".png",
            file_size=path.stat().st_size,
            status=ResourceStatus.INDEXED_WITH_WARNINGS,
        )

        self.client.force_login(self.teacher)
        response = self.client.post(f"/api/kb/resources/{resource.id}/reparse/")

        self.assertEqual(response.status_code, 202)
        payload = response.json()
        self.assertTrue(payload["success"])
        self.assertEqual(payload["job"]["type"], "parse")
        self.assertEqual(payload["job"]["status"], IndexJobStatus.QUEUED)
        self.assertEqual(payload["resource"]["status"], ResourceStatus.INDEXING)
        resource.refresh_from_db()
        self.assertEqual(resource.status, ResourceStatus.INDEXING)
        self.assertEqual(resource.index_job_id, KnowledgeIndexJob.objects.get(celery_task_id="single-parse-task-id").id)

    @patch("backend.apps.knowledge_base.tasks.run_index_job.delay")
    def test_teacher_can_force_reparse_all_own_resources(self, delay):
        delay.return_value.id = "force-parse-task-id"
        own_path = Path(self.temp_dir.name) / "own.pdf"
        other_path = Path(self.temp_dir.name) / "other.pdf"
        own_path.write_bytes(MINIMAL_TEXT_PDF)
        other_path.write_bytes(MINIMAL_TEXT_PDF)
        own_resource = KnowledgeResource.objects.create(
            owner=self.teacher,
            original_name="own.pdf",
            stored_path=str(own_path),
            file_type=".pdf",
            file_size=own_path.stat().st_size,
            status=ResourceStatus.INDEXED,
        )
        other_teacher = User.objects.create_user(phone="13800139014", password="secure-pass-123", role=UserRole.TEACHER)
        KnowledgeResource.objects.create(
            owner=other_teacher,
            original_name="other.pdf",
            stored_path=str(other_path),
            file_type=".pdf",
            file_size=other_path.stat().st_size,
            status=ResourceStatus.INDEXED,
        )

        self.client.force_login(self.teacher)
        with patch.dict(os.environ, {"A3_ALLOW_ONLINE_VISUAL_EMBEDDING": "1", "TRANSFORMERS_OFFLINE": "0"}):
            response = self.client.post("/api/kb/reparse/force/")

        self.assertEqual(response.status_code, 202)
        payload = response.json()
        self.assertTrue(payload["success"])
        self.assertEqual(payload["job"]["type"], "parse")
        self.assertEqual(payload["job"]["status"], IndexJobStatus.QUEUED)
        self.assertEqual(payload["resources"][0]["id"], str(own_resource.id))
        own_resource.refresh_from_db()
        self.assertEqual(own_resource.index_job_id, KnowledgeIndexJob.objects.get(celery_task_id="force-parse-task-id").id)

    @patch("backend.apps.knowledge_base.tasks.run_index_job.delay")
    def test_student_force_reparse_only_touches_own_resources(self, delay):
        delay.return_value.id = "student-force-parse-task-id"
        own_path = Path(self.temp_dir.name) / "student-reparse.pdf"
        other_path = Path(self.temp_dir.name) / "teacher-reparse.pdf"
        own_path.write_bytes(b"student")
        other_path.write_bytes(b"teacher")
        own_resource = KnowledgeResource.objects.create(
            owner=self.student,
            original_name="student-reparse.pdf",
            stored_path=str(own_path),
            file_type=".pdf",
            file_size=own_path.stat().st_size,
            resource_scope=ResourceScope.PERSONAL,
            status=ResourceStatus.INDEXED,
        )
        KnowledgeResource.objects.create(
            owner=self.teacher,
            original_name="teacher-reparse.pdf",
            stored_path=str(other_path),
            file_type=".pdf",
            file_size=other_path.stat().st_size,
            status=ResourceStatus.INDEXED,
        )
        self.client.force_login(self.student)

        response = self.client.post("/api/kb/reparse/force/")

        self.assertEqual(response.status_code, 202)
        payload = response.json()
        self.assertTrue(payload["success"])
        self.assertEqual(payload["job"]["type"], "parse")
        self.assertEqual(payload["resources"][0]["id"], str(own_resource.id))
        own_resource.refresh_from_db()
        self.assertEqual(own_resource.index_job_id, KnowledgeIndexJob.objects.get(celery_task_id="student-force-parse-task-id").id)
