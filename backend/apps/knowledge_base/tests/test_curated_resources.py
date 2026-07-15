from __future__ import annotations

import tempfile
import uuid
from datetime import datetime
from io import StringIO
from pathlib import Path
from unittest.mock import patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.core.exceptions import ValidationError
from django.test import SimpleTestCase, TestCase, override_settings
from django.urls import reverse
from docx import Document
from openpyxl import Workbook

from backend.apps.accounts.models import User
from backend.apps.knowledge_base.docx_preview import (
    DocxPreviewSourceUnavailable,
    DocxPreviewUnavailable,
)
from backend.apps.knowledge_base.models import (
    CuratedLearningResource,
    CuratedResourceDeliveryMode,
    CuratedResourceDifficulty,
    CuratedResourceType,
    CuratedResourceVerificationStatus,
    KnowledgeResource,
    ResourceScope,
)
from backend.apps.knowledge_base.services import save_upload
from backend.apps.learning.models import Chapter, Course, KnowledgePoint
from backend.domains.knowledge_base.document_loader import iter_source_files, load_documents
from backend.domains.knowledge_base.multimodal_parser import parse_resource


def write_sample_docx(path: Path, *, title: str = "进程与线程") -> None:
    document = Document()
    document.add_heading(title, level=1)
    document.add_paragraph("进程是程序的一次动态执行，线程是处理机调度的基本单位。")
    document.add_heading("调度算法", level=2)
    document.add_paragraph("时间片轮转需要在时间片结束后将未完成进程重新入队。")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "算法"
    table.rows[0].cells[1].text = "特点"
    row = table.add_row().cells
    row[0].text = "RR"
    row[1].text = "时间片"
    document.save(path)


def write_catalog_docx(path: Path, metadata: dict[str, str]) -> None:
    document = Document()
    document.add_heading(metadata.get("标题", path.stem), level=1)
    table = document.add_table(rows=0, cols=2)
    for key, value in metadata.items():
        if key == "标题":
            continue
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = value
    document.add_heading("核心讲解", level=1)
    document.add_paragraph(f"{metadata['核心知识点']} 的结构化讲解与例题。")
    document.save(path)


def write_video_catalog(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "视频资源清单"
    worksheet.append([
        "主题", "视频标题", "平台", "发布者/机构", "真实URL", "总时长", "关键知识点时间段",
        "难度", "适用学生画像", "薄弱点触发条件", "推荐理由", "是否免费", "核验日期",
    ])
    rows = [
        ("进程与线程", "进程基本概念", "OS-PROCESS"),
        ("进程与线程", "线程基本概念", "OS-THREAD"),
        ("CPU调度", "先来先服务和短作业优先算法", "OS-SCHEDULING"),
        ("CPU调度", "时间片轮转与多级反馈队列算法", "OS-SCHEDULING"),
        ("进程同步", "信号量机制定义", "OS-SYNC"),
        ("进程同步", "生产者消费者问题", "OS-SYNC"),
        ("死锁", "死锁基本概念", "OS-DEADLOCK"),
        ("死锁", "避免死锁与银行家算法", "OS-DEADLOCK"),
        ("虚拟内存", "按需调页", "OS-MEMORY"),
        ("虚拟内存", "页面置换", "OS-VIRTUAL-MEMORY"),
        ("文件系统", "文件和文件系统", "OS-FILE"),
        ("文件系统", "文件的物理结构", "OS-FILE"),
    ]
    for index, (topic, title, point_code) in enumerate(rows, start=1):
        worksheet.append([
            topic,
            title,
            "哔哩哔哩（B站）",
            "高校公开课程",
            f"https://www.bilibili.com/video/BVTEST?p={index}",
            "18:30",
            "00:00-18:30",
            "基础" if index in {1, 2, 7, 11} else "中等",
            "需要针对性补齐该知识点的学生",
            f"{point_code} 对应题目连续答错",
            f"围绕{title}进行单主题讲解。",
            "是（公开免费观看）",
            datetime(2026, 7, 15),
        ])
    workbook.save(path)


class DocxKnowledgeBaseSupportTests(SimpleTestCase):
    def test_docx_is_discovered_and_loaded_as_section_and_table_documents(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "进程与线程.docx"
            write_sample_docx(source_path)

            discovered = iter_source_files(Path(temp_dir))
            documents, skipped = load_documents([source_path])

        self.assertEqual(discovered, [source_path])
        self.assertEqual(skipped, [])
        self.assertGreaterEqual(len(documents), 3)
        text = "\n".join(document.get_content() for document in documents)
        metadata = [document.metadata for document in documents]
        self.assertIn("时间片轮转", text)
        self.assertIn("| 算法 | 特点 |", text)
        self.assertTrue(any(item.get("section") == "调度算法" for item in metadata))
        self.assertTrue(any(item.get("block_type") == "table" for item in metadata))
        self.assertTrue(all(item.get("source_type") == "docx" for item in metadata))

    def test_docx_multimodal_parse_preserves_heading_and_table_locations(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "进程与线程.docx"
            write_sample_docx(source_path)

            result = parse_resource(source_path)

        self.assertEqual(result.status, "completed")
        blocks = [block for page in result.pages for block in page.blocks]
        self.assertTrue(any(block.modality == "text" and "时间片轮转" in block.text for block in blocks))
        table = next(block for block in blocks if block.modality == "table")
        self.assertIn("| RR | 时间片 |", table.metadata["markdownText"])
        self.assertEqual(table.metadata["section"], "调度算法")
        self.assertEqual(table.metadata["tableIndex"], 1)
        self.assertEqual(result.result["tableCount"], 1)


class CuratedLearningResourceModelTests(TestCase):
    def setUp(self):
        self.owner = User.objects.create_user(phone="13800139201", password="secure-pass-123")

    def test_docx_upload_is_accepted_for_indexing(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "进程与线程.docx"
            write_sample_docx(source_path)
            upload = SimpleUploadedFile(
                source_path.name,
                source_path.read_bytes(),
                content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
            with override_settings(KNOWLEDGE_UPLOAD_ROOT=Path(temp_dir) / "uploads"):
                resource = save_upload(self.owner, upload)

        self.assertEqual(resource.file_type, ".docx")
        self.assertEqual(resource.original_name, "进程与线程.docx")

    def test_internal_document_requires_knowledge_resource(self):
        course = Course.objects.create(code="OS-VALIDATION-INTERNAL", title="站内资源校验")
        resource = CuratedLearningResource(
            resource_code="OS-DOC-INVALID",
            course=course,
            resource_type=CuratedResourceType.LECTURE_DOC,
            delivery_mode=CuratedResourceDeliveryMode.INTERNAL_DOCUMENT,
            title="缺少知识库文件",
        )

        with self.assertRaisesMessage(ValidationError, "站内文档必须关联知识库文件"):
            resource.full_clean()

    def test_external_resource_requires_https_url(self):
        course = Course.objects.create(code="OS-VALIDATION-EXTERNAL", title="外部资源校验")
        resource = CuratedLearningResource(
            resource_code="OS-VIDEO-INVALID",
            course=course,
            resource_type=CuratedResourceType.VIDEO,
            delivery_mode=CuratedResourceDeliveryMode.EXTERNAL_LINK,
            title="不安全的视频链接",
            external_url="http://example.com/video",
        )

        with self.assertRaisesMessage(ValidationError, "外部资源必须使用 HTTPS 链接"):
            resource.full_clean()

    def test_internal_document_links_course_knowledge_points_and_indexed_file(self):
        course = Course.objects.create(code="OS-CATALOG", title="操作系统资源目录")
        chapter = Chapter.objects.create(course=course, order=1, title="进程管理")
        point = KnowledgePoint.objects.create(
            chapter=chapter,
            code="OS-CATALOG-PROCESS",
            order=1,
            title="进程与线程",
            difficulty=2,
            estimated_minutes=40,
        )
        knowledge_resource = KnowledgeResource.objects.create(
            owner=self.owner,
            original_name="进程与线程.docx",
            stored_path="/tmp/process.docx",
            file_type=".docx",
            resource_scope=ResourceScope.GLOBAL,
            course_key=course.code,
            visibility="shared",
        )

        resource = CuratedLearningResource.objects.create(
            resource_code="OS-DOC-TEST-001",
            course=course,
            knowledge_resource=knowledge_resource,
            resource_type=CuratedResourceType.LECTURE_DOC,
            delivery_mode=CuratedResourceDeliveryMode.INTERNAL_DOCUMENT,
            title="进程与线程讲解",
            difficulty=CuratedResourceDifficulty.INTERMEDIATE,
            estimated_minutes=45,
            weakness_triggers=["混淆进程与线程"],
            completion_criteria=["完成 5 道自测题"],
            verification_status=CuratedResourceVerificationStatus.VERIFIED,
        )
        resource.knowledge_points.add(point)

        self.assertEqual(resource.resource_code, "OS-DOC-TEST-001")
        self.assertEqual(resource.knowledge_resource, knowledge_resource)
        self.assertEqual(list(resource.knowledge_points.values_list("code", flat=True)), [point.code])
        self.assertEqual(knowledge_resource.curated_resource, resource)


class CuratedResourcePreviewApiTests(TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.root = Path(self.temp_dir.name)
        self.source = self.root / "managed.docx"
        write_sample_docx(self.source)
        suffix = str(uuid.uuid4().int)[-8:]
        self.user = User.objects.create_user(phone=f"139{suffix}", password="secure-pass-123")
        self.course = Course.objects.create(code=f"OS-PREVIEW-{suffix}", title="操作系统预览")
        self.knowledge_resource = KnowledgeResource.objects.create(
            owner=self.user,
            original_name="managed.docx",
            stored_path=str(self.source),
            file_type=".docx",
            resource_scope=ResourceScope.GLOBAL,
            course_key=self.course.code,
            visibility="shared",
        )
        self.resource = CuratedLearningResource.objects.create(
            resource_code=f"OS-DOC-PREVIEW-{suffix}",
            course=self.course,
            knowledge_resource=self.knowledge_resource,
            resource_type=CuratedResourceType.LECTURE_DOC,
            delivery_mode=CuratedResourceDeliveryMode.INTERNAL_DOCUMENT,
            title="Word 在线预览测试",
            verification_status=CuratedResourceVerificationStatus.VERIFIED,
            content_hash="abcdef123456",
        )
        self.preview_url = reverse("knowledge-base-curated-library-preview", args=[self.resource.resource_code])
        self.content_url = reverse("knowledge-base-curated-library-content", args=[self.resource.resource_code])
        self.detail_url = reverse("knowledge-base-curated-library-detail", args=[self.resource.resource_code])

    def test_authenticated_user_receives_inline_private_pdf(self):
        pdf = self.root / "preview.pdf"
        pdf.write_bytes(b"%PDF-1.4\npreview")
        self.client.force_login(self.user)

        with patch("backend.apps.knowledge_base.views.get_or_create_docx_preview", return_value=pdf):
            response = self.client.get(self.preview_url)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response["Content-Disposition"].startswith("inline"))
        self.assertEqual(response["Cache-Control"], "private, max-age=86400")
        self.assertEqual(response["X-Content-Type-Options"], "nosniff")
        self.assertEqual(b"".join(response.streaming_content), b"%PDF-1.4\npreview")

    def test_preview_requires_authentication(self):
        response = self.client.get(self.preview_url)

        self.assertEqual(response.status_code, 401)
        self.assertEqual(response.json()["error"], "请先登录")

    def test_inactive_or_unverified_resources_are_hidden(self):
        self.client.force_login(self.user)
        self.resource.is_active = False
        self.resource.save(update_fields=["is_active"])

        response = self.client.get(self.preview_url)

        self.assertEqual(response.status_code, 404)

    def test_external_resource_cannot_use_document_preview(self):
        self.client.force_login(self.user)
        self.resource.delivery_mode = CuratedResourceDeliveryMode.EXTERNAL_LINK
        self.resource.external_url = "https://www.bilibili.com/video/BVTEST"
        self.resource.knowledge_resource = None
        self.resource.save(update_fields=["delivery_mode", "external_url", "knowledge_resource"])

        response = self.client.get(self.preview_url)

        self.assertEqual(response.status_code, 400)
        self.assertIn("Word", response.json()["error"])

    def test_missing_source_and_conversion_failure_have_readable_errors(self):
        self.client.force_login(self.user)
        with patch(
            "backend.apps.knowledge_base.views.get_or_create_docx_preview",
            side_effect=DocxPreviewSourceUnavailable("资料原文件暂时不可用"),
        ):
            missing = self.client.get(self.preview_url)
        with patch(
            "backend.apps.knowledge_base.views.get_or_create_docx_preview",
            side_effect=DocxPreviewUnavailable("internal converter detail /secret/path"),
        ):
            unavailable = self.client.get(self.preview_url)

        self.assertEqual(missing.status_code, 404)
        self.assertEqual(missing.json()["error"], "资料原文件暂时不可用")
        self.assertEqual(unavailable.status_code, 503)
        self.assertEqual(unavailable.json()["error"], "Word 在线预览暂时不可用，可下载原文档")
        self.assertNotIn("/secret/path", unavailable.content.decode("utf-8"))

    def test_content_endpoint_downloads_the_original_docx_as_attachment(self):
        self.client.force_login(self.user)

        response = self.client.get(self.content_url)
        downloaded = b"".join(response.streaming_content)

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response["Content-Disposition"].startswith("attachment"))
        self.assertIn("managed.docx", response["Content-Disposition"])
        self.assertEqual(downloaded, self.source.read_bytes())

    def test_detail_splits_preview_and_download_actions_without_storage_path(self):
        self.client.force_login(self.user)

        response = self.client.get(self.detail_url)

        self.assertEqual(response.status_code, 200)
        payload = response.json()["resource"]
        self.assertFalse(payload["action"]["external"])
        self.assertEqual(payload["action"]["preview"]["href"], self.preview_url)
        self.assertEqual(payload["action"]["preview"]["label"], "在线预览")
        self.assertEqual(payload["action"]["download"]["href"], self.content_url)
        self.assertEqual(payload["action"]["download"]["label"], "下载原文档")
        self.assertNotIn("stored_path", response.content.decode("utf-8"))


class ImportOsCuratedResourcesCommandTests(TestCase):
    DOCUMENTS = (
        ("智学引擎A3_OS-DOC-001_进程线程与CPU调度.docx", "OS-DOC-001", "讲解文档", "进程、线程、CPU调度"),
        ("智学引擎A3_OS-DOC-002_虚拟内存与页面置换.docx", "OS-DOC-002", "讲解文档", "虚拟内存、页面置换"),
        ("智学引擎A3_OS-DOC-003_进程同步信号量与死锁.docx", "OS-DOC-003", "讲解文档", "信号量、死锁"),
        ("智学引擎A3_OS-LAB-001_Linux进程调度观察实验.docx", "OS-LAB-001", "实操案例", "Linux进程、CPU调度"),
        ("智学引擎A3_OS-LAB-002_用Python实现页面置换算法模拟器.docx", "OS-LAB-002", "实操案例", "FIFO、LRU、页面置换"),
        ("智学引擎A3_OS-LAB-003_生产者消费者同步问题.docx", "OS-LAB-003", "实操案例", "生产者消费者、信号量"),
        ("智学引擎A3_操作系统_进程状态与状态转换.docx", "OS-001-LD-02-01", "讲解文档", "进程状态、状态转换"),
    )

    def test_command_imports_nineteen_resources_and_is_idempotent(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source_dir = Path(temp_dir) / "资料"
            upload_dir = Path(temp_dir) / "uploads"
            source_dir.mkdir()
            for file_name, resource_code, resource_type, points in self.DOCUMENTS:
                write_catalog_docx(
                    source_dir / file_name,
                    {
                        "标题": Path(file_name).stem,
                        "资源编号": resource_code,
                        "课程编号": "OS-001",
                        "资源类型": resource_type,
                        "所属章节": "操作系统",
                        "核心知识点": points,
                        "难度": "中等",
                        "前置知识": "操作系统基础；计算机组成基础",
                        "适用学生画像": "需要补齐基础概念的本科生",
                        "薄弱点触发条件": "对应知识点连续答错；无法完成基础推演",
                        "预计学习时间": "45～60 分钟",
                        "学习目标": "理解核心概念；完成对应练习",
                        "完成判据": "自测题至少答对 4 道；能够解释关键步骤",
                        "推荐理由模板": "检测到对应薄弱点，建议完成该资源。",
                        "检索关键词": points,
                    },
                )
            write_video_catalog(source_dir / "操作系统免费视频资源_12个精准推荐视频.xlsx")

            first_stdout = StringIO()
            second_stdout = StringIO()
            with override_settings(KNOWLEDGE_UPLOAD_ROOT=upload_dir):
                call_command(
                    "import_os_curated_resources",
                    "--source-dir",
                    str(source_dir),
                    "--owner-phone",
                    "13800139202",
                    stdout=first_stdout,
                )
                first_ids = set(CuratedLearningResource.objects.values_list("id", flat=True))
                call_command(
                    "import_os_curated_resources",
                    "--source-dir",
                    str(source_dir),
                    "--owner-phone",
                    "13800139202",
                    stdout=second_stdout,
                )
                with patch(
                    "backend.apps.knowledge_base.management.commands.import_os_curated_resources.create_upload_index_job"
                ) as create_job:
                    create_job.return_value.id = "index-job-1"
                    call_command(
                        "import_os_curated_resources",
                        "--source-dir",
                        str(source_dir),
                        "--owner-phone",
                        "13800139202",
                        "--enqueue-index",
                        stdout=StringIO(),
                    )

        self.assertEqual(CuratedLearningResource.objects.count(), 19)
        self.assertEqual(KnowledgeResource.objects.filter(file_type=".docx").count(), 7)
        self.assertEqual(set(CuratedLearningResource.objects.values_list("id", flat=True)), first_ids)
        self.assertEqual(
            set(CuratedLearningResource.objects.get(resource_code="OS-DOC-001").knowledge_points.values_list("code", flat=True)),
            {"OS-PROCESS", "OS-THREAD", "OS-SCHEDULING"},
        )
        video = CuratedLearningResource.objects.get(resource_code="OS-VIDEO-004")
        self.assertEqual(video.delivery_mode, CuratedResourceDeliveryMode.EXTERNAL_LINK)
        self.assertEqual(video.external_url, "https://www.bilibili.com/video/BVTEST?p=4")
        self.assertEqual(list(video.knowledge_points.values_list("code", flat=True)), ["OS-SCHEDULING"])
        self.assertIn("新增 19", first_stdout.getvalue())
        self.assertIn("新增 0", second_stdout.getvalue())
        create_job.assert_called_once()
        self.assertEqual(len(create_job.call_args.args[1]), 7)
