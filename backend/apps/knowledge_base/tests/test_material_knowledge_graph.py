from __future__ import annotations

import tempfile
from pathlib import Path
from unittest.mock import patch

from django.test import Client, TestCase

from backend.apps.accounts.models import User
from backend.apps.compat.models import AgentRun
from backend.apps.compat.agent_workflows import get_agent_workflow_contract
from backend.apps.knowledge_base.graph_builder import rebuild_resource_graph
from backend.apps.knowledge_base.models import (
    KnowledgeBlock,
    KnowledgeGraphEntity,
    KnowledgeGraphRelation,
    KnowledgeParseRun,
    KnowledgePage,
    KnowledgeResource,
    ResourceScope,
    ResourceStatus,
)
from backend.domains.knowledge_base.multimodal_parser import VisionUnderstanding, parse_resource


class MaterialKnowledgeGraphTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(phone="13800139111", password="secure-pass-123")
        self.other = User.objects.create_user(phone="13800139222", password="secure-pass-123")
        self.client = Client()
        self.client.force_login(self.user)
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)

    def test_builds_graph_from_existing_blocks_with_source_evidence(self):
        resource = self._resource("control.pdf")
        page = KnowledgePage.objects.create(resource=resource, page_number=3, summary="根轨迹方法说明")
        KnowledgeBlock.objects.create(
            resource=resource,
            page=page,
            modality="text",
            text="根轨迹方法用于分析闭环极点。传递函数是描述线性系统输入输出关系的数学模型。",
            confidence=0.86,
        )

        result = rebuild_resource_graph(resource)

        self.assertEqual(result.status, "ready")
        self.assertGreaterEqual(KnowledgeGraphEntity.objects.filter(resource=resource).count(), 2)
        relation = KnowledgeGraphRelation.objects.filter(resource=resource).first()
        self.assertIsNotNone(relation)
        self.assertEqual(relation.evidence[0]["resourceTitle"], "control.pdf")
        self.assertIn("第 3 页", relation.evidence[0]["location"])
        self.assertNotIn(str(Path(self.temp_dir.name)), relation.evidence[0]["location"])

    def test_material_graph_evidence_exposes_markdown_snippet(self):
        resource = self._resource("formula.pdf")
        page = KnowledgePage.objects.create(resource=resource, page_number=1, summary="公式证据")
        KnowledgeBlock.objects.create(
            resource=resource,
            page=page,
            modality="text",
            text="Root Locus 方程为：s² + 2s + Kg = 0。Closed Loop Pole 用于稳定性分析。",
            confidence=0.88,
            metadata={
                "markdownText": "Root Locus 方程为：$s^2 + 2s + Kg = 0$。Closed Loop Pole 用于稳定性分析。",
                "contentFormat": "markdown",
                "formatFeatures": ["math"],
            },
        )

        rebuild_resource_graph(resource)

        entity = KnowledgeGraphEntity.objects.filter(resource=resource).first()
        self.assertIsNotNone(entity)
        evidence = entity.source_locations[0]
        self.assertEqual(evidence["contentFormat"], "markdown")
        self.assertEqual(
            evidence["markdownSnippet"],
            "Root Locus 方程为：$s^2 + 2s + Kg = 0$。Closed Loop Pole 用于稳定性分析。",
        )
        self.assertIn("Root Locus 方程为：s² + 2s + Kg = 0。Closed Loop Pole 用于稳定性分析。", evidence["snippet"])
        self.assertEqual(evidence["formatFeatures"], ["math"])

    def test_graph_api_returns_material_contract_and_ppt_slide_evidence(self):
        resource = self._resource("lesson.pptx", file_type=".pptx")
        page = KnowledgePage.objects.create(resource=resource, slide_number=5, summary="根轨迹课件")
        KnowledgeBlock.objects.create(
            resource=resource,
            page=page,
            modality="text",
            text="闭环极点影响系统稳定性。根轨迹图用于展示极点变化。",
            confidence=0.77,
            metadata={"blockType": "bullet"},
        )
        rebuild_resource_graph(resource)

        response = self.client.get("/api/kb/graph/")

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload["success"])
        graph = payload["graph"]
        self.assertEqual(graph["schemaVersion"], "a3.material-knowledge-graph.v1")
        self.assertEqual(graph["source"], "knowledge_base_materials")
        self.assertGreaterEqual(graph["summary"]["entityCount"], 1)
        self.assertIn("quality", graph)
        self.assertIn("evidenceCoverage", graph["quality"])
        self.assertIn("relationDensity", graph["quality"])
        evidence = graph["nodes"][0]["evidence"][0]
        self.assertEqual(evidence["resourceTitle"], "lesson.pptx")
        self.assertEqual(evidence["slideIndex"], 5)
        self.assertIn("PPT 第 5 页", evidence["location"])

    def test_graph_api_repairs_legacy_ppt_symbol_font_evidence(self):
        resource = self._resource("legacy-symbol.pptx", file_type=".pptx")
        KnowledgeGraphEntity.objects.create(
            owner=self.user,
            resource=resource,
            label="闭环特征方程",
            normalized_label="闭环特征方程",
            entity_type="formula",
            summary="参数从0\uf0ae\uf0a5，s2=\uf02d2，用\uf0b4表示。",
            confidence=0.8,
            source_count=1,
            source_locations=[
                {
                    "resourceId": str(resource.id),
                    "resourceTitle": resource.original_name,
                    "resourceType": "pptx",
                    "blockId": "legacy-block",
                    "blockType": "text_box",
                    "pageNumber": None,
                    "slideIndex": 3,
                    "location": "PPT 第 3 页 · 文本框",
                    "snippet": "参数从0\uf0ae\uf0a5；s2从（\uf02d2,j0）开始；用\uf0b4表示。",
                    "markdownSnippet": "参数从0\uf0ae\uf0a5；s2从（\uf02d2,j0）开始；用\uf0b4表示。",
                    "contentFormat": "markdown",
                    "formatFeatures": [],
                    "confidence": 0.8,
                    "degraded": False,
                    "warnings": [],
                }
            ],
            degraded=False,
        )

        response = self.client.get("/api/kb/graph/")

        self.assertEqual(response.status_code, 200)
        node = response.json()["graph"]["nodes"][0]
        evidence = node["evidence"][0]
        self.assertNotRegex(node["summary"], r"[\ue000-\uf8ff]")
        self.assertNotRegex(evidence["markdownSnippet"], r"[\ue000-\uf8ff]")
        self.assertIn("0→∞", evidence["markdownSnippet"])
        self.assertIn("（-2,j0）", evidence["markdownSnippet"])
        self.assertIn("×", evidence["markdownSnippet"])

    def test_graph_api_is_scoped_to_accessible_resources(self):
        own = self._resource("mine.pdf")
        own_page = KnowledgePage.objects.create(resource=own, page_number=1)
        KnowledgeBlock.objects.create(resource=own, page=own_page, modality="text", text="根轨迹方法用于控制系统分析。", confidence=0.8)
        rebuild_resource_graph(own)

        other_resource = self._resource("other.pdf", owner=self.other, scope=ResourceScope.PERSONAL)
        other_page = KnowledgePage.objects.create(resource=other_resource, page_number=1)
        KnowledgeBlock.objects.create(resource=other_resource, page=other_page, modality="text", text="状态空间模型用于现代控制理论。", confidence=0.8)
        rebuild_resource_graph(other_resource)

        response = self.client.get("/api/kb/graph/")
        labels = {node["label"] for node in response.json()["graph"]["nodes"]}

        self.assertTrue(any("根轨迹" in label for label in labels))
        self.assertFalse(any("状态空间" in label for label in labels))

    def test_graph_api_returns_focused_neighborhood_by_center_depth_and_limit(self):
        resource = self._resource("focused-neighborhood.pdf")
        center = self._graph_entity(resource, "中心实体")
        first = self._graph_entity(resource, "一跳实体")
        second = self._graph_entity(resource, "二跳实体")
        unrelated = self._graph_entity(resource, "无关实体")
        self._graph_relation(resource, center, first, "depends_on")
        self._graph_relation(resource, first, second, "used_for")
        self._graph_relation(resource, unrelated, self._graph_entity(resource, "无关邻居"), "contains")

        response = self.client.get(f"/api/kb/graph/?centerId={center.id}&depth=1&limit=3")

        self.assertEqual(response.status_code, 200)
        graph = response.json()["graph"]
        labels = {node["label"] for node in graph["nodes"]}
        self.assertEqual(labels, {"中心实体", "一跳实体"})
        self.assertEqual(len(graph["edges"]), 1)
        self.assertEqual(graph["summary"]["entityCount"], 5)
        self.assertEqual(graph["focus"]["centerId"], str(center.id))
        self.assertEqual(graph["focus"]["depth"], "one")
        self.assertEqual(graph["focus"]["limit"], 3)
        self.assertEqual(graph["focus"]["returnedNodeCount"], 2)
        self.assertFalse(graph["focus"]["hasMore"])
        self.assertTrue(graph["focus"]["partial"])

    def test_graph_api_defaults_to_highest_degree_neighborhood_instead_of_full_graph(self):
        resource = self._resource("default-focused.pdf")
        center = self._graph_entity(resource, "最高连接实体")
        for index in range(4):
            neighbor = self._graph_entity(resource, f"邻居{index}")
            self._graph_relation(resource, center, neighbor, "co_occurs")
        self._graph_entity(resource, "孤立实体")

        response = self.client.get("/api/kb/graph/?limit=3")

        self.assertEqual(response.status_code, 200)
        graph = response.json()["graph"]
        labels = {node["label"] for node in graph["nodes"]}
        self.assertIn("最高连接实体", labels)
        self.assertLessEqual(len(labels), 3)
        self.assertNotIn("孤立实体", labels)
        self.assertEqual(graph["summary"]["entityCount"], 6)
        self.assertEqual(graph["focus"]["centerId"], str(center.id))
        self.assertTrue(graph["focus"]["hasMore"])

    def test_rebuild_replaces_old_rows_without_duplicates(self):
        resource = self._resource("reparse.pdf")
        page = KnowledgePage.objects.create(resource=resource, page_number=1)
        KnowledgeBlock.objects.create(resource=resource, page=page, modality="text", text="根轨迹方法用于控制系统分析。", confidence=0.8)
        rebuild_resource_graph(resource)
        first_count = KnowledgeGraphEntity.objects.filter(resource=resource).count()

        KnowledgeBlock.objects.filter(resource=resource).delete()
        KnowledgeBlock.objects.create(resource=resource, page=page, modality="text", text="传递函数是线性系统模型。", confidence=0.8)
        rebuild_resource_graph(resource)

        self.assertLessEqual(KnowledgeGraphEntity.objects.filter(resource=resource).count(), max(first_count, 3))
        labels = {entity.label for entity in KnowledgeGraphEntity.objects.filter(resource=resource)}
        self.assertFalse(any("根轨迹" in label for label in labels))

    def test_pptx_parser_creates_granular_slide_blocks(self):
        import base64

        from pptx import Presentation
        from pptx.util import Inches

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = "Transfer Function"
        textbox = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(5), Inches(1))
        textbox.text_frame.text = "Control System 是 Dynamic Model"
        bullet_box = slide.shapes.add_textbox(Inches(1), Inches(2.4), Inches(5), Inches(1))
        bullet_box.text_frame.text = "• Root Locus 用于 Closed Loop Pole"
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(3.4), Inches(5), Inches(1)).table
        table.cell(0, 0).text = "Concept"
        table.cell(0, 1).text = "Meaning"
        table.cell(1, 0).text = "Gain"
        table.cell(1, 1).text = "Stability"
        image_path = Path(self.temp_dir.name) / "diagram.png"
        image_path.write_bytes(
            base64.b64decode(
                "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
            )
        )
        slide.shapes.add_picture(str(image_path), Inches(5.7), Inches(1.5), Inches(1), Inches(1))
        slide.notes_slide.notes_text_frame.text = "Speaker note: emphasize stability margin."

        pptx_path = Path(self.temp_dir.name) / "granular.pptx"
        presentation.save(pptx_path)

        result = parse_resource(pptx_path)
        block_types = {
            str(block.metadata.get("blockType") or "")
            for page in result.pages
            for block in page.blocks
        }

        self.assertIn("slide_title", block_types)
        self.assertIn("text_box", block_types)
        self.assertIn("bullet", block_types)
        self.assertIn("table", block_types)
        self.assertIn("image_caption", block_types)
        self.assertIn("speaker_note", block_types)
        self.assertGreaterEqual(result.result["pptxBlockDiagnostics"]["tableBlockCount"], 1)
        self.assertGreaterEqual(result.result["pptxBlockDiagnostics"]["imageCaptionBlockCount"], 1)

    @patch("backend.domains.knowledge_base.multimodal_parser._understand_image")
    def test_pptx_picture_caption_uses_vision_provider_metadata(self, understand_image):
        import base64

        from pptx import Presentation
        from pptx.util import Inches

        understand_image.return_value = VisionUnderstanding(
            figures=[{"caption": "root locus diagram", "description": "closed loop poles move along a plotted curve"}],
            keywords=["root locus", "closed loop poles"],
            raw={"provider": "test-vision"},
        )
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        image_path = Path(self.temp_dir.name) / "diagram.png"
        image_path.write_bytes(
            base64.b64decode(
                "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
            )
        )
        slide.shapes.add_picture(str(image_path), Inches(1), Inches(1), Inches(2), Inches(2))
        pptx_path = Path(self.temp_dir.name) / "vision-caption.pptx"
        presentation.save(pptx_path)

        result = parse_resource(pptx_path)
        caption_block = next(
            block
            for page in result.pages
            for block in page.blocks
            if block.metadata.get("blockType") == "image_caption"
        )

        self.assertIn("root locus diagram", caption_block.text)
        self.assertEqual(caption_block.metadata["captionMethod"], "vision_provider")
        self.assertGreaterEqual(caption_block.metadata["captionConfidence"], 0.7)
        self.assertFalse(caption_block.warnings)

    @patch.dict(
        "os.environ",
        {
            "A3_LLM_API_KEY": "test-key",
            "A3_LLM_MODEL": "global-chat-model",
            "A3_GRAPH_LLM_PROVIDER": "dashscope",
            "A3_GRAPH_LLM_MODEL": "graph-extractor-model",
            "A3_GRAPH_LLM_GLEANING_ROUNDS": "0",
        },
    )
    @patch("backend.apps.knowledge_base.graph_builder.chat_completion")
    def test_graph_build_uses_llm_graph_extraction_when_key_is_available(self, chat_completion):
        captured = {}

        def fake_chat(messages, **kwargs):
            captured["messages"] = messages
            captured["kwargs"] = kwargs
            captured["env_model"] = __import__("os").environ.get("A3_LLM_MODEL")
            return {
                "content": """
                {
                  "entities": [
                    {
                      "label": "Root Locus",
                      "aliases": ["RL"],
                      "type": "concept",
                      "summary": "Plots closed-loop pole movement.",
                      "confidence": 0.93,
                      "evidenceBlockIds": []
                    },
                    {
                      "label": "Closed Loop Pole",
                      "aliases": [],
                      "type": "concept",
                      "summary": "Pole of the closed-loop system.",
                      "confidence": 0.9,
                      "evidenceBlockIds": []
                    }
                  ],
                  "relations": [
                    {
                      "source": "Root Locus",
                      "target": "Closed Loop Pole",
                      "type": "used_for",
                      "label": "used for",
                      "confidence": 0.88,
                      "evidenceBlockIds": []
                    }
                  ],
                  "claims": [
                    {
                      "text": "Root Locus is used for Closed Loop Pole.",
                      "entities": ["Root Locus", "Closed Loop Pole"],
                      "evidenceBlockIds": []
                    }
                  ]
                }
                """,
                "provider": "dashscope",
                "model": "qwen-plus",
                "fallback": False,
            }

        chat_completion.side_effect = fake_chat
        resource = self._resource("llm-graph.pdf")
        page = KnowledgePage.objects.create(resource=resource, page_number=1, summary="llm graph")
        KnowledgeParseRun.objects.create(resource=resource, status="completed", parser="native", pages=1)
        KnowledgeBlock.objects.create(
            resource=resource,
            page=page,
            modality="text",
            text="Root Locus is used for Closed Loop Pole.",
            confidence=0.84,
        )

        result = rebuild_resource_graph(resource)

        self.assertEqual(result.status, "ready")
        self.assertFalse(result.degraded)
        run = AgentRun.objects.filter(user=self.user, workflow="material_knowledge_graph").order_by("-started_at").first()
        self.assertTrue(run.steps.filter(name="LLMGraphExtractionAgent", status="completed").exists())
        self.assertTrue(KnowledgeGraphEntity.objects.filter(resource=resource, extraction_method="llm_graph", degraded=False).exists())
        relation = KnowledgeGraphRelation.objects.get(resource=resource, extraction_method="llm_graph")
        self.assertEqual(relation.label, "used for")
        entity = KnowledgeGraphEntity.objects.get(resource=resource, label="Root Locus")
        self.assertEqual(entity.aliases, ["RL"])
        parse_run = KnowledgeParseRun.objects.filter(resource=resource).latest("created_at")
        self.assertEqual(parse_run.result["materialGraph"]["method"], "llm_graph")
        self.assertFalse(parse_run.result["materialGraph"]["degraded"])
        prompt_text = "\n".join(str(message["content"]) for message in captured["messages"])
        self.assertIn("chunk", prompt_text.lower())
        self.assertIn("aliases", prompt_text)
        self.assertIn("evidenceBlockIds", prompt_text)
        self.assertIn("entity_types", prompt_text)
        self.assertIn("relation_types", prompt_text)
        self.assertIn("same or adjacent source blocks", prompt_text)
        self.assertEqual(captured["kwargs"]["response_format"], {"type": "json_object"})
        self.assertEqual(captured["env_model"], "graph-extractor-model")

    @patch.dict(
        "os.environ",
        {
            "A3_LLM_API_KEY": "test-key",
            "A3_GRAPH_LLM_PROVIDER": "dashscope",
            "A3_GRAPH_LLM_GLEANING_ROUNDS": "0",
        },
    )
    @patch("backend.apps.knowledge_base.graph_builder.chat_completion")
    def test_graph_llm_merges_aliases_and_accumulates_chunk_evidence(self, chat_completion):
        chat_completion.return_value = {
            "content": """
            {
              "entities": [
                {
                  "label": "Root Locus",
                  "aliases": ["RL"],
                  "type": "concept",
                  "summary": "Plots closed-loop pole movement.",
                  "confidence": 0.91,
                  "evidenceBlockIds": ["1"]
                },
                {
                  "label": "RL",
                  "aliases": ["Root Locus"],
                  "type": "concept",
                  "summary": "Abbreviation of Root Locus.",
                  "confidence": 0.89,
                  "evidenceBlockIds": ["2"]
                },
                {
                  "label": "Closed Loop Pole",
                  "aliases": ["CL Pole"],
                  "type": "concept",
                  "summary": "Pole of the closed-loop system.",
                  "confidence": 0.9,
                  "evidenceBlockIds": ["2"]
                }
              ],
              "relations": [
                {
                  "source": "RL",
                  "target": "Closed Loop Pole",
                  "type": "influences",
                  "label": "affects",
                  "confidence": 0.86,
                  "evidenceBlockIds": ["2"]
                }
              ],
              "claims": []
            }
            """,
            "provider": "dashscope",
            "model": "qwen-plus",
            "fallback": False,
        }
        resource = self._resource("llm-merge.pdf")
        page = KnowledgePage.objects.create(resource=resource, page_number=1, summary="llm merge")
        KnowledgeParseRun.objects.create(resource=resource, status="completed", parser="native", pages=1)
        KnowledgeBlock.objects.create(
            resource=resource,
            page=page,
            modality="text",
            text="Root Locus plots closed-loop pole movement.",
            confidence=0.84,
        )
        KnowledgeBlock.objects.create(
            resource=resource,
            page=page,
            modality="text",
            text="RL affects Closed Loop Pole stability.",
            confidence=0.86,
        )

        result = rebuild_resource_graph(resource)

        self.assertEqual(result.method, "llm_graph")
        root_locus_entities = KnowledgeGraphEntity.objects.filter(resource=resource, normalized_label="root locus")
        self.assertEqual(root_locus_entities.count(), 1)
        root_locus = root_locus_entities.get()
        self.assertEqual(root_locus.aliases, ["RL"])
        self.assertEqual(root_locus.source_count, 2)
        relation = KnowledgeGraphRelation.objects.get(resource=resource, relation_type="influences")
        self.assertEqual(relation.source_entity.normalized_label, "root locus")
        self.assertEqual(relation.target_entity.normalized_label, "closed loop pole")
        parse_run = KnowledgeParseRun.objects.filter(resource=resource).latest("created_at")
        quality = parse_run.result["materialGraph"]["quality"]
        self.assertEqual(quality["aliasMergeCount"], 1)
        self.assertEqual(quality["isolatedEntityRate"], 0.0)
        self.assertEqual(quality["evidenceCoverage"], 1.0)
        self.assertGreater(quality["relationDensity"], 0)

    @patch.dict(
        "os.environ",
        {
            "A3_LLM_API_KEY": "test-key",
            "A3_GRAPH_LLM_PROVIDER": "dashscope",
            "A3_GRAPH_LLM_GLEANING_ROUNDS": "1",
        },
    )
    @patch("backend.apps.knowledge_base.graph_builder.chat_completion")
    def test_graph_llm_gleaning_rounds_add_missing_relations(self, chat_completion):
        chat_completion.side_effect = [
            {
                "content": """
                {
                  "entities": [
                    {"label": "Root Locus", "aliases": [], "type": "concept", "summary": "Plots pole movement.", "confidence": 0.91, "evidenceBlockIds": ["1"]},
                    {"label": "Gain Margin", "aliases": [], "type": "metric", "summary": "Stability margin.", "confidence": 0.9, "evidenceBlockIds": ["1"]}
                  ],
                  "relations": [],
                  "claims": []
                }
                """,
                "provider": "dashscope",
                "model": "qwen-plus",
            },
            {
                "content": """
                {
                  "entities": [],
                  "relations": [
                    {"source": "Root Locus", "target": "Gain Margin", "type": "influences", "label": "affects", "confidence": 0.86, "evidenceBlockIds": ["1"]}
                  ],
                  "claims": [
                    {"text": "Root Locus affects Gain Margin.", "entities": ["Root Locus", "Gain Margin"], "evidenceBlockIds": ["1"]}
                  ]
                }
                """,
                "provider": "dashscope",
                "model": "qwen-plus",
            },
            {
                "content": """
                {
                  "entities": [],
                  "relations": [],
                  "claims": []
                }
                """,
                "provider": "dashscope",
                "model": "qwen-plus",
            },
        ]
        resource = self._resource("llm-gleaning.pdf")
        page = KnowledgePage.objects.create(resource=resource, page_number=1, summary="llm gleaning")
        KnowledgeParseRun.objects.create(resource=resource, status="completed", parser="native", pages=1)
        KnowledgeBlock.objects.create(
            resource=resource,
            page=page,
            modality="text",
            text="Root Locus affects Gain Margin.",
            confidence=0.87,
        )

        result = rebuild_resource_graph(resource)

        self.assertEqual(result.method, "llm_graph")
        self.assertEqual(chat_completion.call_count, 2)
        second_prompt = "\n".join(str(message["content"]) for message in chat_completion.call_args_list[1].args[0])
        self.assertIn("gleaning", second_prompt.lower())
        self.assertTrue(KnowledgeGraphRelation.objects.filter(resource=resource, relation_type="influences").exists())
        parse_run = KnowledgeParseRun.objects.filter(resource=resource).latest("created_at")
        self.assertEqual(parse_run.result["materialGraph"]["quality"]["gleaningRounds"], 1)

    @patch.dict(
        "os.environ",
        {
            "A3_LLM_API_KEY": "test-key",
            "A3_GRAPH_LLM_PROVIDER": "dashscope",
            "A3_GRAPH_LLM_GLEANING_ROUNDS": "",
        },
    )
    @patch("backend.apps.knowledge_base.graph_builder.chat_completion")
    def test_graph_llm_default_gleaning_repairs_zero_relation_extraction(self, chat_completion):
        chat_completion.side_effect = [
            {
                "content": """
                {
                  "entities": [
                    {"label": "作业", "aliases": [], "type": "concept", "summary": "用户提交给系统的任务实体。", "confidence": 0.91, "evidenceBlockIds": ["1"]},
                    {"label": "进程", "aliases": [], "type": "process", "summary": "程序的一次执行过程。", "confidence": 0.9, "evidenceBlockIds": ["1"]}
                  ],
                  "relations": [],
                  "claims": []
                }
                """,
                "provider": "dashscope",
                "model": "qwen-plus",
            },
            {
                "content": """
                {
                  "entities": [],
                  "relations": [
                    {"source": "作业", "target": "进程", "type": "depends_on", "label": "转换为", "confidence": 0.88, "evidenceBlockIds": ["1"]}
                  ],
                  "claims": []
                }
                """,
                "provider": "dashscope",
                "model": "qwen-plus",
            },
        ]
        resource = self._resource("llm-zero-relations.pptx", file_type=".pptx")
        page = KnowledgePage.objects.create(resource=resource, slide_number=5, summary="作业和进程")
        KnowledgeParseRun.objects.create(resource=resource, status="completed", parser="native", pages=1)
        KnowledgeBlock.objects.create(
            resource=resource,
            page=page,
            modality="text",
            text="作业由程序、数据和作业说明书组成；系统调度后可转换为进程。",
            confidence=0.88,
        )

        result = rebuild_resource_graph(resource)

        self.assertEqual(result.method, "llm_graph")
        self.assertEqual(chat_completion.call_count, 3)
        self.assertGreater(KnowledgeGraphRelation.objects.filter(resource=resource).count(), 0)
        parse_run = KnowledgeParseRun.objects.filter(resource=resource).latest("created_at")
        self.assertEqual(parse_run.result["materialGraph"]["quality"]["gleaningRounds"], 1)
        self.assertLess(parse_run.result["materialGraph"]["quality"]["isolatedEntityRate"], 1.0)

    @patch.dict(
        "os.environ",
        {
            "A3_LLM_API_KEY": "test-key",
            "A3_GRAPH_LLM_PROVIDER": "dashscope",
            "A3_GRAPH_LLM_GLEANING_ROUNDS": "0",
        },
    )
    @patch("backend.apps.knowledge_base.graph_builder.chat_completion")
    def test_graph_llm_processes_all_material_blocks_in_batches(self, chat_completion):
        def response_for_call(messages, **kwargs):
            prompt = "\n".join(str(message["content"]) for message in messages)
            entity = "前部概念" if "front block" in prompt else "后部概念"
            evidence = "1" if entity == "前部概念" else "chunk-26"
            return {
                "content": f"""
                {{
                  "entities": [
                    {{"label": "{entity}", "aliases": [], "type": "concept", "summary": "{entity}", "confidence": 0.9, "evidenceBlockIds": ["{evidence}"]}}
                  ],
                  "relations": [],
                  "claims": []
                }}
                """,
                "provider": "dashscope",
                "model": "qwen-plus",
            }

        chat_completion.side_effect = response_for_call
        resource = self._resource("llm-all-blocks.pdf")
        page = KnowledgePage.objects.create(resource=resource, page_number=1, summary="many blocks")
        KnowledgeParseRun.objects.create(resource=resource, status="completed", parser="native", pages=1)
        for index in range(30):
            text = "front block" if index == 0 else ("tail block unique concept" if index == 25 else f"middle block {index}")
            KnowledgeBlock.objects.create(resource=resource, page=page, modality="text", text=text, confidence=0.8)

        rebuild_resource_graph(resource)

        labels = set(KnowledgeGraphEntity.objects.filter(resource=resource).values_list("label", flat=True))
        self.assertIn("前部概念", labels)
        self.assertIn("后部概念", labels)
        self.assertGreaterEqual(chat_completion.call_count, 2)

    @patch.dict(
        "os.environ",
        {
            "A3_LLM_API_KEY": "test-key",
            "A3_GRAPH_LLM_PROVIDER": "dashscope",
            "A3_GRAPH_LLM_GLEANING_ROUNDS": "0",
        },
    )
    @patch("backend.apps.knowledge_base.graph_builder.chat_completion")
    def test_graph_llm_repairs_non_json_model_response_before_fallback(self, chat_completion):
        chat_completion.side_effect = [
            {
                "content": "I extracted the graph, but here is an explanation instead of JSON.",
                "provider": "dashscope",
                "model": "qwen-plus",
                "fallback": False,
            },
            {
                "content": """
                {
                  "entities": [
                    {"label": "Root Locus", "type": "concept", "summary": "Plots closed-loop pole movement.", "confidence": 0.92},
                    {"label": "Closed Loop Pole", "type": "concept", "summary": "Closed-loop system pole.", "confidence": 0.9}
                  ],
                  "relations": [
                    {"source": "Root Locus", "target": "Closed Loop Pole", "type": "used_for", "label": "used for", "confidence": 0.87}
                  ],
                  "claims": []
                }
                """,
                "provider": "dashscope",
                "model": "qwen-plus",
                "fallback": False,
            },
        ]
        resource = self._resource("llm-repair.pdf")
        page = KnowledgePage.objects.create(resource=resource, page_number=1, summary="llm repair")
        KnowledgeParseRun.objects.create(resource=resource, status="completed", parser="native", pages=1)
        KnowledgeBlock.objects.create(
            resource=resource,
            page=page,
            modality="text",
            text="Root Locus is used for Closed Loop Pole.",
            confidence=0.84,
        )

        result = rebuild_resource_graph(resource)

        self.assertFalse(result.degraded)
        self.assertEqual(result.method, "llm_graph")
        self.assertEqual(chat_completion.call_count, 2)
        self.assertTrue(KnowledgeGraphRelation.objects.filter(resource=resource, extraction_method="llm_graph").exists())

    @patch.dict(
        "os.environ",
        {
            "A3_LLM_API_KEY": "test-key",
            "A3_GRAPH_LLM_PROVIDER": "dashscope",
            "A3_GRAPH_LLM_GLEANING_ROUNDS": "0",
        },
    )
    @patch("backend.apps.knowledge_base.graph_builder.chat_completion")
    def test_graph_llm_normalizes_common_non_schema_json_shapes(self, chat_completion):
        chat_completion.return_value = {
            "content": """
            {
              "graph": {
                "nodes": [
                  {"name": "作业", "entityType": "concept", "description": "用户提交给系统的任务实体。", "confidence": 0.91, "evidence": ["chunk-1"]},
                  {"name": "进程", "entityType": "process", "description": "程序的一次执行过程。", "confidence": 0.9, "evidence": ["chunk-1"]}
                ],
                "edges": [
                  {"from": "作业", "to": "进程", "relation": "depends_on", "description": "调度后转换为", "confidence": 0.88, "evidence": ["chunk-1"]}
                ],
                "facts": [
                  {"claim": "作业经调度后可转换为进程。", "entityNames": ["作业", "进程"], "evidence": ["chunk-1"]}
                ]
              }
            }
            """,
            "provider": "dashscope",
            "model": "qwen-plus",
            "fallback": False,
        }
        resource = self._resource("llm-common-shape.pptx", file_type=".pptx")
        page = KnowledgePage.objects.create(resource=resource, slide_number=5, summary="作业和进程")
        KnowledgeParseRun.objects.create(resource=resource, status="completed", parser="native", pages=1)
        KnowledgeBlock.objects.create(
            resource=resource,
            page=page,
            modality="text",
            text="作业由程序、数据和作业说明书组成；系统调度后可转换为进程。",
            confidence=0.88,
        )

        result = rebuild_resource_graph(resource)

        self.assertFalse(result.degraded)
        self.assertEqual(result.method, "llm_graph")
        self.assertTrue(KnowledgeGraphEntity.objects.filter(resource=resource, label="作业").exists())
        self.assertTrue(KnowledgeGraphRelation.objects.filter(resource=resource, extraction_method="llm_graph").exists())
        self.assertTrue(KnowledgeGraphRelation.objects.filter(resource=resource, label="调度后转换为").exists())
        self.assertTrue(KnowledgeGraphRelation.objects.filter(resource=resource, target_entity__label="进程").exists())

    @patch.dict(
        "os.environ",
        {
            "A3_LLM_API_KEY": "test-key",
            "A3_GRAPH_LLM_PROVIDER": "dashscope",
            "A3_GRAPH_LLM_GLEANING_ROUNDS": "0",
        },
    )
    @patch("backend.apps.knowledge_base.graph_builder.chat_completion")
    def test_graph_llm_keeps_successful_batches_when_one_batch_is_unrepairable(self, chat_completion):
        def response_for_call(messages, **kwargs):
            prompt = "\n".join(str(message["content"]) for message in messages)
            if "front invalid block" in prompt:
                return {"content": "not json and not repairable", "provider": "dashscope", "model": "qwen-plus"}
            return {
                "content": """
                {
                  "entities": [
                    {"label": "后部概念", "aliases": [], "type": "concept", "summary": "后部块中的概念。", "confidence": 0.9, "evidenceBlockIds": ["chunk-26"]},
                    {"label": "后部方法", "aliases": [], "type": "method", "summary": "后部块中的方法。", "confidence": 0.88, "evidenceBlockIds": ["chunk-26"]}
                  ],
                  "relations": [
                    {"source": "后部方法", "target": "后部概念", "type": "used_for", "label": "用于", "confidence": 0.86, "evidenceBlockIds": ["chunk-26"]}
                  ],
                  "claims": []
                }
                """,
                "provider": "dashscope",
                "model": "qwen-plus",
            }

        chat_completion.side_effect = response_for_call
        resource = self._resource("llm-partial-batches.pdf")
        page = KnowledgePage.objects.create(resource=resource, page_number=1, summary="partial batches")
        KnowledgeParseRun.objects.create(resource=resource, status="completed", parser="native", pages=1)
        for index in range(30):
            text = "front invalid block" if index == 0 else ("tail valid block" if index == 25 else f"middle block {index}")
            KnowledgeBlock.objects.create(resource=resource, page=page, modality="text", text=text, confidence=0.8)

        result = rebuild_resource_graph(resource)

        self.assertFalse(result.degraded)
        self.assertEqual(result.method, "llm_graph")
        self.assertTrue(KnowledgeGraphEntity.objects.filter(resource=resource, label="后部概念").exists())
        self.assertTrue(KnowledgeGraphRelation.objects.filter(resource=resource, extraction_method="llm_graph").exists())
        parse_run = KnowledgeParseRun.objects.filter(resource=resource).latest("created_at")
        quality = parse_run.result["materialGraph"]["quality"]
        self.assertEqual(quality["processedBatchCount"], 2)
        self.assertEqual(quality["successfulBatchCount"], 1)
        self.assertEqual(quality["failedBatchCount"], 1)

    @patch.dict("os.environ", {"A3_GRAPH_LLM_API_KEY": "", "A3_LLM_API_KEY": "", "DASHSCOPE_API_KEY": "", "OPENAI_API_KEY": ""})
    def test_graph_build_keeps_rule_fallback_when_llm_key_is_missing(self):
        resource = self._resource("no-llm-graph.pdf")
        page = KnowledgePage.objects.create(resource=resource, page_number=1, summary="fallback graph")
        KnowledgeParseRun.objects.create(resource=resource, status="completed", parser="native", pages=1)
        KnowledgeBlock.objects.create(
            resource=resource,
            page=page,
            modality="text",
            text="Root Locus is used for Closed Loop Pole.",
            confidence=0.84,
        )

        result = rebuild_resource_graph(resource)

        self.assertEqual(result.status, "ready")
        self.assertTrue(result.degraded)
        run = AgentRun.objects.filter(user=self.user, workflow="material_knowledge_graph").order_by("-started_at").first()
        self.assertTrue(run.steps.filter(name="LLMGraphExtractionAgent", status="degraded").exists())
        parse_run = KnowledgeParseRun.objects.filter(resource=resource).latest("created_at")
        self.assertEqual(parse_run.result["materialGraph"]["method"], "rule_fallback")

    @patch.dict("os.environ", {"A3_GRAPH_LLM_API_KEY": "", "A3_LLM_API_KEY": "", "DASHSCOPE_API_KEY": "", "OPENAI_API_KEY": ""})
    def test_image_caption_confidence_is_preserved_in_graph_evidence(self):
        resource = self._resource("caption-evidence.pptx", file_type=".pptx")
        page = KnowledgePage.objects.create(resource=resource, slide_number=2, summary="caption evidence")
        KnowledgeParseRun.objects.create(resource=resource, status="completed", parser="native", pages=1)
        KnowledgeBlock.objects.create(
            resource=resource,
            page=page,
            modality="image",
            text="Root Locus diagram shows Closed Loop Pole movement.",
            confidence=0.78,
            metadata={
                "blockType": "image_caption",
                "captionMethod": "vision_provider",
                "captionConfidence": 0.78,
                "captionProvider": "dashscope",
            },
        )

        rebuild_resource_graph(resource)

        evidence_rows = [
            evidence
            for entity in KnowledgeGraphEntity.objects.filter(resource=resource)
            for evidence in entity.source_locations
        ]
        self.assertTrue(any(row.get("captionConfidence") == 0.78 for row in evidence_rows))
        self.assertTrue(any(row.get("captionMethod") == "vision_provider" for row in evidence_rows))

    def test_graph_build_uses_agent_kernel_workflow_metadata(self):
        resource = self._resource("kernel.pdf")
        page = KnowledgePage.objects.create(resource=resource, page_number=1, summary="kernel")
        KnowledgeParseRun.objects.create(resource=resource, status="completed", parser="native", pages=1)
        KnowledgeBlock.objects.create(
            resource=resource,
            page=page,
            modality="text",
            text="Transfer Function 是 Control Model. Root Locus 用于 Closed Loop Pole.",
            confidence=0.83,
        )

        result = rebuild_resource_graph(resource)

        self.assertEqual(result.status, "ready")
        run = AgentRun.objects.filter(user=self.user, workflow="material_knowledge_graph").order_by("-started_at").first()
        self.assertIsNotNone(run)
        self.assertEqual(run.status, "completed")
        self.assertTrue(run.steps.filter(name="GraphExtractionAgent").exists())
        parse_run = KnowledgeParseRun.objects.filter(resource=resource).latest("created_at")
        material_graph = parse_run.result["materialGraph"]
        self.assertEqual(material_graph["workflow"], "material_knowledge_graph")
        self.assertEqual(material_graph["workflowSchemaVersion"], "a3.agent-workflow.v1")
        self.assertTrue(material_graph["traceId"])
        self.assertIn("degraded", material_graph)

    def test_material_graph_workflow_has_formal_contract(self):
        contract = get_agent_workflow_contract("material_knowledge_graph")

        self.assertIsNotNone(contract)
        self.assertEqual(
            contract.required_steps,
            ("GraphInputAgent", "GraphExtractionAgent", "LLMGraphExtractionAgent", "GraphPersistenceAgent"),
        )
        self.assertIn("materialGraph", contract.output_schema["properties"])

    def test_demo_status_exposes_material_graph_readiness(self):
        resource = self._resource("status.pdf")
        page = KnowledgePage.objects.create(resource=resource, page_number=1)
        KnowledgeParseRun.objects.create(resource=resource, status="completed", parser="native", pages=1)
        KnowledgeBlock.objects.create(
            resource=resource,
            page=page,
            modality="text",
            text="Transfer Function 是 Control Model. Root Locus 方程为：s² + 2s + Kg = 0. Closed Loop Pole.",
            confidence=0.82,
            metadata={
                "markdownText": "Transfer Function 是 Control Model. Root Locus 方程为：$s^2 + 2s + Kg = 0$. Closed Loop Pole.",
                "contentFormat": "markdown",
                "formatFeatures": ["math"],
            },
        )
        rebuild_resource_graph(resource)

        response = self.client.get("/api/demo/status/")

        self.assertEqual(response.status_code, 200)
        material_graph = response.json()["materialKnowledgeGraph"]
        self.assertEqual(material_graph["schemaVersion"], "a3.material-knowledge-graph.v1")
        self.assertGreaterEqual(material_graph["entityCount"], 1)
        self.assertGreaterEqual(material_graph["relationCount"], 1)
        self.assertEqual(material_graph["workflow"], "material_knowledge_graph")
        self.assertIn("reason", material_graph)
        self.assertEqual(material_graph["markdownEvidence"]["contentFormat"], "markdown")
        self.assertGreaterEqual(material_graph["markdownEvidence"]["mathCount"], 1)
        self.assertTrue(material_graph["markdownEvidence"]["rawTextPreserved"])

    def _graph_entity(self, resource: KnowledgeResource, label: str) -> KnowledgeGraphEntity:
        return KnowledgeGraphEntity.objects.create(
            owner=resource.owner,
            resource=resource,
            label=label,
            normalized_label=label.lower(),
            entity_type="concept",
            summary=f"{label} summary",
            confidence=0.9,
            source_count=1,
            source_locations=[{"resourceId": str(resource.id), "resourceTitle": resource.original_name, "location": "测试资料", "snippet": label}],
            degraded=False,
        )

    def _graph_relation(self, resource: KnowledgeResource, source: KnowledgeGraphEntity, target: KnowledgeGraphEntity, relation_type: str) -> KnowledgeGraphRelation:
        return KnowledgeGraphRelation.objects.create(
            owner=resource.owner,
            resource=resource,
            source_entity=source,
            target_entity=target,
            relation_type=relation_type,
            label=relation_type,
            weight=0.9,
            evidence=[{"resourceId": str(resource.id), "resourceTitle": resource.original_name, "location": "测试资料", "snippet": relation_type}],
            degraded=False,
        )

    def _resource(self, name: str, *, owner=None, file_type=".pdf", scope=ResourceScope.PERSONAL) -> KnowledgeResource:
        path = Path(self.temp_dir.name) / name
        path.write_text("demo", encoding="utf-8")
        return KnowledgeResource.objects.create(
            owner=owner or self.user,
            original_name=name,
            stored_path=str(path.resolve()),
            file_type=file_type,
            file_size=path.stat().st_size,
            resource_scope=scope,
            status=ResourceStatus.INDEXED,
        )
