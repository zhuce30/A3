from __future__ import annotations

import hashlib
import math
import re
import shutil
from datetime import date, datetime, time
from pathlib import Path

from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils import timezone

from backend.apps.accounts.models import UserRole
from backend.apps.knowledge_base.models import (
    CuratedLearningResource,
    CuratedResourceDeliveryMode,
    CuratedResourceDifficulty,
    CuratedResourceType,
    CuratedResourceVerificationStatus,
    KnowledgeResource,
    ResourceScope,
    ResourceStatus,
)
from backend.apps.knowledge_base.services import create_upload_index_job
from backend.apps.learning.models import KnowledgePoint
from backend.apps.learning.services import DEFAULT_COURSE_CODE, _repair_default_course_catalog
from backend.core.settings import DATA_ROOT


VIDEO_CATALOG_NAME = "操作系统免费视频资源_12个精准推荐视频.xlsx"
VIDEO_SHEET_NAME = "视频资源清单"
EXPECTED_DOCUMENT_CODES = {
    "OS-DOC-001": ("OS-PROCESS", "OS-THREAD", "OS-SCHEDULING"),
    "OS-DOC-002": ("OS-MEMORY", "OS-VIRTUAL-MEMORY"),
    "OS-DOC-003": ("OS-SYNC", "OS-DEADLOCK"),
    "OS-LAB-001": ("OS-PROCESS", "OS-SCHEDULING"),
    "OS-LAB-002": ("OS-VIRTUAL-MEMORY",),
    "OS-LAB-003": ("OS-SYNC",),
    "OS-001-LD-02-01": ("OS-PROCESS",),
}
REQUIRED_METADATA_FIELDS = {
    "资源编号",
    "课程编号",
    "资源类型",
    "核心知识点",
    "难度",
    "适用学生画像",
    "薄弱点触发条件",
    "预计学习时间",
    "学习目标",
    "完成判据",
    "推荐理由模板",
    "检索关键词",
}
VIDEO_COLUMNS = (
    "主题",
    "视频标题",
    "平台",
    "发布者/机构",
    "真实URL",
    "总时长",
    "关键知识点时间段",
    "难度",
    "适用学生画像",
    "薄弱点触发条件",
    "推荐理由",
    "是否免费",
    "核验日期",
)


class Command(BaseCommand):
    help = "将本地操作系统讲解文档、实操案例和视频清单导入 A3 结构化资源目录。"

    def add_arguments(self, parser):
        parser.add_argument("--source-dir", default=str(DATA_ROOT / "操作系统" / "资料"))
        parser.add_argument("--course-code", default=DEFAULT_COURSE_CODE)
        parser.add_argument("--owner-phone", default="13800000002")
        parser.add_argument("--scope", choices=[ResourceScope.GLOBAL, ResourceScope.TEACHER], default=ResourceScope.GLOBAL)
        parser.add_argument("--enqueue-index", action="store_true")

    def handle(self, *args, **options):
        source_dir = Path(options["source_dir"]).resolve()
        if not source_dir.is_dir():
            raise CommandError(f"资源目录不存在：{source_dir}")

        video_catalog = source_dir / VIDEO_CATALOG_NAME
        if not video_catalog.is_file():
            raise CommandError(f"缺少视频资源清单：{video_catalog.name}")

        document_rows = self._read_document_rows(source_dir)
        video_rows = _read_video_rows(video_catalog)
        if len(video_rows) != 12:
            raise CommandError(f"视频资源清单应包含 12 条记录，当前为 {len(video_rows)} 条")

        course = _repair_default_course_catalog()
        if course.code != options["course_code"]:
            raise CommandError(f"当前默认课程编号为 {course.code}，与参数 {options['course_code']} 不一致")
        points = {
            point.code: point
            for point in KnowledgePoint.objects.filter(chapter__course=course).select_related("chapter")
        }
        owner = _material_owner(options["owner_phone"])
        target_root = Path(getattr(settings, "KNOWLEDGE_UPLOAD_ROOT", DATA_ROOT / "uploads")).resolve()
        target_root = target_root / "course-materials" / course.code / "curated"
        target_root.mkdir(parents=True, exist_ok=True)

        counters = {"created": 0, "updated": 0, "unchanged": 0}
        resources_to_index: list[KnowledgeResource] = []
        with transaction.atomic():
            for source_path, metadata in document_rows:
                catalog, state, needs_index = _upsert_document(
                    owner=owner,
                    course=course,
                    points=points,
                    source_path=source_path,
                    target_root=target_root,
                    metadata=metadata,
                    resource_scope=options["scope"],
                )
                counters[state] += 1
                if needs_index and catalog.knowledge_resource:
                    resources_to_index.append(catalog.knowledge_resource)

            for index, row in enumerate(video_rows, start=1):
                _, state = _upsert_video(course=course, points=points, index=index, row=row)
                counters[state] += 1

        job = None
        if options["enqueue_index"] and resources_to_index:
            job = create_upload_index_job(owner, resources_to_index)

        summary = (
            f"操作系统精准资源导入完成：新增 {counters['created']}，更新 {counters['updated']}，"
            f"未变化 {counters['unchanged']}，待索引 {len(resources_to_index)}"
        )
        if job:
            summary += f"，索引任务 {job.id}"
        self.stdout.write(self.style.SUCCESS(summary))

    def _read_document_rows(self, source_dir: Path) -> list[tuple[Path, dict[str, str]]]:
        rows: list[tuple[Path, dict[str, str]]] = []
        codes: set[str] = set()
        for source_path in sorted(source_dir.glob("*.docx")):
            metadata = _read_docx_metadata(source_path)
            missing_fields = sorted(REQUIRED_METADATA_FIELDS - set(metadata))
            if missing_fields:
                raise CommandError(f"{source_path.name} 缺少元数据字段：{'、'.join(missing_fields)}")
            resource_code = metadata["资源编号"].strip()
            if resource_code not in EXPECTED_DOCUMENT_CODES:
                raise CommandError(f"{source_path.name} 使用了未登记的资源编号：{resource_code}")
            if resource_code in codes:
                raise CommandError(f"资源编号重复：{resource_code}")
            if metadata["课程编号"].strip() != DEFAULT_COURSE_CODE:
                raise CommandError(f"{resource_code} 的课程编号不是 {DEFAULT_COURSE_CODE}")
            codes.add(resource_code)
            rows.append((source_path, metadata))

        missing_codes = sorted(set(EXPECTED_DOCUMENT_CODES) - codes)
        if missing_codes:
            raise CommandError(f"缺少操作系统资源文档：{'、'.join(missing_codes)}")
        return rows


def _upsert_document(*, owner, course, points: dict, source_path: Path, target_root: Path, metadata: dict, resource_scope: str):
    resource_code = metadata["资源编号"].strip()
    point_codes = EXPECTED_DOCUMENT_CODES[resource_code]
    linked_points = _required_points(points, point_codes, resource_code)
    content_hash = _sha256(source_path)
    target_path = target_root / source_path.name
    if not target_path.exists() or _sha256(target_path) != content_hash:
        shutil.copy2(source_path, target_path)

    existing = CuratedLearningResource.objects.filter(resource_code=resource_code).select_related("knowledge_resource").first()
    knowledge_resource = existing.knowledge_resource if existing else None
    if knowledge_resource is None:
        knowledge_resource = KnowledgeResource.objects.filter(
            course_key=course.code,
            resource_scope=resource_scope,
            student_profile_snapshot__curatedResourceCode=resource_code,
        ).first()

    knowledge_defaults = {
        "owner": owner,
        "original_name": source_path.name,
        "stored_path": str(target_path),
        "file_type": ".docx",
        "file_size": target_path.stat().st_size,
        "resource_scope": resource_scope,
        "course_key": course.code,
        "visibility": "shared",
        "student_profile_snapshot": {
            "source": "curated_os_import",
            "curatedResourceCode": resource_code,
            "courseCode": course.code,
            "materialRole": "code_lab" if "实操" in metadata["资源类型"] else "lecture_doc",
        },
    }
    knowledge_created = knowledge_resource is None
    knowledge_changed = knowledge_created or any(getattr(knowledge_resource, key) != value for key, value in knowledge_defaults.items())
    if knowledge_created:
        knowledge_resource = KnowledgeResource.objects.create(**knowledge_defaults, status=ResourceStatus.UPLOADED)
    elif knowledge_changed or existing is None or existing.content_hash != content_hash:
        for key, value in knowledge_defaults.items():
            setattr(knowledge_resource, key, value)
        knowledge_resource.status = ResourceStatus.UPLOADED
        knowledge_resource.index_job = None
        knowledge_resource.document_count = 0
        knowledge_resource.chunk_count = 0
        knowledge_resource.error_message = ""
        knowledge_resource.save()

    resource_type = CuratedResourceType.CODE_LAB if "实操" in metadata["资源类型"] else CuratedResourceType.LECTURE_DOC
    defaults = {
        "course": course,
        "knowledge_resource": knowledge_resource,
        "resource_type": resource_type,
        "delivery_mode": CuratedResourceDeliveryMode.INTERNAL_DOCUMENT,
        "title": metadata.get("标题") or source_path.stem,
        "summary": metadata.get("核心知识点", ""),
        "external_url": "",
        "provider": "A3 本地知识库",
        "difficulty": _difficulty(metadata.get("难度")),
        "estimated_minutes": _estimated_minutes(metadata.get("预计学习时间")),
        "key_segment": "",
        "prerequisites": _split_values(metadata.get("前置知识")),
        "learner_profiles": _split_values(metadata.get("适用学生画像")),
        "weakness_triggers": _split_values(metadata.get("薄弱点触发条件")),
        "learning_objectives": _split_values(metadata.get("学习目标")),
        "completion_criteria": _split_values(metadata.get("完成判据")),
        "recommendation_template": metadata.get("推荐理由模板", ""),
        "keywords": _split_values(metadata.get("检索关键词")),
        "is_free": True,
        "verification_status": CuratedResourceVerificationStatus.VERIFIED,
        "verified_at": date.fromtimestamp(source_path.stat().st_mtime),
        "content_hash": content_hash,
        "metadata": {"sourceMetadata": metadata, "sourceFile": source_path.name},
        "is_active": True,
    }
    catalog, state = _upsert_catalog(resource_code, defaults, linked_points)
    needs_index = (
        knowledge_created
        or knowledge_changed
        or existing is None
        or existing.content_hash != content_hash
        or knowledge_resource.status in {ResourceStatus.UPLOADED, ResourceStatus.FAILED}
    )
    return catalog, state, needs_index


def _upsert_video(*, course, points: dict, index: int, row: dict[str, object]):
    resource_code = f"OS-VIDEO-{index:03d}"
    point_codes = _video_point_codes(str(row["主题"]), str(row["视频标题"]))
    linked_points = _required_points(points, point_codes, resource_code)
    verified_at = _date_value(row.get("核验日期"))
    defaults = {
        "course": course,
        "knowledge_resource": None,
        "resource_type": CuratedResourceType.VIDEO,
        "delivery_mode": CuratedResourceDeliveryMode.EXTERNAL_LINK,
        "title": str(row["视频标题"]).strip(),
        "summary": str(row["推荐理由"]).strip(),
        "external_url": str(row["真实URL"]).strip(),
        "provider": " · ".join(filter(None, [str(row["平台"]).strip(), str(row["发布者/机构"]).strip()])),
        "difficulty": _difficulty(row.get("难度")),
        "estimated_minutes": _video_minutes(row.get("总时长")),
        "key_segment": str(row["关键知识点时间段"]).strip(),
        "prerequisites": [],
        "learner_profiles": _split_values(row.get("适用学生画像")),
        "weakness_triggers": _split_values(row.get("薄弱点触发条件")),
        "learning_objectives": [],
        "completion_criteria": [],
        "recommendation_template": str(row["推荐理由"]).strip(),
        "keywords": list(dict.fromkeys([str(row["主题"]).strip(), str(row["视频标题"]).strip()])),
        "is_free": str(row["是否免费"]).strip().startswith("是"),
        "verification_status": CuratedResourceVerificationStatus.VERIFIED,
        "verified_at": verified_at,
        "content_hash": hashlib.sha256(str(row["真实URL"]).strip().encode("utf-8")).hexdigest(),
        "metadata": {"source": "curated_os_video_catalog", "catalogRow": index + 1, "raw": _json_safe_row(row)},
        "is_active": True,
    }
    return _upsert_catalog(resource_code, defaults, linked_points)


def _upsert_catalog(resource_code: str, defaults: dict, linked_points: list):
    existing = CuratedLearningResource.objects.filter(resource_code=resource_code).first()
    existing_point_ids = set(existing.knowledge_points.values_list("id", flat=True)) if existing else set()
    desired_point_ids = {point.id for point in linked_points}
    changed = existing is None or existing_point_ids != desired_point_ids or any(
        getattr(existing, key) != value for key, value in defaults.items()
    )
    catalog, created = CuratedLearningResource.objects.update_or_create(resource_code=resource_code, defaults=defaults)
    catalog.knowledge_points.set(linked_points)
    if created:
        return catalog, "created"
    return catalog, "updated" if changed else "unchanged"


def _read_docx_metadata(path: Path) -> dict[str, str]:
    try:
        from docx import Document
    except ModuleNotFoundError as exc:
        raise CommandError("缺少 python-docx，无法读取 DOCX 元数据") from exc

    document = Document(str(path))
    if not document.tables:
        raise CommandError(f"{path.name} 没有资源元数据表")
    metadata: dict[str, str] = {"标题": _document_title(document, path)}
    for row in document.tables[0].rows:
        values = [_clean_cell(cell.text) for cell in row.cells]
        for index in range(0, len(values) - 1, 2):
            key = values[index]
            value = values[index + 1]
            if key in REQUIRED_METADATA_FIELDS or key in {"所属章节"}:
                metadata[key] = value
        if len(values) == 2 and values[0]:
            metadata[values[0]] = values[1]
    return metadata


def _read_video_rows(path: Path) -> list[dict[str, object]]:
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as exc:
        raise CommandError("缺少 openpyxl，无法读取视频资源清单") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    if VIDEO_SHEET_NAME not in workbook.sheetnames:
        workbook.close()
        raise CommandError(f"视频清单缺少工作表：{VIDEO_SHEET_NAME}")
    worksheet = workbook[VIDEO_SHEET_NAME]
    rows = worksheet.iter_rows(values_only=True)
    headers = [_clean_cell(value) for value in next(rows, [])]
    missing = [column for column in VIDEO_COLUMNS if column not in headers]
    if missing:
        workbook.close()
        raise CommandError(f"视频清单缺少字段：{'、'.join(missing)}")
    indices = {column: headers.index(column) for column in VIDEO_COLUMNS}
    payload: list[dict[str, object]] = []
    urls: set[str] = set()
    for values in rows:
        row = {column: values[index] if index < len(values) else None for column, index in indices.items()}
        url = _clean_cell(row["真实URL"])
        title = _clean_cell(row["视频标题"])
        if not url and not title:
            continue
        if not url.startswith("https://"):
            workbook.close()
            raise CommandError(f"视频链接必须使用 HTTPS：{title or url}")
        if url in urls:
            workbook.close()
            raise CommandError(f"视频链接重复：{url}")
        urls.add(url)
        payload.append(row)
    workbook.close()
    return payload


def _material_owner(phone: str):
    User = get_user_model()
    owner, _ = User.objects.get_or_create(
        phone=phone,
        defaults={
            "role": UserRole.TEACHER,
            "display_name": "课程资料管理员",
            "school": "智学引擎 A3",
            "privacy_accepted": True,
            "profile_completed": True,
            "onboarding_step": "complete",
        },
    )
    if owner.role not in {UserRole.TEACHER, UserRole.ADMIN}:
        owner.role = UserRole.TEACHER
        owner.save(update_fields=["role", "updated_at"])
    return owner


def _required_points(points: dict, codes: tuple[str, ...], resource_code: str) -> list:
    missing = [code for code in codes if code not in points]
    if missing:
        raise CommandError(f"{resource_code} 无法映射知识点：{'、'.join(missing)}")
    return [points[code] for code in codes]


def _video_point_codes(topic: str, title: str) -> tuple[str, ...]:
    text = f"{topic} {title}"
    if "线程" in text:
        return ("OS-THREAD",)
    if "进程基本" in text:
        return ("OS-PROCESS",)
    if any(token in text for token in ("CPU调度", "先来先服务", "短作业", "时间片", "多级队列")):
        return ("OS-SCHEDULING",)
    if any(token in text for token in ("信号量", "生产者", "消费者")):
        return ("OS-SYNC",)
    if any(token in text for token in ("死锁", "银行家")):
        return ("OS-DEADLOCK",)
    if "页面置换" in text:
        return ("OS-VIRTUAL-MEMORY",)
    if "按需调页" in text:
        return ("OS-MEMORY", "OS-VIRTUAL-MEMORY")
    if "文件" in text:
        return ("OS-FILE",)
    raise CommandError(f"视频资源无法映射知识点：{title}")


def _difficulty(value: object) -> str:
    text = _clean_cell(value)
    if any(token in text for token in ("基础", "入门")):
        return CuratedResourceDifficulty.FOUNDATION
    if any(token in text for token in ("进阶", "高级", "困难")):
        return CuratedResourceDifficulty.ADVANCED
    return CuratedResourceDifficulty.INTERMEDIATE


def _estimated_minutes(value: object) -> int:
    numbers = [int(item) for item in re.findall(r"\d+", _clean_cell(value))]
    return max(numbers, default=0)


def _video_minutes(value: object) -> int:
    if isinstance(value, time):
        return max(1, math.ceil((value.hour * 3600 + value.minute * 60 + value.second) / 60))
    text = _clean_cell(value)
    match = re.fullmatch(r"(\d+):(\d{1,2})(?::(\d{1,2}))?", text)
    if match:
        parts = [int(item or 0) for item in match.groups()]
        seconds = parts[0] * 60 + parts[1] if match.group(3) is None else parts[0] * 3600 + parts[1] * 60 + parts[2]
        return max(1, math.ceil(seconds / 60))
    return _estimated_minutes(text)


def _date_value(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _clean_cell(value)
    if not text:
        return None
    try:
        return date.fromisoformat(text[:10])
    except ValueError:
        return timezone.localdate()


def _split_values(value: object) -> list[str]:
    return [item.strip() for item in re.split(r"[；;、，,\n]+", _clean_cell(value)) if item.strip()]


def _clean_cell(value: object) -> str:
    return "\n".join(line.strip() for line in str(value or "").replace("\r", "\n").split("\n") if line.strip())


def _document_title(document, path: Path) -> str:
    for paragraph in document.paragraphs:
        text = _clean_cell(paragraph.text)
        if text:
            return text
    return path.stem


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _json_safe_row(row: dict[str, object]) -> dict[str, object]:
    return {
        key: value.isoformat() if isinstance(value, (date, datetime, time)) else value
        for key, value in row.items()
    }
