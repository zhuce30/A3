from __future__ import annotations

import json
import math
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, TYPE_CHECKING

if TYPE_CHECKING:
    from llama_index.core import Document
else:
    Document = object

SUPPORTED_EXTENSIONS = {".docx", ".xlsx", ".pdf", ".pptx"}
SKIPPED_EXTENSIONS: set[str] = set()


@dataclass(frozen=True)
class LexicalDocument:
    text: str
    metadata: dict

    def get_content(self) -> str:
        return self.text


def iter_source_files(data_root: Path, limit_files: int | None = None) -> list[Path]:
    candidates: list[Path] = []
    for path in data_root.rglob("*"):
        if not path.is_file():
            continue
        suffix = path.suffix.lower()
        if suffix in SUPPORTED_EXTENSIONS:
            candidates.append(path)

    files = sorted(candidates, key=_file_priority)
    if limit_files:
        files = files[:limit_files]
    return files


def load_documents(paths: Iterable[Path]) -> tuple[list[Document], list[dict]]:
    documents: list[Document] = []
    skipped: list[dict] = []

    for path in paths:
        suffix = path.suffix.lower()
        try:
            if suffix == ".xlsx":
                documents.extend(_load_qa_xlsx(path))
            elif suffix == ".docx":
                documents.extend(_load_docx(path))
            elif suffix == ".pdf":
                documents.extend(_load_pdf(path))
            elif suffix == ".pptx":
                documents.extend(_load_pptx(path))
            elif suffix == ".ppt":
                skipped.append({
                    "path": str(path),
                    "reason": "Legacy .ppt is not supported by python-pptx. Convert to .pptx first.",
                })
        except Exception as exc:
            skipped.append({"path": str(path), "reason": f"{type(exc).__name__}: {exc}"})

    return documents, skipped


def _document(text: str, metadata: dict) -> Document | LexicalDocument:
    try:
        from llama_index.core import Document
    except ModuleNotFoundError:
        return LexicalDocument(text=text, metadata=metadata)
    return Document(text=text, metadata=metadata)


def _missing_loader_dependency(package: str, feature: str, exc: ModuleNotFoundError) -> RuntimeError:
    return RuntimeError(
        f"{package} is required to {feature}. Install backend requirements before building the knowledge base."
    )


def write_build_report(path: Path, report: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")


def _file_priority(path: Path) -> tuple[int, str]:
    text = str(path)
    if path.name in {"\u6559\u5b66\u95ee\u7b54\u6570\u636e\u96c6.xlsx", "鏁欏闂瓟鏁版嵁闆?xlsx"}:
        return (0, text)
    if "\u6750\u6599\u529b\u5b66" in text or "鏉愭枡鍔涘" in text:
        return (1, text)
    if path.suffix.lower() == ".xlsx":
        return (2, text)
    if path.suffix.lower() in {".docx", ".pdf", ".pptx"}:
        return (3, text)
    return (9, text)


def _load_qa_xlsx(path: Path) -> list[Document]:
    try:
        import pandas as pd
    except ModuleNotFoundError as exc:
        raise _missing_loader_dependency("pandas", "load XLSX question-bank files", exc) from exc

    df = pd.read_excel(path)
    documents: list[Document] = []
    question_col = _find_column(df, ["\u6807\u51c6\u95ee\u9898", "鏍囧噯闂"])
    answer_col = _find_column(df, ["\u7b54\u6848", "绛旀"])
    category_col = _find_column(df, ["\u89c4\u5219\u5206\u7c7b", "瑙勫垯"])
    similar_cols = [
        column
        for column in df.columns
        if any(token in str(column) for token in ["\u76f8\u4f3c\u95ee\u6cd5", "鐩镐技闂"])
    ][:6]

    for idx, row in df.iterrows():
        question = _clean_cell(row.get(question_col)) if question_col is not None else ""
        answer = _clean_cell(row.get(answer_col)) if answer_col is not None else ""
        if not question or not answer:
            continue

        similar_questions = [_clean_cell(row.get(column)) for column in similar_cols]
        similar_questions = [q for q in similar_questions if q]
        category = (_clean_cell(row.get(category_col)) if category_col is not None else "") or "\u9ed8\u8ba4\u5206\u7c7b"

        text = "\n".join([
            f"\u6807\u51c6\u95ee\u9898\uff1a{question}",
            f"\u76f8\u4f3c\u95ee\u6cd5\uff1a{'；'.join(similar_questions) if similar_questions else '\u65e0'}",
            f"\u7b54\u6848\uff1a{answer}",
        ])
        documents.append(
            _document(
                text=text,
                metadata={
                    "source_type": "qa_xlsx",
                    "source_path": str(path),
                    "row": int(idx) + 2,
                    "course": _infer_course(path),
                    "category": category,
                    "title": question,
                    "original_name": path.name,
                    "resource_type": "question_bank",
                },
            )
        )

    return documents


def _find_column(df, tokens: list[str]) -> object | None:
    for column in df.columns:
        text = str(column)
        if any(token in text for token in tokens):
            return column
    return None


def _load_pdf(path: Path) -> list[Document]:
    try:
        from pypdf import PdfReader
    except ModuleNotFoundError as exc:
        raise _missing_loader_dependency("pypdf", "load PDF courseware", exc) from exc

    reader = PdfReader(str(path))
    documents: list[Document] = []

    for page_index, page in enumerate(reader.pages, start=1):
        text = _normalize_text(page.extract_text() or "")
        if not text:
            continue
        documents.append(
            _document(
                text=text,
                metadata={
                    "source_type": "pdf",
                    "source_path": str(path),
                    "page": page_index,
                    "course": _infer_course(path),
                    "title": path.stem,
                    "original_name": path.name,
                    "resource_type": "document",
                },
            )
        )

    return documents


def _load_pptx(path: Path) -> list[Document]:
    try:
        from pptx import Presentation
    except ModuleNotFoundError as exc:
        raise _missing_loader_dependency("python-pptx", "load PPTX courseware", exc) from exc

    presentation = Presentation(str(path))
    documents: list[Document] = []

    for slide_index, slide in enumerate(presentation.slides, start=1):
        text_parts: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                text_parts.append(shape.text)
        text = _normalize_text("\n".join(text_parts))
        if not text:
            continue
        documents.append(
            _document(
                text=text,
                metadata={
                    "source_type": "pptx",
                    "source_path": str(path),
                    "slide": slide_index,
                    "course": _infer_course(path),
                    "title": path.stem,
                    "original_name": path.name,
                    "resource_type": "courseware",
                },
            )
        )

    return documents


def _load_docx(path: Path) -> list[Document]:
    from backend.domains.knowledge_base.docx_parser import parse_docx_blocks

    documents: list[Document] = []
    for block in parse_docx_blocks(path):
        metadata = {
            "source_type": "docx",
            "source_path": str(path),
            "course": _infer_course(path),
            "title": block.section or path.stem,
            "section": block.section,
            "heading_level": block.heading_level,
            "block_index": block.block_index,
            "block_type": block.block_type,
            "paragraph_start": block.paragraph_start,
            "paragraph_end": block.paragraph_end,
            "table_index": block.table_index,
            "original_name": path.name,
            "resource_type": "document",
        }
        documents.append(_document(text=block.markdown, metadata=metadata))
    return documents


def _clean_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    return _normalize_text(str(value))


def _normalize_text(text: str) -> str:
    lines = [line.strip() for line in text.replace("\r", "\n").split("\n")]
    return "\n".join(line for line in lines if line)


def _infer_course(path: Path) -> str:
    parts = list(path.parts)
    if "PPT" in parts:
        idx = parts.index("PPT")
        if idx + 1 < len(parts):
            return parts[idx + 1].replace("-\u8bfe\u4ef620260303", "")
    if path.name.endswith(".xlsx"):
        return "\u6559\u5b66\u95ee\u7b54\u6570\u636e\u96c6"
    return path.parent.name
