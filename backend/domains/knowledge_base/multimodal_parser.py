from __future__ import annotations

import base64
import hashlib
import json
import mimetypes
import os
import shutil
import subprocess
import tempfile
import time
import urllib.error
import urllib.parse
import urllib.request
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path
from typing import Callable

from backend.domains.knowledge_base.document_parse_v2 import finalize_parse_result
from backend.domains.knowledge_base.markdown_normalizer import normalize_material_markdown


IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}
DEFAULT_PPTX_RENDER_MAX_SLIDES = 120
DEFAULT_PPTX_OCR_MAX_SLIDES = 40
DEFAULT_PDF_RENDER_MAX_PAGES = 120
QWEN_OCR_PROVIDERS = {"qwen_ocr", "qwen-vl-ocr", "qwen_vl_ocr", "qwen_ocr_vl", "qwen_vl"}

ProgressCallback = Callable[[dict], None]


@dataclass(frozen=True)
class ParsedBlock:
    modality: str
    text: str = ""
    confidence: float = 0
    warnings: list[str] = field(default_factory=list)
    metadata: dict = field(default_factory=dict)

    def __post_init__(self) -> None:
        if not str(self.text or "").strip():
            return
        metadata = dict(self.metadata or {})
        normalized = normalize_material_markdown(self.text)
        metadata.setdefault("contentFormat", normalized.content_format)
        metadata.setdefault("markdownText", normalized.markdown)
        metadata.setdefault("rawTextPreserved", True)
        metadata.setdefault("formatFeatures", normalized.features)
        object.__setattr__(self, "metadata", metadata)


@dataclass(frozen=True)
class VisionUnderstanding:
    text: str = ""
    formulas: list[str] = field(default_factory=list)
    tables: list[dict] = field(default_factory=list)
    figures: list[dict] = field(default_factory=list)
    keywords: list[str] = field(default_factory=list)
    raw: dict = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)

    def to_search_text(self) -> str:
        parts: list[str] = []
        if self.text.strip():
            parts.append(self.text.strip())
        parts.extend(item.strip() for item in self.formulas if str(item).strip())
        for table in self.tables:
            text = str(table.get("markdown") or table.get("text") or table.get("caption") or "").strip()
            if text:
                parts.append(text)
        for figure in self.figures:
            text = " ".join(
                str(figure.get(key) or "").strip()
                for key in ("caption", "description")
                if str(figure.get(key) or "").strip()
            )
            if text:
                parts.append(text)
        parts.extend(keyword.strip() for keyword in self.keywords if str(keyword).strip())
        return "\n".join(dict.fromkeys(parts))


@dataclass(frozen=True)
class ParsedAsset:
    kind: str
    storage_path: str
    mime_type: str = ""
    sha256: str = ""
    width: int = 0
    height: int = 0
    metadata: dict = field(default_factory=dict)


@dataclass(frozen=True)
class ParsedPage:
    page_number: int | None = None
    slide_number: int | None = None
    ocr_text: str = ""
    summary: str = ""
    quality_score: float = 0
    warnings: list[str] = field(default_factory=list)
    metadata: dict = field(default_factory=dict)
    assets: list[ParsedAsset] = field(default_factory=list)
    blocks: list[ParsedBlock] = field(default_factory=list)


@dataclass(frozen=True)
class MultimodalParseResult:
    parser: str
    status: str
    pages: list[ParsedPage] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    quality_score: float = 0
    result: dict = field(default_factory=dict)

    @property
    def has_searchable_text(self) -> bool:
        return any(block.text.strip() for page in self.pages for block in page.blocks)

    @property
    def has_visual_evidence(self) -> bool:
        return any(
            block.modality in {"image", "page_image", "slide_image"}
            for page in self.pages
            for block in page.blocks
        )


def parse_resource(source_path: Path, progress_callback: ProgressCallback | None = None) -> MultimodalParseResult:
    source_path = source_path.resolve()
    if not source_path.exists():
        return finalize_parse_result(MultimodalParseResult(
            parser="native",
            status="failed",
            warnings=[f"Source file does not exist: {source_path.name}"],
            quality_score=0,
        ))
    suffix = source_path.suffix.lower()
    if suffix == ".pptx":
        return finalize_parse_result(_parse_pptx(source_path, progress_callback=progress_callback))
    if suffix in IMAGE_EXTENSIONS:
        return finalize_parse_result(_parse_image(source_path, progress_callback=progress_callback))
    if suffix == ".xlsx":
        return finalize_parse_result(_parse_xlsx(source_path))
    if suffix == ".pdf":
        return finalize_parse_result(_parse_pdf(source_path, progress_callback=progress_callback))
    return finalize_parse_result(MultimodalParseResult(
        parser="native",
        status="failed",
        warnings=[f"Unsupported multimodal resource type: {suffix or 'unknown'}"],
        quality_score=0,
    ))


def _parse_image(source_path: Path, progress_callback: ProgressCallback | None = None) -> MultimodalParseResult:
    width = 0
    height = 0
    image_warnings: list[str] = []
    try:
        from PIL import Image

        with Image.open(source_path) as image:
            width, height = image.size
    except Exception as exc:
        image_warnings.append(f"Image dimensions unavailable: {type(exc).__name__}: {exc}")
    _emit_progress(progress_callback, {
        "phase": "ocr",
        "phaseLabel": "正在识别图片",
        "page": {"current": 1, "total": 1},
        "event": "开始图片 OCR/VLM 理解",
        "phaseProgress": 28,
    })
    understanding = _understand_image(source_path)
    ocr_text = understanding.text
    image_warnings.extend(understanding.warnings)
    image_warnings.append("Visual embedding is not configured yet; image is registered as visual evidence only.")
    asset = ParsedAsset(
        kind="original",
        storage_path=str(source_path),
        mime_type=mimetypes.guess_type(source_path.name)[0] or "application/octet-stream",
        sha256=_sha256(source_path),
        width=width,
        height=height,
    )
    blocks = _vision_understanding_blocks(
        understanding,
        visual_modality="image",
        visual_fallback=source_path.stem,
        warnings=image_warnings,
        metadata={"sourceType": "image"},
    )
    page = ParsedPage(
        page_number=1,
        ocr_text=ocr_text,
        summary=f"Image resource: {source_path.name}",
        quality_score=0.65 if ocr_text else 0.45,
        warnings=image_warnings,
        metadata={"width": width, "height": height},
        assets=[asset],
        blocks=blocks,
    )
    return MultimodalParseResult(
        parser="native",
        status="completed_with_warnings",
        pages=[page],
        warnings=image_warnings,
        quality_score=0.65 if ocr_text else 0.45,
        result={"imageCount": 1, "ocrChars": len(ocr_text)},
    )


def _parse_pptx(source_path: Path, progress_callback: ProgressCallback | None = None) -> MultimodalParseResult:
    try:
        from pptx import Presentation
    except ModuleNotFoundError as exc:
        return MultimodalParseResult(
            parser="native",
            status="failed",
            warnings=[f"python-pptx is required to parse PPTX: {exc}"],
            quality_score=0,
        )

    try:
        presentation = Presentation(str(source_path))
    except Exception as exc:
        return MultimodalParseResult(
            parser="native",
            status="failed",
            warnings=[f"PPTX parse failed: {type(exc).__name__}: {exc}"],
            quality_score=0,
        )
    pages: list[ParsedPage] = []
    warnings: list[str] = []
    slide_total = len(presentation.slides)
    _emit_progress(progress_callback, {
        "phase": "render",
        "phaseLabel": f"正在渲染 PPTX 幻灯片 0/{slide_total}",
        "page": {"current": 0, "total": slide_total},
        "event": f"开始渲染 {slide_total} 页 PPTX",
        "phaseProgress": 12,
    })
    rendered_slides, render_warnings = _render_pptx_slides(source_path, progress_callback=progress_callback, total_hint=slide_total)
    warnings.extend(render_warnings)
    text_slide_count = 0
    ocr_slide_count = 0
    ocr_candidates_used = 0
    max_ocr_slides = _env_int("A3_PPTX_OCR_MAX_SLIDES", DEFAULT_PPTX_OCR_MAX_SLIDES)
    ocr_text_slides = _env_bool("A3_PPTX_OCR_TEXT_SLIDES", False)
    block_diagnostics = {
        "titleBlockCount": 0,
        "textBlockCount": 0,
        "bulletBlockCount": 0,
        "tableBlockCount": 0,
        "imageCaptionBlockCount": 0,
        "speakerNoteBlockCount": 0,
        "captionDegradedCount": 0,
    }
    for slide_index, slide in enumerate(presentation.slides, start=1):
        _emit_progress(progress_callback, {
            "phase": "parse",
            "phaseLabel": f"正在解析 PPTX 第 {slide_index}/{slide_total} 页",
            "page": {"current": slide_index, "total": slide_total},
            "metrics": {"pagesParsed": slide_index - 1, "blocksCreated": sum(len(page.blocks) for page in pages), "warnings": len(warnings)},
            "event": f"开始解析第 {slide_index} 页",
            "phaseProgress": min(55, 30 + round((slide_index - 1) / max(slide_total, 1) * 25)),
        })
        extracted_blocks = _extract_pptx_slide_blocks(slide, slide_index)
        for block in extracted_blocks:
            block_type = str(block.metadata.get("blockType") or "")
            if block_type == "slide_title":
                block_diagnostics["titleBlockCount"] += 1
            elif block_type == "text_box":
                block_diagnostics["textBlockCount"] += 1
            elif block_type == "bullet":
                block_diagnostics["bulletBlockCount"] += 1
            elif block_type == "table":
                block_diagnostics["tableBlockCount"] += 1
            elif block_type == "image_caption":
                block_diagnostics["imageCaptionBlockCount"] += 1
                if block.warnings:
                    block_diagnostics["captionDegradedCount"] += 1
            elif block_type == "speaker_note":
                block_diagnostics["speakerNoteBlockCount"] += 1
        text_parts = [block.text for block in extracted_blocks if block.text.strip()]
        text = "\n".join(part for part in text_parts if part)
        rendered_asset = rendered_slides.get(slide_index)
        ocr_text = ""
        understanding = VisionUnderstanding()
        page_warnings: list[str] = []
        if rendered_asset:
            should_ocr = (not text or ocr_text_slides) and ocr_candidates_used < max_ocr_slides
            if should_ocr:
                ocr_candidates_used += 1
                _emit_progress(progress_callback, {
                    "phase": "ocr",
                    "phaseLabel": f"正在 OCR PPTX 第 {slide_index}/{slide_total} 页",
                    "page": {"current": slide_index, "total": slide_total},
                    "metrics": {"pagesParsed": slide_index - 1, "blocksCreated": sum(len(page.blocks) for page in pages), "warnings": len(warnings)},
                    "event": f"开始 OCR 第 {slide_index} 页",
                    "phaseProgress": min(62, 40 + round(slide_index / max(slide_total, 1) * 22)),
                })
                understanding = _understand_image(Path(rendered_asset.storage_path))
                ocr_text = understanding.text
                page_warnings.extend(understanding.warnings)
                warnings.extend(f"slide {slide_index}: {warning}" for warning in understanding.warnings)
                if ocr_text:
                    ocr_slide_count += 1
            elif text:
                page_warnings.append("OCR skipped because PPTX slide has extractable text.")
            else:
                page_warnings.append(f"PPTX OCR limited to first {max_ocr_slides} image-only slides.")
        blocks: list[ParsedBlock] = []
        if extracted_blocks:
            text_slide_count += 1
            blocks.extend(extracted_blocks)
        if rendered_asset:
            blocks.extend(
                _vision_understanding_blocks(
                    understanding,
                    visual_modality="slide_image",
                    visual_fallback=f"Rendered PPTX slide {slide_index}",
                    warnings=page_warnings,
                    metadata={
                        "slide": slide_index,
                        "sourceType": "pptx",
                        "assetKind": "rendered_slide",
                    },
                )
            )
        if not text and not ocr_text:
            warning = "PPTX slide has no extractable text; rendering/OCR is required for full recognition."
            warnings.append(f"slide {slide_index}: {warning}")
            page_warnings.append(warning)
            if not blocks:
                blocks.append(
                    ParsedBlock(
                        modality="slide_image",
                        text=f"Slide {slide_index} visual evidence placeholder",
                        confidence=0.3,
                        warnings=[warning],
                        metadata={"slide": slide_index, "sourceType": "pptx"},
                    )
                )
        quality_score = 0.35
        if text and ocr_text:
            quality_score = 0.85
        elif text:
            quality_score = 0.8
        elif ocr_text:
            quality_score = 0.65
        summary = text[:300] if text else (ocr_text[:300] if ocr_text else f"Slide {slide_index} requires rendering/OCR.")
        page = ParsedPage(
            page_number=slide_index,
            slide_number=slide_index,
            ocr_text=ocr_text,
            summary=summary,
            quality_score=quality_score,
            warnings=page_warnings,
            assets=[rendered_asset] if rendered_asset else [],
            blocks=blocks,
            )
        pages.append(page)
        _emit_progress(progress_callback, {
            "phase": "chunk",
            "phaseLabel": f"正在生成 PPTX 知识块 第 {slide_index}/{slide_total} 页",
            "page": {"current": slide_index, "total": slide_total},
            "block": {"current": sum(len(item.blocks) for item in pages), "total": max(sum(len(item.blocks) for item in pages), 1)},
            "metrics": {
                "pagesParsed": slide_index,
                "blocksCreated": sum(len(item.blocks) for item in pages),
                "visualAssets": len(rendered_slides),
                "warnings": len(warnings),
            },
            "event": f"第 {slide_index} 页已生成 {len(blocks)} 个知识块",
            "phaseProgress": min(68, 46 + round(slide_index / max(slide_total, 1) * 22)),
        })

    if not pages:
        warnings.append("PPTX has no slides.")
    status = "completed" if pages and not warnings else "completed_with_warnings"
    quality_score = round(sum(page.quality_score for page in pages) / len(pages), 3) if pages else 0
    return MultimodalParseResult(
        parser="native",
        status=status,
        pages=pages,
        warnings=warnings,
        quality_score=quality_score,
        result={
            "slides": len(pages),
            "textSlides": text_slide_count,
            "renderedSlides": len(rendered_slides),
            "ocrSlides": ocr_slide_count,
            "ocrLimit": max_ocr_slides,
            "ocrTextSlides": ocr_text_slides,
            "pptxBlockDiagnostics": block_diagnostics,
        },
    )


def _extract_pptx_slide_blocks(slide, slide_index: int) -> list[ParsedBlock]:
    blocks: list[ParsedBlock] = []
    title_shape = getattr(slide.shapes, "title", None)
    for shape_index, shape in enumerate(slide.shapes, start=1):
        if _shape_has_table(shape):
            table_text = _pptx_table_markdown(shape.table)
            if table_text:
                blocks.append(
                    ParsedBlock(
                        modality="table",
                        text=table_text,
                        confidence=0.88,
                        metadata={
                            "slide": slide_index,
                            "sourceType": "pptx",
                            "blockType": "table",
                            "shapeIndex": shape_index,
                            "bbox": _shape_bbox(shape),
                        },
                    )
                )
            continue
        if _shape_is_picture(shape):
            caption, confidence, caption_warnings, caption_metadata = _pptx_picture_caption(shape, slide_index)
            blocks.append(
                ParsedBlock(
                    modality="image",
                    text=caption,
                    confidence=confidence,
                    warnings=caption_warnings,
                    metadata={
                        "slide": slide_index,
                        "sourceType": "pptx",
                        "blockType": "image_caption",
                        "shapeIndex": shape_index,
                        "bbox": _shape_bbox(shape),
                        **caption_metadata,
                    },
                )
            )
            continue
        text = str(getattr(shape, "text", "") or "").strip()
        if not text:
            continue
        block_type = "slide_title" if title_shape is not None and shape == title_shape else _pptx_text_block_type(shape, text)
        blocks.append(
            ParsedBlock(
                modality="text",
                text=text,
                confidence=0.9 if block_type == "slide_title" else 0.86,
                metadata={
                    "slide": slide_index,
                    "sourceType": "pptx",
                    "blockType": block_type,
                    "shapeIndex": shape_index,
                    "bbox": _shape_bbox(shape),
                },
            )
        )
    speaker_notes = _pptx_speaker_notes(slide)
    if speaker_notes:
        blocks.append(
            ParsedBlock(
                modality="text",
                text=speaker_notes,
                confidence=0.82,
                metadata={
                    "slide": slide_index,
                    "sourceType": "pptx",
                    "blockType": "speaker_note",
                },
            )
        )
    return blocks


def _shape_has_table(shape) -> bool:
    try:
        return bool(getattr(shape, "has_table", False))
    except Exception:
        return False


def _shape_is_picture(shape) -> bool:
    try:
        return str(getattr(getattr(shape, "shape_type", None), "name", "")).upper() == "PICTURE"
    except Exception:
        return False


def _pptx_picture_caption(shape, slide_index: int) -> tuple[str, float, list[str], dict]:
    alt_text = _shape_alt_text(shape)
    image_path: Path | None = None
    try:
        image = shape.image
        suffix = f".{str(getattr(image, 'ext', '') or 'png').lstrip('.')}"
        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as handle:
            handle.write(image.blob)
            image_path = Path(handle.name)
        understanding = _understand_image(image_path)
        visual_text = _visual_block_text(understanding)
        if visual_text:
            confidence = 0.78 if understanding.raw else 0.72
            return (
                visual_text,
                confidence,
                understanding.warnings,
                {
                    "captionMethod": "vision_provider",
                    "captionConfidence": confidence,
                    "captionProvider": os.getenv("A3_OCR_PROVIDER", "local").strip().lower() or "local",
                    "vlm": bool(understanding.raw or understanding.figures or understanding.tables or understanding.formulas or understanding.keywords),
                    "keywords": understanding.keywords,
                    "figures": understanding.figures,
                },
            )
        warnings = list(understanding.warnings)
    except Exception as exc:
        warnings = [f"PPTX image caption failed: {type(exc).__name__}: {exc}"]
    finally:
        if image_path:
            try:
                image_path.unlink(missing_ok=True)
            except Exception:
                pass
    caption = alt_text or f"Slide {slide_index} image: {getattr(shape, 'name', 'picture')}"
    return (
        caption,
        0.52 if alt_text else 0.45,
        [*warnings, "No vision provider caption was available; using image alt text or placeholder."],
        {
            "captionMethod": "alt_text_or_placeholder",
            "captionConfidence": 0.52 if alt_text else 0.45,
            "captionProvider": "fallback",
        },
    )


def _pptx_text_block_type(shape, text: str) -> str:
    normalized = text.strip()
    if normalized.startswith(("•", "-", "*")):
        return "bullet"
    try:
        text_frame = shape.text_frame
        paragraphs = list(text_frame.paragraphs)
        if any(getattr(paragraph, "level", 0) > 0 for paragraph in paragraphs):
            return "bullet"
        if len([paragraph for paragraph in paragraphs if str(paragraph.text or "").strip()]) > 1:
            return "bullet"
    except Exception:
        pass
    return "text_box"


def _pptx_table_markdown(table) -> str:
    rows: list[list[str]] = []
    for row in table.rows:
        values = [" ".join(str(cell.text or "").split()) for cell in row.cells]
        if any(values):
            rows.append(values)
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    padded = [row + [""] * (width - len(row)) for row in rows]
    header = padded[0]
    lines = ["| " + " | ".join(header) + " |", "| " + " | ".join(["---"] * width) + " |"]
    for row in padded[1:]:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


def _pptx_speaker_notes(slide) -> str:
    try:
        notes_slide = slide.notes_slide
        notes = str(notes_slide.notes_text_frame.text or "").strip()
    except Exception:
        return ""
    cleaned_lines = []
    for line in notes.splitlines():
        cleaned = line.strip()
        if cleaned and cleaned.lower() not in {"click to add notes"}:
            cleaned_lines.append(cleaned)
    return "\n".join(cleaned_lines)


def _shape_alt_text(shape) -> str:
    try:
        element = shape._element
        c_nv_pr = element.xpath(".//p:cNvPr")
        for item in c_nv_pr:
            text = str(item.get("descr") or item.get("title") or "").strip()
            if text:
                return text
    except Exception:
        return ""
    return ""


def _shape_bbox(shape) -> list[int]:
    values = []
    for attr in ("left", "top", "width", "height"):
        try:
            values.append(int(getattr(shape, attr)))
        except Exception:
            values.append(0)
    return values


def _parse_xlsx(source_path: Path) -> MultimodalParseResult:
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as exc:
        return MultimodalParseResult(
            parser="native",
            status="failed",
            warnings=[f"openpyxl is required to parse XLSX: {exc}"],
            quality_score=0,
        )
    try:
        workbook = load_workbook(source_path, read_only=True, data_only=True)
    except Exception as exc:
        return MultimodalParseResult(
            parser="native",
            status="failed",
            warnings=[f"XLSX parse failed: {type(exc).__name__}: {exc}"],
            quality_score=0,
        )

    pages: list[ParsedPage] = []
    warnings: list[str] = []
    for sheet_index, worksheet in enumerate(workbook.worksheets, start=1):
        rows: list[list[str]] = []
        for row in worksheet.iter_rows(max_row=80, values_only=True):
            values = ["" if value is None else str(value).strip() for value in row]
            if any(values):
                rows.append(values)
            if len(rows) >= 40:
                break
        if not rows:
            warnings.append(f"sheet {worksheet.title}: no extractable cells")
            continue
        text = "\n".join(" | ".join(cell for cell in row if cell) for row in rows)
        pages.append(
            ParsedPage(
                page_number=sheet_index,
                summary=f"Worksheet {worksheet.title}",
                quality_score=0.75,
                metadata={"sheet": worksheet.title},
                blocks=[
                    ParsedBlock(
                        modality="table",
                        text=text,
                        confidence=0.8,
                        metadata={"sheet": worksheet.title, "sourceType": "xlsx"},
                    )
                ],
            )
        )
    workbook.close()
    if not pages:
        return MultimodalParseResult(
            parser="native",
            status="failed",
            warnings=warnings or ["XLSX has no extractable sheets."],
            quality_score=0,
        )
    return MultimodalParseResult(
        parser="native",
        status="completed_with_warnings" if warnings else "completed",
        pages=pages,
        warnings=warnings,
        quality_score=0.75 if not warnings else 0.6,
        result={"sheets": len(workbook.worksheets), "parsedSheets": len(pages)},
    )


def _parse_pdf(source_path: Path, progress_callback: ProgressCallback | None = None) -> MultimodalParseResult:
    try:
        from pypdf import PdfReader
    except ModuleNotFoundError as exc:
        return MultimodalParseResult(
            parser="native",
            status="failed",
            warnings=[f"pypdf is required to parse PDF: {exc}"],
            quality_score=0,
        )
    try:
        reader = PdfReader(str(source_path))
    except Exception as exc:
        return MultimodalParseResult(
            parser="native",
            status="failed",
            warnings=[f"PDF parse failed: {type(exc).__name__}: {exc}"],
            quality_score=0,
        )

    total_pages = len(reader.pages)
    _emit_progress(progress_callback, {
        "phase": "render",
        "phaseLabel": f"正在渲染 PDF 页面 0/{total_pages}",
        "page": {"current": 0, "total": total_pages},
        "event": f"开始渲染 {total_pages} 页 PDF",
        "phaseProgress": 12,
    })
    rendered_pages, render_warnings = _render_pdf_pages(source_path, progress_callback=progress_callback, total_hint=total_pages)
    pages: list[ParsedPage] = []
    warnings: list[str] = []
    warnings.extend(render_warnings)
    text_pages = 0
    ocr_pages = 0
    for page_index, page in enumerate(reader.pages, start=1):
        _emit_progress(progress_callback, {
            "phase": "parse",
            "phaseLabel": f"正在解析 PDF 第 {page_index}/{total_pages} 页",
            "page": {"current": page_index, "total": total_pages},
            "metrics": {"pagesParsed": page_index - 1, "blocksCreated": sum(len(item.blocks) for item in pages), "warnings": len(warnings)},
            "event": f"开始解析第 {page_index} 页",
            "phaseProgress": min(54, 30 + round((page_index - 1) / max(total_pages, 1) * 24)),
        })
        text = "\n".join(line.strip() for line in (page.extract_text() or "").splitlines() if line.strip())
        rendered_asset = rendered_pages.get(page_index)
        ocr_text = ""
        understanding = VisionUnderstanding()
        page_warnings: list[str] = []
        if rendered_asset:
            _emit_progress(progress_callback, {
                "phase": "ocr",
                "phaseLabel": f"正在 OCR PDF 第 {page_index}/{total_pages} 页",
                "page": {"current": page_index, "total": total_pages},
                "metrics": {"pagesParsed": page_index - 1, "blocksCreated": sum(len(item.blocks) for item in pages), "warnings": len(warnings)},
                "event": f"开始 OCR 第 {page_index} 页",
                "phaseProgress": min(62, 40 + round(page_index / max(total_pages, 1) * 22)),
            })
            understanding = _understand_image(Path(rendered_asset.storage_path))
            ocr_text = understanding.text
            page_warnings.extend(understanding.warnings)
            if ocr_text:
                ocr_pages += 1
            warnings.extend(f"page {page_index}: {warning}" for warning in understanding.warnings)
        visual_blocks = (
            _vision_understanding_blocks(
                understanding,
                visual_modality="page_image",
                visual_fallback=f"Rendered PDF page {page_index}",
                warnings=page_warnings,
                metadata={"page": page_index, "sourceType": "pdf", "assetKind": "rendered_page"},
            )
            if rendered_asset
            else []
        )
        if text:
            text_pages += 1
            blocks = [
                ParsedBlock(
                    modality="text",
                    text=text,
                    confidence=0.8,
                    metadata={"page": page_index, "sourceType": "pdf"},
                )
            ]
            blocks.extend(visual_blocks)
            pages.append(
                ParsedPage(
                    page_number=page_index,
                    ocr_text=ocr_text,
                    summary=text[:300],
                    quality_score=0.8 if ocr_text else 0.75,
                    warnings=page_warnings,
                    assets=[rendered_asset] if rendered_asset else [],
                    blocks=blocks,
                )
            )
            _emit_progress(progress_callback, {
                "phase": "chunk",
                "phaseLabel": f"正在生成 PDF 知识块 第 {page_index}/{total_pages} 页",
                "page": {"current": page_index, "total": total_pages},
                "block": {"current": sum(len(item.blocks) for item in pages), "total": max(sum(len(item.blocks) for item in pages), 1)},
                "metrics": {
                    "pagesParsed": page_index,
                    "blocksCreated": sum(len(item.blocks) for item in pages),
                    "visualAssets": len(rendered_pages),
                    "warnings": len(warnings),
                },
                "event": f"第 {page_index} 页已生成 {len(blocks)} 个知识块",
                "phaseProgress": min(68, 46 + round(page_index / max(total_pages, 1) * 22)),
            })
        else:
            warning = "PDF page has no extractable text; rendering/OCR is required for full recognition."
            if not ocr_text:
                warnings.append(f"page {page_index}: {warning}")
            blocks = [
                *visual_blocks
            ] or [
                ParsedBlock(
                    modality="page_image",
                    text=f"Page {page_index} visual evidence placeholder",
                    confidence=0.25,
                    warnings=[warning],
                    metadata={"page": page_index, "sourceType": "pdf"},
                )
            ]
            pages.append(
                ParsedPage(
                    page_number=page_index,
                    ocr_text=ocr_text,
                    summary=ocr_text[:300] if ocr_text else f"Page {page_index} requires rendering/OCR.",
                    quality_score=0.6 if ocr_text else 0.3,
                    warnings=page_warnings if ocr_text else [warning, *page_warnings],
                    assets=[rendered_asset] if rendered_asset else [],
                    blocks=blocks,
                )
            )
            _emit_progress(progress_callback, {
                "phase": "chunk",
                "phaseLabel": f"正在生成 PDF 知识块 第 {page_index}/{total_pages} 页",
                "page": {"current": page_index, "total": total_pages},
                "block": {"current": sum(len(item.blocks) for item in pages), "total": max(sum(len(item.blocks) for item in pages), 1)},
                "metrics": {
                    "pagesParsed": page_index,
                    "blocksCreated": sum(len(item.blocks) for item in pages),
                    "visualAssets": len(rendered_pages),
                    "warnings": len(warnings),
                },
                "event": f"第 {page_index} 页已生成 {len(blocks)} 个知识块",
                "phaseProgress": min(68, 46 + round(page_index / max(total_pages, 1) * 22)),
            })
    if not pages:
        warnings.append("PDF has no pages.")
    return MultimodalParseResult(
        parser="native",
        status="completed" if pages and not warnings else "completed_with_warnings",
        pages=pages,
        warnings=warnings,
        quality_score=0.75 if pages and not warnings else (0.5 if pages else 0),
        result={"pages": len(pages), "textPages": text_pages, "ocrPages": ocr_pages},
    )


def _vision_understanding_blocks(
    understanding: VisionUnderstanding,
    *,
    visual_modality: str,
    visual_fallback: str,
    warnings: list[str],
    metadata: dict,
) -> list[ParsedBlock]:
    metadata = dict(metadata)
    provider = str((understanding.raw or {}).get("provider") or "").strip()
    structured_metadata = {
        **metadata,
        "ocr": bool(understanding.text),
        "vlm": bool(understanding.raw or understanding.formulas or understanding.tables or understanding.figures or understanding.keywords),
        "keywords": understanding.keywords,
        "figures": understanding.figures,
    }
    if provider:
        structured_metadata["provider"] = provider
    visual_text = _visual_block_text(understanding) or visual_fallback
    blocks = [
        ParsedBlock(
            modality=visual_modality,
            text=visual_text,
            confidence=0.7 if understanding.to_search_text() else 0.35,
            warnings=warnings,
            metadata=structured_metadata,
        )
    ]
    if understanding.text.strip():
        blocks.append(
            ParsedBlock(
                modality="text",
                text=understanding.text.strip(),
                confidence=0.72,
                metadata={
                    **metadata,
                    **({"provider": provider} if provider else {}),
                    "ocr": True,
                    "vlm": bool(structured_metadata["vlm"]),
                    "keywords": understanding.keywords,
                },
            )
        )
    for formula in understanding.formulas:
        text = str(formula).strip()
        if text:
            blocks.append(
                ParsedBlock(
                    modality="formula",
                    text=text,
                    confidence=0.7,
                    metadata={
                        **metadata,
                        **({"provider": provider} if provider else {}),
                        "vlm": bool(structured_metadata["vlm"]),
                        "keywords": understanding.keywords,
                    },
                )
            )
    for table in understanding.tables:
        text = str(table.get("markdown") or table.get("text") or "").strip()
        if text:
            table_metadata = {
                **metadata,
                **({"provider": provider} if provider else {}),
                "vlm": bool(structured_metadata["vlm"]),
                "caption": str(table.get("caption") or "").strip(),
            }
            blocks.append(
                ParsedBlock(
                    modality="table",
                    text=text,
                    confidence=0.7,
                    metadata=table_metadata,
                )
            )
    return blocks


def _visual_block_text(understanding: VisionUnderstanding) -> str:
    figure_parts = []
    for figure in understanding.figures:
        text = " ".join(
            str(figure.get(key) or "").strip()
            for key in ("caption", "description")
            if str(figure.get(key) or "").strip()
        )
        if text:
            figure_parts.append(text)
    return "\n".join(dict.fromkeys([*figure_parts, understanding.to_search_text()])).strip()


def _understand_image(image_path: Path) -> VisionUnderstanding:
    provider = os.getenv("A3_OCR_PROVIDER", "local").strip().lower()
    if provider in QWEN_OCR_PROVIDERS:
        understanding = QwenOcrDocumentParserProvider(provider=provider).understand_image(image_path)
        if understanding.to_search_text():
            return understanding
        local_text, local_warnings = _ocr_image_local(image_path)
        return VisionUnderstanding(text=local_text, warnings=[*understanding.warnings, *local_warnings])
    if provider in {"baidu", "baidu_paddleocr_vl", "paddleocr_vl"}:
        understanding = _understand_image_with_baidu_document_provider(image_path)
        if understanding.to_search_text():
            return understanding
        local_text, local_warnings = _ocr_image_local(image_path)
        return VisionUnderstanding(text=local_text, warnings=[*understanding.warnings, *local_warnings])
    if provider and provider not in {"local", "rapidocr", "none"}:
        understanding = _understand_image_with_openai_compatible_provider(image_path, provider=provider)
        if understanding.to_search_text():
            return understanding
        local_text, local_warnings = _ocr_image_local(image_path)
        return VisionUnderstanding(text=local_text, warnings=[*understanding.warnings, *local_warnings])
    if provider == "none":
        return VisionUnderstanding(warnings=["OCR disabled by A3_OCR_PROVIDER=none."])
    text, warnings = _ocr_image(image_path)
    return VisionUnderstanding(text=text, warnings=warnings)


def _ocr_image(image_path: Path) -> tuple[str, list[str]]:
    provider = os.getenv("A3_OCR_PROVIDER", "local").strip().lower()
    if provider in QWEN_OCR_PROVIDERS:
        understanding = QwenOcrDocumentParserProvider(provider=provider).understand_image(image_path)
        text = understanding.to_search_text()
        if text:
            return text, understanding.warnings
        local_text, local_warnings = _ocr_image_local(image_path)
        return local_text, [*understanding.warnings, *local_warnings]
    if provider in {"baidu", "baidu_paddleocr_vl", "paddleocr_vl"}:
        understanding = _understand_image_with_baidu_document_provider(image_path)
        text = understanding.to_search_text()
        if text:
            return text, understanding.warnings
        local_text, local_warnings = _ocr_image_local(image_path)
        return local_text, [*understanding.warnings, *local_warnings]
    if provider and provider not in {"local", "rapidocr", "none"}:
        text, warnings = _ocr_image_with_openai_compatible_provider(image_path, provider=provider)
        if text:
            return text, warnings
        local_text, local_warnings = _ocr_image_local(image_path)
        return local_text, [*warnings, *local_warnings]
    if provider == "none":
        return "", ["OCR disabled by A3_OCR_PROVIDER=none."]
    return _ocr_image_local(image_path)


def _ocr_image_local(image_path: Path) -> tuple[str, list[str]]:
    try:
        engine = _get_ocr_engine()
        result, _ = engine(str(image_path))
    except ModuleNotFoundError:
        return "", ["OCR engine is not installed; install rapidocr-onnxruntime to enable OCR."]
    except Exception as exc:
        return "", [f"OCR failed: {type(exc).__name__}: {exc}"]
    lines: list[str] = []
    for item in result or []:
        if not isinstance(item, (list, tuple)) or len(item) < 2:
            continue
        text = str(item[1]).strip()
        if text:
            lines.append(text)
    return "\n".join(lines), []


def _ocr_image_with_openai_compatible_provider(image_path: Path, *, provider: str) -> tuple[str, list[str]]:
    api_key = (_resolve_env_reference(os.getenv("A3_OCR_API_KEY", "")) or os.getenv("A3_LLM_API_KEY") or "").strip()
    base_url = (_resolve_env_reference(os.getenv("A3_OCR_BASE_URL", "")) or os.getenv("A3_LLM_BASE_URL") or "").strip()
    model = os.getenv("A3_OCR_MODEL", "").strip()
    if not api_key or not base_url or not model:
        return "", [f"OCR API provider {provider} is configured but A3_OCR_API_KEY/A3_OCR_BASE_URL/A3_OCR_MODEL is incomplete."]
    try:
        endpoint = _openai_compatible_chat_completions_url(base_url)
        payload = {
            "model": model,
            "temperature": 0,
            "messages": [
                {
                    "role": "user",
                    "content": [
                        {
                            "type": "text",
                            "text": (
                                "请只提取图片中的可见文字、公式、表格文字和关键图表标签。"
                                "保持原有顺序；如果没有可读文字，返回空字符串。"
                            ),
                        },
                        {
                            "type": "image_url",
                            "image_url": {"url": _image_data_url(image_path)},
                        },
                    ],
                }
            ],
        }
        request = urllib.request.Request(
            endpoint,
            data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
            headers={
                "Authorization": f"Bearer {api_key}",
                "Content-Type": "application/json",
            },
            method="POST",
        )
        timeout = _env_int("A3_OCR_API_TIMEOUT_SECONDS", 60)
        with urllib.request.urlopen(request, timeout=timeout) as response:
            raw = response.read().decode("utf-8")
        data = json.loads(raw)
        text = _extract_openai_compatible_text(data).strip()
        return text, [] if text else [f"OCR API provider {provider} returned no text."]
    except (OSError, urllib.error.URLError, urllib.error.HTTPError, json.JSONDecodeError, KeyError, RuntimeError, ValueError) as exc:
        return "", [f"OCR API failed ({provider}): {type(exc).__name__}: {exc}"]


def _understand_image_with_openai_compatible_provider(image_path: Path, *, provider: str) -> VisionUnderstanding:
    api_key = (_resolve_env_reference(os.getenv("A3_OCR_API_KEY", "")) or os.getenv("A3_LLM_API_KEY") or "").strip()
    base_url = (_resolve_env_reference(os.getenv("A3_OCR_BASE_URL", "")) or os.getenv("A3_LLM_BASE_URL") or "").strip()
    model = os.getenv("A3_OCR_MODEL", "").strip()
    if not api_key or not base_url or not model:
        return VisionUnderstanding(
            warnings=[f"OCR API provider {provider} is configured but A3_OCR_API_KEY/A3_OCR_BASE_URL/A3_OCR_MODEL is incomplete."]
        )
    try:
        endpoint = _openai_compatible_chat_completions_url(base_url)
        payload = {
            "model": model,
            "temperature": 0,
            "response_format": {"type": "json_object"},
            "messages": [
                {
                    "role": "system",
                    "content": "You extract teaching-document structure for multimodal RAG and return only valid JSON.",
                },
                {
                    "role": "user",
                    "content": [
                        {
                            "type": "text",
                            "text": (
                                "Return strict JSON only. Understand this teaching image/page for multimodal RAG. "
                                "Schema: {\"text\": string, \"formulas\": string[], "
                                "\"tables\": [{\"caption\": string, \"markdown\": string}], "
                                "\"figures\": [{\"caption\": string, \"description\": string}], "
                                "\"keywords\": string[]}. Preserve visible text order, formulas, table structure, "
                                "and describe diagrams needed for retrieval. Use empty arrays when absent."
                            ),
                        },
                        {
                            "type": "image_url",
                            "image_url": {"url": _image_data_url(image_path)},
                        },
                    ],
                }
            ],
        }
        request = urllib.request.Request(
            endpoint,
            data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
            headers={
                "Authorization": f"Bearer {api_key}",
                "Content-Type": "application/json",
            },
            method="POST",
        )
        timeout = _env_int("A3_OCR_API_TIMEOUT_SECONDS", 60)
        with urllib.request.urlopen(request, timeout=timeout) as response:
            raw = response.read().decode("utf-8")
        data = json.loads(raw)
        content = _extract_openai_compatible_text(data).strip()
        return _vision_understanding_from_text(content, provider=provider)
    except (OSError, urllib.error.URLError, urllib.error.HTTPError, json.JSONDecodeError, KeyError, RuntimeError, ValueError) as exc:
        return VisionUnderstanding(warnings=[f"OCR API failed ({provider}): {type(exc).__name__}: {exc}"])


class QwenOcrDocumentParserProvider:
    provider = "qwen_ocr"
    default_base_url = "https://dashscope.aliyuncs.com/compatible-mode/v1"
    default_model = "qwen-vl-ocr-latest"

    def __init__(self, *, provider: str = "qwen_ocr") -> None:
        self.provider = provider or self.provider
        self.api_key = (
            _resolve_env_reference(os.getenv("A3_OCR_API_KEY", ""))
            or os.getenv("DASHSCOPE_API_KEY")
            or os.getenv("A3_LLM_API_KEY")
            or ""
        ).strip()
        self.base_url = (
            _resolve_env_reference(os.getenv("A3_OCR_BASE_URL", ""))
            or os.getenv("A3_LLM_BASE_URL")
            or self.default_base_url
        ).strip()
        self.model = (os.getenv("A3_OCR_MODEL") or self.default_model).strip()
        self.timeout = _env_int("A3_OCR_API_TIMEOUT_SECONDS", 120)

    def understand_image(self, image_path: Path) -> VisionUnderstanding:
        if not self.api_key:
            return VisionUnderstanding(warnings=["QwenOcrDocumentParserProvider is missing A3_OCR_API_KEY/DASHSCOPE_API_KEY; falling back to local OCR."])
        started = time.monotonic()
        try:
            endpoint = _openai_compatible_chat_completions_url(self.base_url)
            payload = {
                "model": self.model,
                "temperature": 0,
                "messages": [
                    {
                        "role": "system",
                        "content": (
                            "You are a document OCR engine for teaching materials. "
                            "Extract visible text, formulas, tables, figure captions, and reading order. "
                            "Return Markdown when possible; keep tables as Markdown tables and formulas as LaTeX."
                        ),
                    },
                    {
                        "role": "user",
                        "content": [
                            {
                                "type": "text",
                                "text": (
                                    "Extract this teaching-material page. Preserve reading order; use Markdown tables for tables; use LaTeX for formulas. "
                                    "For diagrams or flowcharts, add one searchable caption. Do not invent invisible content."
                                ),
                            },
                            {
                                "type": "image_url",
                                "image_url": {"url": _image_data_url(image_path)},
                            },
                        ],
                    }
                ],
            }
            request = urllib.request.Request(
                endpoint,
                data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
                headers={
                    "Authorization": f"Bearer {self.api_key}",
                    "Content-Type": "application/json",
                },
                method="POST",
            )
            with urllib.request.urlopen(request, timeout=self.timeout) as response:
                raw = response.read().decode("utf-8")
            data = json.loads(raw)
            content = _extract_openai_compatible_text(data).strip()
            understanding = _vision_understanding_from_text(content, provider=self.provider)
            object.__setattr__(understanding, "raw", {
                **(understanding.raw or {}),
                "provider": self.provider,
                "model": self.model,
                "latencyMs": int((time.monotonic() - started) * 1000),
            })
            if understanding.to_search_text():
                return understanding
            return VisionUnderstanding(
                raw={"provider": self.provider, "model": self.model, "response": data},
                warnings=[f"Qwen OCR provider {self.model} returned no readable text."],
            )
        except (OSError, urllib.error.URLError, urllib.error.HTTPError, json.JSONDecodeError, KeyError, RuntimeError, ValueError) as exc:
            return VisionUnderstanding(warnings=[f"Qwen OCR failed ({self.model}); falling back to local OCR: {type(exc).__name__}: {exc}"])


class BaiduDocumentParserProvider:
    provider = "baidu_paddleocr_vl"

    def __init__(self) -> None:
        self.token = (os.getenv("A3_BAIDU_PADDLEOCR_TOKEN") or os.getenv("BAIDU_PADDLEOCR_TOKEN") or "").strip()
        self.submit_url = os.getenv("A3_BAIDU_PADDLEOCR_JOB_URL", "https://paddleocr.aistudio-app.com/api/v2/ocr/jobs").strip()
        self.result_url_template = os.getenv(
            "A3_BAIDU_PADDLEOCR_RESULT_URL_TEMPLATE",
            "https://paddleocr.aistudio-app.com/api/v2/ocr/jobs/{job_id}",
        ).strip()
        self.timeout = _env_int("A3_BAIDU_PADDLEOCR_TIMEOUT_SECONDS", 180)
        self.poll_interval = max(1, _env_int("A3_BAIDU_PADDLEOCR_POLL_INTERVAL_SECONDS", 2))

    def understand_image(self, image_path: Path) -> VisionUnderstanding:
        if not self.token:
            return VisionUnderstanding(warnings=["BaiduDocumentParserProvider 未配置 A3_BAIDU_PADDLEOCR_TOKEN，已降级为本地 OCR。"])
        started = time.monotonic()
        try:
            submit_data = _submit_baidu_paddleocr_job(
                image_path,
                submit_url=self.submit_url,
                token=self.token,
                optional_payload=_baidu_optional_payload(enhanced=False),
            )
            job_id = _extract_baidu_job_id(submit_data)
            if not job_id:
                return VisionUnderstanding(warnings=[f"BaiduDocumentParserProvider 未返回 job id：{submit_data}"])
            result_data = _poll_baidu_paddleocr_job(
                job_id,
                result_url_template=self.result_url_template,
                token=self.token,
                timeout=self.timeout,
                poll_interval=self.poll_interval,
            )
            understanding = _baidu_result_to_vision_understanding(result_data, job_id=job_id)
            if not _has_readable_vision_content(understanding) and _env_bool("A3_BAIDU_PADDLEOCR_RETRY_EMPTY", True):
                retry_data = _submit_baidu_paddleocr_job(
                    image_path,
                    submit_url=self.submit_url,
                    token=self.token,
                    optional_payload=_baidu_optional_payload(enhanced=True),
                )
                retry_job_id = _extract_baidu_job_id(retry_data)
                if retry_job_id:
                    retry_result = _poll_baidu_paddleocr_job(
                        retry_job_id,
                        result_url_template=self.result_url_template,
                        token=self.token,
                        timeout=self.timeout,
                        poll_interval=self.poll_interval,
                    )
                    retry_understanding = _baidu_result_to_vision_understanding(retry_result, job_id=retry_job_id)
                    if _has_readable_vision_content(retry_understanding):
                        object.__setattr__(retry_understanding, "raw", {
                            **(retry_understanding.raw or {}),
                            "retry": {
                                "reason": "first_empty",
                                "firstJobId": job_id,
                                "retryJobId": retry_job_id,
                                "enhanced": True,
                            },
                        })
                        understanding = retry_understanding
            object.__setattr__(understanding, "raw", {
                **(understanding.raw or {}),
                "provider": self.provider,
                "latencyMs": int((time.monotonic() - started) * 1000),
            })
            return understanding
        except (OSError, urllib.error.URLError, urllib.error.HTTPError, json.JSONDecodeError, TimeoutError, ValueError) as exc:
            return VisionUnderstanding(warnings=[f"BaiduDocumentParserProvider 解析失败，已降级为本地 OCR：{type(exc).__name__}: {exc}"])


def _understand_image_with_baidu_document_provider(image_path: Path) -> VisionUnderstanding:
    return BaiduDocumentParserProvider().understand_image(image_path)


def _baidu_optional_payload(*, enhanced: bool) -> dict:
    return {
        "useDocOrientationClassify": _env_bool("A3_BAIDU_PADDLEOCR_USE_DOC_ORIENTATION_CLASSIFY", enhanced),
        "useDocUnwarping": _env_bool("A3_BAIDU_PADDLEOCR_USE_DOC_UNWARPING", enhanced),
        "useChartRecognition": _env_bool("A3_BAIDU_PADDLEOCR_USE_CHART_RECOGNITION", enhanced),
    }


def _submit_baidu_paddleocr_job(image_path: Path, *, submit_url: str, token: str, optional_payload: dict | None = None) -> dict:
    fields = {
        "model": os.getenv("A3_BAIDU_PADDLEOCR_MODEL", "PaddleOCR-VL-1.6"),
        "optionalPayload": json.dumps(optional_payload or _baidu_optional_payload(enhanced=False), ensure_ascii=False),
        "fileType": image_path.suffix.lstrip(".") or "png",
    }
    body, content_type = _multipart_form_data(fields, file_field="file", file_path=image_path)
    request = urllib.request.Request(
        submit_url,
        data=body,
        headers={
            "Authorization": f"Bearer {token}",
            "Content-Type": content_type,
        },
        method="POST",
    )
    with urllib.request.urlopen(request, timeout=_env_int("A3_BAIDU_PADDLEOCR_SUBMIT_TIMEOUT_SECONDS", 60)) as response:
        return json.loads(response.read().decode("utf-8"))


def _poll_baidu_paddleocr_job(job_id: str, *, result_url_template: str, token: str, timeout: int, poll_interval: int) -> dict:
    deadline = time.time() + timeout
    last_data: dict = {}
    while time.time() <= deadline:
        url = result_url_template.format(job_id=urllib.parse.quote(job_id))
        request = urllib.request.Request(url, headers={"Authorization": f"Bearer {token}"}, method="GET")
        with urllib.request.urlopen(request, timeout=_env_int("A3_BAIDU_PADDLEOCR_RESULT_TIMEOUT_SECONDS", 30)) as response:
            last_data = json.loads(response.read().decode("utf-8"))
        status = str(
            _deep_get(last_data, "data.state")
            or _deep_get(last_data, "data.status")
            or last_data.get("state")
            or last_data.get("status")
            or ""
        ).lower()
        if status in {"succeeded", "success", "completed", "done", "finished"}:
            return last_data
        if status in {"failed", "error", "canceled", "cancelled"}:
            raise ValueError(f"job {job_id} failed: {last_data}")
        if _extract_baidu_markdown(last_data) or _extract_baidu_json_payload(last_data):
            return last_data
        time.sleep(poll_interval)
    raise TimeoutError(f"job {job_id} timed out after {timeout}s; last={last_data}")


def _baidu_result_to_vision_understanding(data: dict, *, job_id: str) -> VisionUnderstanding:
    markdown = _extract_baidu_markdown(data)
    json_payload = _extract_baidu_json_payload(data)
    text = markdown or _baidu_json_to_text(json_payload)
    plain = _vision_understanding_from_plain_text(text)
    tables = plain.tables
    formulas = plain.formulas
    if isinstance(json_payload, dict):
        tables = [*tables, *_extract_baidu_tables(json_payload)]
        formulas = [*formulas, *_extract_baidu_formulas(json_payload)]
    return VisionUnderstanding(
        text=text,
        formulas=list(dict.fromkeys(formulas)),
        tables=tables,
        figures=[],
        keywords=[],
        raw={"provider": "baidu_paddleocr_vl", "jobId": job_id, "result": data},
        warnings=[] if text else [f"Baidu PaddleOCR-VL job {job_id} returned no readable text."],
    )


def _has_readable_vision_content(understanding: VisionUnderstanding) -> bool:
    return bool(
        understanding.text.strip()
        or understanding.formulas
        or understanding.tables
        or understanding.figures
    )


def _extract_baidu_job_id(data: dict) -> str:
    for path in ("data.task_id", "data.taskId", "data.job_id", "data.jobId", "task_id", "taskId", "job_id", "jobId"):
        value = _deep_get(data, path)
        if value:
            return str(value)
    return ""


def _extract_baidu_markdown(data: dict) -> str:
    for path in ("data.markdown", "data.result.markdown", "result.markdown", "markdown"):
        value = _deep_get(data, path)
        if isinstance(value, str) and value.strip():
            return value.strip()
    json_payload = _extract_baidu_json_payload(data)
    markdown = _extract_baidu_markdown_from_payload(json_payload)
    if markdown:
        return markdown
    markdown_url = _deep_get(data, "data.resultUrl.markdownUrl") or _deep_get(data, "resultUrl.markdownUrl")
    return _download_text(markdown_url) if markdown_url else ""


def _extract_baidu_json_payload(data: dict):
    for path in ("data.json", "data.result.json", "result.json", "json"):
        value = _deep_get(data, path)
        if isinstance(value, (dict, list)):
            return value
    json_url = _deep_get(data, "data.resultUrl.jsonUrl") or _deep_get(data, "resultUrl.jsonUrl")
    if json_url:
        raw = _download_text(json_url)
        if raw:
            try:
                return json.loads(raw)
            except json.JSONDecodeError:
                jsonl_payload = _parse_jsonl_payload(raw)
                return jsonl_payload if jsonl_payload else {"text": raw}
    return {}


def _parse_jsonl_payload(raw: str) -> list[dict]:
    items: list[dict] = []
    for line in raw.splitlines():
        line = line.strip()
        if not line:
            continue
        try:
            parsed = json.loads(line)
        except json.JSONDecodeError:
            continue
        if isinstance(parsed, dict):
            items.append(parsed)
    return items


def _extract_baidu_markdown_from_payload(payload) -> str:
    parts: list[str] = []
    if isinstance(payload, dict):
        markdown = payload.get("markdown")
        if isinstance(markdown, dict):
            text = str(markdown.get("text") or "").strip()
            if text:
                parts.append(text)
        elif isinstance(markdown, str) and markdown.strip():
            parts.append(markdown.strip())
        result = payload.get("result")
        if isinstance(result, dict):
            parts.append(_extract_baidu_markdown_from_payload(result))
        elif isinstance(result, list):
            for item in result:
                parts.append(_extract_baidu_markdown_from_payload(item))
        layout_results = payload.get("layoutParsingResults")
        if isinstance(layout_results, list):
            for item in layout_results:
                parts.append(_extract_baidu_markdown_from_payload(item))
        for key in ("pages", "blocks", "results"):
            value = payload.get(key)
            if isinstance(value, list):
                for item in value:
                    parts.append(_extract_baidu_markdown_from_payload(item))
    elif isinstance(payload, list):
        for item in payload:
            parts.append(_extract_baidu_markdown_from_payload(item))
    return "\n\n".join(part for part in parts if part).strip()


def _baidu_json_to_text(payload) -> str:
    if isinstance(payload, dict):
        parts = []
        for key in ("markdown", "text", "content"):
            value = payload.get(key)
            if isinstance(value, str) and value.strip():
                parts.append(value.strip())
        for key in ("pages", "blocks", "results"):
            value = payload.get(key)
            if isinstance(value, list):
                parts.extend(_baidu_json_to_text(item) for item in value)
        return "\n".join(part for part in parts if part)
    if isinstance(payload, list):
        return "\n".join(_baidu_json_to_text(item) for item in payload)
    return str(payload).strip() if payload else ""


def _extract_baidu_tables(payload) -> list[dict]:
    tables: list[dict] = []
    if isinstance(payload, dict):
        block_type = str(payload.get("type") or payload.get("blockType") or "").lower()
        markdown = str(payload.get("markdown") or payload.get("html") or payload.get("text") or "").strip()
        if "table" in block_type and markdown:
            tables.append({"caption": str(payload.get("caption") or ""), "markdown": markdown})
        for value in payload.values():
            tables.extend(_extract_baidu_tables(value))
    elif isinstance(payload, list):
        for item in payload:
            tables.extend(_extract_baidu_tables(item))
    return tables


def _extract_baidu_formulas(payload) -> list[str]:
    formulas: list[str] = []
    if isinstance(payload, dict):
        block_type = str(payload.get("type") or payload.get("blockType") or "").lower()
        text = str(payload.get("latex") or payload.get("formula") or payload.get("text") or "").strip()
        if "formula" in block_type and text:
            formulas.append(text)
        for value in payload.values():
            formulas.extend(_extract_baidu_formulas(value))
    elif isinstance(payload, list):
        for item in payload:
            formulas.extend(_extract_baidu_formulas(item))
    return formulas


def _download_text(url: str) -> str:
    if not url:
        return ""
    with urllib.request.urlopen(url, timeout=_env_int("A3_BAIDU_PADDLEOCR_DOWNLOAD_TIMEOUT_SECONDS", 30)) as response:
        return response.read().decode("utf-8", errors="replace")


def _deep_get(data: dict, path: str):
    current = data
    for part in path.split("."):
        if not isinstance(current, dict):
            return None
        current = current.get(part)
    return current


def _multipart_form_data(fields: dict[str, str], *, file_field: str, file_path: Path) -> tuple[bytes, str]:
    boundary = f"----A3Boundary{hashlib.sha256(str(file_path).encode()).hexdigest()[:16]}"
    lines: list[bytes] = []
    for name, value in fields.items():
        lines.extend([
            f"--{boundary}\r\n".encode(),
            f'Content-Disposition: form-data; name="{name}"\r\n\r\n'.encode(),
            str(value).encode("utf-8"),
            b"\r\n",
        ])
    mime_type = mimetypes.guess_type(file_path.name)[0] or "application/octet-stream"
    lines.extend([
        f"--{boundary}\r\n".encode(),
        f'Content-Disposition: form-data; name="{file_field}"; filename="{file_path.name}"\r\n'.encode("utf-8"),
        f"Content-Type: {mime_type}\r\n\r\n".encode(),
        file_path.read_bytes(),
        b"\r\n",
        f"--{boundary}--\r\n".encode(),
    ])
    return b"".join(lines), f"multipart/form-data; boundary={boundary}"


def _vision_understanding_from_text(content: str, *, provider: str) -> VisionUnderstanding:
    if not content:
        return VisionUnderstanding(warnings=[f"OCR API provider {provider} returned no text."])
    parsed = _parse_json_object_from_model_text(content)
    if not parsed:
        return _vision_understanding_from_plain_text(content)
    text = str(parsed.get("text") or "").strip()
    formulas = _coerce_str_list(parsed.get("formulas"))
    tables = _coerce_dict_list(parsed.get("tables"), allowed_keys={"caption", "markdown", "text"})
    figures = _coerce_dict_list(parsed.get("figures"), allowed_keys={"caption", "description"})
    keywords = _coerce_str_list(parsed.get("keywords"))
    if text and (not formulas or not tables):
        fallback = _vision_understanding_from_plain_text(text)
        if not formulas:
            formulas = fallback.formulas
        if not tables:
            tables = fallback.tables
    return VisionUnderstanding(
        text=text,
        formulas=formulas,
        tables=tables,
        figures=figures,
        keywords=keywords,
        raw=parsed,
    )


def _vision_understanding_from_plain_text(content: str) -> VisionUnderstanding:
    formulas: list[str] = []
    table_lines: list[str] = []
    for raw_line in content.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        lowered = line.lower()
        if lowered.startswith("formula:"):
            formula = line.split(":", 1)[1].strip()
            if formula:
                formulas.append(formula)
            continue
        if "|" in line:
            if lowered.startswith("table:"):
                line = line.split(":", 1)[1].strip()
            table_lines.append(line)
    tables = [{"caption": "", "markdown": "\n".join(table_lines)}] if table_lines else []
    return VisionUnderstanding(
        text=content.strip(),
        formulas=list(dict.fromkeys(formulas)),
        tables=tables,
    )


def _parse_json_object_from_model_text(content: str) -> dict:
    text = content.strip()
    if text.startswith("```"):
        lines = text.splitlines()
        if lines and lines[0].strip().startswith("```"):
            lines = lines[1:]
        if lines and lines[-1].strip() == "```":
            lines = lines[:-1]
        text = "\n".join(lines).strip()
    try:
        data = json.loads(text)
    except json.JSONDecodeError:
        start = text.find("{")
        end = text.rfind("}")
        if start < 0:
            return {}
        candidates = [text[start:]]
        if end > start:
            candidates.append(text[start : end + 1])
        repaired_candidates = [item for candidate in candidates for item in (candidate, _complete_json_fragment(candidate))]
        for value in repaired_candidates:
            try:
                data = json.loads(value)
                break
            except json.JSONDecodeError:
                data = {}
        else:
            data = {}
    return data if isinstance(data, dict) else {}


def _complete_json_fragment(text: str) -> str:
    stack: list[str] = []
    in_string = False
    escape = False
    pairs = {"{": "}", "[": "]"}
    for char in text:
        if in_string:
            if escape:
                escape = False
            elif char == "\\":
                escape = True
            elif char == '"':
                in_string = False
            continue
        if char == '"':
            in_string = True
        elif char in pairs:
            stack.append(pairs[char])
        elif char in {"}", "]"} and stack and stack[-1] == char:
            stack.pop()
    return text + "".join(reversed(stack))


def _coerce_str_list(value) -> list[str]:
    if not isinstance(value, list):
        return []
    return [str(item).strip() for item in value if str(item).strip()]


def _coerce_dict_list(value, *, allowed_keys: set[str]) -> list[dict]:
    if not isinstance(value, list):
        return []
    items: list[dict] = []
    for item in value:
        if isinstance(item, dict):
            normalized = {key: str(item.get(key) or "").strip() for key in allowed_keys}
        else:
            normalized = {"text": str(item).strip()}
        if any(normalized.values()):
            items.append(normalized)
    return items


def _openai_compatible_chat_completions_url(base_url: str) -> str:
    normalized = base_url.rstrip("/")
    if normalized.endswith("/chat/completions"):
        return normalized
    return f"{normalized}/chat/completions"


def _image_data_url(image_path: Path) -> str:
    mime_type = mimetypes.guess_type(image_path.name)[0] or "image/png"
    encoded = base64.b64encode(image_path.read_bytes()).decode("ascii")
    return f"data:{mime_type};base64,{encoded}"


def _extract_openai_compatible_text(data: dict) -> str:
    choices = data.get("choices")
    if not isinstance(choices, list) or not choices:
        return ""
    message = choices[0].get("message") if isinstance(choices[0], dict) else None
    content = message.get("content") if isinstance(message, dict) else ""
    if isinstance(content, str):
        return content
    if isinstance(content, list):
        parts = []
        for item in content:
            if isinstance(item, dict) and isinstance(item.get("text"), str):
                parts.append(item["text"])
        return "\n".join(parts)
    return ""


@lru_cache(maxsize=1)
def _get_ocr_engine():
    from rapidocr_onnxruntime import RapidOCR

    return RapidOCR()


def _render_pdf_pages(
    source_path: Path,
    max_pages: int | None = None,
    progress_callback: ProgressCallback | None = None,
    total_hint: int | None = None,
) -> tuple[dict[int, ParsedAsset], list[str]]:
    max_pages = max_pages if max_pages is not None else _env_int("A3_PDF_RENDER_MAX_PAGES", DEFAULT_PDF_RENDER_MAX_PAGES)
    output_dir = source_path.parent / ".a3_rendered" / source_path.stem
    return _render_pdf_to_assets(
        source_path,
        output_dir=output_dir,
        kind="rendered_page",
        metadata_name="page",
        max_pages=max_pages,
        progress_callback=progress_callback,
        total_hint=total_hint,
    )


def _render_pptx_slides(
    source_path: Path,
    max_slides: int | None = None,
    progress_callback: ProgressCallback | None = None,
    total_hint: int | None = None,
) -> tuple[dict[int, ParsedAsset], list[str]]:
    max_slides = max_slides if max_slides is not None else _env_int("A3_PPTX_RENDER_MAX_SLIDES", DEFAULT_PPTX_RENDER_MAX_SLIDES)
    soffice = _find_soffice()
    if not soffice:
        return {}, ["LibreOffice/soffice is not installed; PPTX slide rendering is unavailable."]
    output_dir = source_path.parent / ".a3_rendered" / source_path.stem
    pdf_dir = output_dir / "pdf"
    pdf_dir.mkdir(parents=True, exist_ok=True)
    try:
        completed = subprocess.run(
            [
                soffice,
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(pdf_dir),
                str(source_path),
            ],
            check=False,
            capture_output=True,
            text=True,
            timeout=180,
        )
    except Exception as exc:
        return {}, [f"PPTX rendering failed to start LibreOffice: {type(exc).__name__}: {exc}"]
    if completed.returncode != 0:
        message = (completed.stderr or completed.stdout or "unknown LibreOffice error").strip()
        return {}, [f"PPTX rendering failed: {message[:500]}"]
    candidates = sorted(pdf_dir.glob("*.pdf"), key=lambda path: path.stat().st_mtime, reverse=True)
    if not candidates:
        return {}, ["PPTX rendering failed: LibreOffice did not produce a PDF."]
    return _render_pdf_to_assets(
        candidates[0],
        output_dir=output_dir / "slides",
        kind="rendered_slide",
        metadata_name="slide",
        max_pages=max_slides,
        progress_callback=progress_callback,
        total_hint=total_hint,
    )


def _find_soffice() -> str:
    explicit = os.getenv("A3_SOFFICE_PATH", "").strip()
    if explicit and Path(explicit).exists():
        return explicit
    configured = shutil.which("soffice") or shutil.which("libreoffice")
    if configured:
        return configured
    workspace_tool = Path.cwd() / ".tools" / "LibreOffice" / "program" / "soffice.exe"
    for candidate in [
        str(workspace_tool),
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ]:
        if Path(candidate).exists():
            return candidate
    return ""


def _render_pdf_to_assets(
    source_path: Path,
    *,
    output_dir: Path,
    kind: str,
    metadata_name: str,
    max_pages: int,
    progress_callback: ProgressCallback | None = None,
    total_hint: int | None = None,
) -> tuple[dict[int, ParsedAsset], list[str]]:
    try:
        import pypdfium2 as pdfium
    except ModuleNotFoundError:
        return {}, ["pypdfium2 is not installed; PDF page rendering is unavailable."]
    output_dir.mkdir(parents=True, exist_ok=True)
    rendered: dict[int, ParsedAsset] = {}
    warnings: list[str] = []
    try:
        pdf = pdfium.PdfDocument(str(source_path))
    except Exception as exc:
        return {}, [f"PDF rendering failed to open document: {type(exc).__name__}: {exc}"]
    total_pages = len(pdf)
    try:
        page_count = min(total_pages, max_pages)
        for zero_index in range(page_count):
            page_number = zero_index + 1
            try:
                _emit_progress(progress_callback, {
                    "phase": "render",
                    "phaseLabel": f"正在渲染第 {page_number}/{total_hint or total_pages} 页",
                    "page": {"current": page_number, "total": total_hint or total_pages},
                    "event": f"开始渲染第 {page_number} 页",
                    "phaseProgress": min(30, 12 + round(page_number / max(page_count, 1) * 18)),
                })
                page = pdf[zero_index]
                bitmap = page.render(scale=1.5)
                image = bitmap.to_pil()
                output_path = output_dir / f"page-{page_number}.png"
                image.save(output_path, format="PNG")
                width, height = image.size
                image.close()
                rendered[page_number] = ParsedAsset(
                    kind=kind,
                    storage_path=str(output_path.resolve()),
                    mime_type="image/png",
                    sha256=_sha256(output_path),
                    width=width,
                    height=height,
                    metadata={metadata_name: page_number, "renderer": "pypdfium2"},
                )
            except Exception as exc:
                warnings.append(f"{metadata_name} {page_number}: PDF render failed: {type(exc).__name__}: {exc}")
    finally:
        pdf.close()
    if total_pages > max_pages:
        warnings.append(f"PDF rendering limited to first {max_pages} pages.")
    return rendered, warnings


def _emit_progress(progress_callback: ProgressCallback | None, detail: dict) -> None:
    if progress_callback is None:
        return
    progress_callback(detail)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _env_int(name: str, default: int) -> int:
    try:
        return max(0, int(os.getenv(name, str(default))))
    except ValueError:
        return default


def _resolve_env_reference(value: str) -> str:
    raw = (value or "").strip()
    if raw.startswith("${") and raw.endswith("}") and len(raw) > 3:
        return os.getenv(raw[2:-1], "").strip()
    return raw


def _env_bool(name: str, default: bool) -> bool:
    raw = os.getenv(name)
    if raw is None:
        return default
    return raw.strip().lower() in {"1", "true", "yes", "on"}
