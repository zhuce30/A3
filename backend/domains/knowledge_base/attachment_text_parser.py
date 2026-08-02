from __future__ import annotations

from dataclasses import dataclass, field
from itertools import islice
from pathlib import Path
from typing import Any


MAX_PDF_PAGES = 120
MAX_PPTX_SLIDES = 120
MAX_SHAPES_PER_SLIDE = 200
MAX_DOCX_BLOCKS = 2_000
MAX_XLSX_SHEETS = 32
MAX_XLSX_ROWS_PER_SHEET = 200
MAX_XLSX_CELLS_PER_ROW = 100
MAX_XLSX_TOTAL_CELLS = 20_000
MAX_IMAGE_PIXELS = 100_000_000


@dataclass(frozen=True)
class AttachmentTextResult:
    text: str
    metadata: dict[str, Any] = field(default_factory=dict)


class _BoundedText:
    def __init__(self, maximum: int) -> None:
        self.maximum = max(0, int(maximum))
        self.parts: list[str] = []
        self.length = 0
        self.truncated = False

    @property
    def full(self) -> bool:
        return self.length >= self.maximum

    def add(self, value: object, *, separator: str = "\n") -> None:
        if self.full:
            self.truncated = True
            return
        text = str(value or "").strip()
        if not text:
            return
        prefix = separator if self.parts else ""
        remaining = self.maximum - self.length
        combined = prefix + text
        self.parts.append(combined[:remaining])
        self.length += min(len(combined), remaining)
        if len(combined) > remaining:
            self.truncated = True

    def result(self) -> str:
        return "".join(self.parts)


def parse_attachment_text(path: Path, mime_type: str, *, max_chars: int) -> AttachmentTextResult:
    normalized_mime = str(mime_type or "").split(";", 1)[0].strip().lower()
    if normalized_mime in {"text/plain", "text/markdown", "text/x-markdown"}:
        return _parse_text(path, max_chars)
    if normalized_mime == "application/pdf":
        return _parse_pdf_text(path, max_chars)
    if normalized_mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return _parse_docx_text(path, max_chars)
    if normalized_mime == "application/vnd.openxmlformats-officedocument.presentationml.presentation":
        return _parse_pptx_text(path, max_chars)
    if normalized_mime == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        return _parse_xlsx_text(path, max_chars)
    if normalized_mime in {"image/png", "image/jpeg", "image/webp"}:
        return _parse_image_metadata(path, normalized_mime)
    raise ValueError("unsupported persisted attachment MIME type")


def _parse_text(path: Path, max_chars: int) -> AttachmentTextResult:
    with path.open("r", encoding="utf-8") as source:
        text = source.read(max(0, int(max_chars)) + 1)
    truncated = len(text) > max_chars
    return AttachmentTextResult(
        text=text[:max_chars],
        metadata={"parser": "attachment_safe_text", "truncated": truncated},
    )


def _parse_pdf_text(path: Path, max_chars: int) -> AttachmentTextResult:
    from pypdf import PdfReader

    collector = _BoundedText(max_chars)
    reader = PdfReader(str(path))
    pages_read = 0
    for page in reader.pages[:MAX_PDF_PAGES]:
        if collector.full:
            collector.truncated = True
            break
        collector.add(page.extract_text() or "", separator="\n\n")
        pages_read += 1
    if len(reader.pages) > MAX_PDF_PAGES:
        collector.truncated = True
    return AttachmentTextResult(
        text=collector.result(),
        metadata={"parser": "attachment_safe_pdf", "pagesRead": pages_read, "truncated": collector.truncated},
    )


def _parse_docx_text(path: Path, max_chars: int) -> AttachmentTextResult:
    from backend.domains.knowledge_base.docx_parser import parse_docx_blocks

    collector = _BoundedText(max_chars)
    blocks = parse_docx_blocks(path, max_blocks=MAX_DOCX_BLOCKS, max_chars=max_chars)
    blocks_read = 0
    for block in blocks[:MAX_DOCX_BLOCKS]:
        if collector.full:
            collector.truncated = True
            break
        collector.add(block.markdown, separator="\n\n")
        blocks_read += 1
    if len(blocks) > MAX_DOCX_BLOCKS:
        collector.truncated = True
    return AttachmentTextResult(
        text=collector.result(),
        metadata={"parser": "attachment_safe_docx", "blocksRead": blocks_read, "truncated": collector.truncated},
    )


def _parse_pptx_text(path: Path, max_chars: int) -> AttachmentTextResult:
    from pptx import Presentation

    collector = _BoundedText(max_chars)
    with path.open("rb") as source:
        presentation = Presentation(source)
        slides_read = 0
        for slide_index, slide in enumerate(presentation.slides):
            if slide_index >= MAX_PPTX_SLIDES:
                collector.truncated = True
                break
            if collector.full:
                collector.truncated = True
                break
            slides_read += 1
            shapes = iter(islice(slide.shapes, MAX_SHAPES_PER_SLIDE))
            while not collector.full:
                try:
                    shape = next(shapes)
                except StopIteration:
                    break
                if getattr(shape, "has_text_frame", False):
                    collector.add(shape.text)
                if getattr(shape, "has_table", False):
                    for row in islice(shape.table.rows, MAX_XLSX_ROWS_PER_SHEET):
                        collector.add(
                            " | ".join(
                                cell.text for cell in islice(row.cells, MAX_XLSX_CELLS_PER_ROW)
                            )
                        )
                        if collector.full:
                            break
    if len(presentation.slides) > MAX_PPTX_SLIDES:
        collector.truncated = True
    return AttachmentTextResult(
        text=collector.result(),
        metadata={"parser": "attachment_safe_pptx", "slidesRead": slides_read, "truncated": collector.truncated},
    )


def _parse_xlsx_text(path: Path, max_chars: int) -> AttachmentTextResult:
    from openpyxl import load_workbook

    collector = _BoundedText(max_chars)
    cells_read = 0
    sheets_read = 0
    with path.open("rb") as source:
        workbook = load_workbook(source, read_only=True, data_only=True, keep_links=False)
        try:
            for worksheet in workbook.worksheets[:MAX_XLSX_SHEETS]:
                if collector.full or cells_read >= MAX_XLSX_TOTAL_CELLS:
                    collector.truncated = True
                    break
                sheets_read += 1
                collector.add(worksheet.title, separator="\n\n")
                for row in worksheet.iter_rows(
                    max_row=MAX_XLSX_ROWS_PER_SHEET,
                    max_col=MAX_XLSX_CELLS_PER_ROW,
                    values_only=True,
                ):
                    values: list[str] = []
                    for value in row[:MAX_XLSX_CELLS_PER_ROW]:
                        cells_read += 1
                        if value is not None:
                            values.append(str(value))
                        if cells_read >= MAX_XLSX_TOTAL_CELLS:
                            break
                    collector.add(" | ".join(values))
                    if collector.full or cells_read >= MAX_XLSX_TOTAL_CELLS:
                        collector.truncated = True
                        break
        finally:
            workbook.close()
    return AttachmentTextResult(
        text=collector.result(),
        metadata={
            "parser": "attachment_safe_xlsx",
            "sheetsRead": sheets_read,
            "cellsRead": cells_read,
            "truncated": collector.truncated,
        },
    )


def _parse_image_metadata(path: Path, mime_type: str) -> AttachmentTextResult:
    from PIL import Image

    with Image.open(path) as image:
        width, height = image.size
        if width <= 0 or height <= 0 or width * height > MAX_IMAGE_PIXELS:
            raise ValueError("image dimensions exceed attachment parser limits")
        image_format = str(image.format or "").upper()
    return AttachmentTextResult(
        text="",
        metadata={
            "parser": "attachment_safe_image_metadata",
            "mimeType": mime_type,
            "format": image_format,
            "width": width,
            "height": height,
            "truncated": False,
        },
    )
